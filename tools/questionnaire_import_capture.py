#!/usr/bin/env python3
"""Capture questionnaire-import fixtures from the vendored 18.6.6 unit (ADR 0018, chunk 4).

AIA imports a filled-in questionnaire template with
``aia_core.infrastructure.questionnaire_file`` and
``aia_core.domain.questionnaire_import``, a port of the unit's
``ui_server.import_questionnaire_payload`` and the ``normalize_project`` rules it ends
with. These fixtures pin the unit's own answer on fictional files, so the port is
compared with the function it replaces.

Three steps, in two environments:

    python tools/questionnaire_import_capture.py inputs
        (an environment with openpyxl 3.1.5) write the fictional input files to
        packages/aia_core/tests/fixtures/questionnaire_import/inputs/

    python tools/questionnaire_import_capture.py template
        (the repo environment) write AIA's own template as one more input, so the
        capture shows the unit accepts what AIA offers for download

    python tools/questionnaire_import_capture.py capture
        (any Python with the unit's app directory importable: its research_project
        needs only the standard library) run the unit's own, unmodified
        ``import_questionnaire_payload`` (taken from the vendored ``ui_server.py`` by
        its source, with ``_xlsx_rows``; ``research_project`` imported from the unit
        with bytecode writing off) on every input and on the unit's own template,
        read in place; write expected.json and index.json

The workbooks ``inputs`` writes carry openpyxl's timestamps, so it writes new bytes
each time; index.json pins whatever ``capture`` last read, and the test checks the pins.
Nothing here writes under ``legacy/``. No network, no real data.
"""

from __future__ import annotations

import argparse
import ast
import base64
import hashlib
import io
import json
import sys
import zipfile
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
UNIT = ROOT / "legacy" / "npc-panel-18.6.6" / "app"
UNIT_SERVER = UNIT / "ui_server.py"
UNIT_TEMPLATE = UNIT / "NPC_QUESTIONNAIRE_TEMPLATE_v17_6_0.xlsx"
OUT = ROOT / "packages" / "aia_core" / "tests" / "fixtures" / "questionnaire_import"
INPUTS = OUT / "inputs"
UNIT_TEMPLATE_KEY = "@unit/NPC_QUESTIONNAIRE_TEMPLATE_v17_6_0.xlsx"
HEADER = "id,otazka,typ,moznosti,skala_min,skala_max,blok,sledovana_sada,povolit_nevim,poznamka"

CSV_INPUTS: dict[str, str] = {
    # Every row kind, a repeated id, a missing id, an id that is a number, markup in a
    # question, a Czech decimal comma in a scale, a blank row, a block that comes back,
    # and a tracked set whose duplicates leave three objects (a design warning).
    "plny.csv": "\n".join(
        [
            HEADER,
            "Q1,Jak často kupujete ranní nápoj?,vyber,Denně|Týdně|Méně často,,,Nákup,,ano,",
            'Q1,Které příchutě znáte?,multi,"Citron\nMáta\nZázvor",,,Nákup,,0,',
            ",Jak hodnotíte cenu?,skala,,0,5,Cena,,TRUE,",
            "3,Co vám chybí?,otevrena,,,,Cena,,,",
            "OBJ_ZNACKY,Jak blízká je vám značka {object}?,objektova_sada,"
            "Alfa|Beta|Gama|Delta|Epsilon,1,7,Značky,Nápojové značky,,Fiktivní značky",
            'Q6,"Jak <b>důležitá</b> je dostupnost?",skala,,1,"5,5",Nákup,,,',
            ",,,,,,,,,",
            "OBJ2,Jak vnímáte médium,objektova_sada,TV|Rádio|TV|Web,,,Média,,,",
        ]
    )
    + "\n",
    "bom.csv": "﻿" + HEADER + "\nQ1,Souhlasíte?,vyber,Ano|Ne,,,,,,\n",
    "chybi_sloupce.csv": "id,text,typ\nQ1,Otázka,vyber\n",
    "neznamy_typ.csv": HEADER + "\nQ1,Otázka?,matice,,,,,,,\n",
    "mala_sada.csv": HEADER + "\nQ1,Jak {object}?,objektova_sada,A|B|C,,,,,,\n",
    "jedna_moznost.csv": HEADER + "\nQ1,Otázka?,vyber,Jen jedna,,,,,,\n",
    "bez_zneni.csv": HEADER + "\nQ1,,vyber,A|B,,,,,,\n",
    "kolize.csv": HEADER
    + "\nS1,Jak {object}?,objektova_sada,Coca-Cola|Coca Cola|Kofola|Voda,,,,,,\n",
    "znacky_html.csv": HEADER + "\nQ1,Otázka?,vyber,<a>|<b>|Ano,,,,,,\n",
    "sada_duplicity.csv": HEADER + "\nS1,Jak {object}?,objektova_sada,A|A|A|B,,,Blok,,,\n",
    "html_rodina.csv": HEADER + "\nS1,Jak {object}?,objektova_sada,A|B|C|D,,,Blok,<i>,,\n",
    "html_zneni.csv": HEADER + "\nQ1,<br>,skala,,,,,,,\n",
}


def _openpyxl_inputs() -> dict[str, bytes]:
    from openpyxl import Workbook

    def save(wb: Any) -> bytes:
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # DOTAZNIK is not the first sheet; numbers, floats and booleans are typed cells.
    wb = Workbook()
    wb.active.title = "Pokyny"
    wb.active.append(["Tento list se nečte."])
    ws = wb.create_sheet("DOTAZNIK")
    ws.append(HEADER.split(","))
    ws.append(
        [
            "Q1",
            "Jak pravděpodobné je, že nápoj koupíte?",
            "skala",
            None,
            1,
            10.0,
            "Záměr",
            None,
            False,
            None,
        ]
    )
    ws.append(
        [
            "Q2",
            "Kde nakupujete?",
            "multi",
            "Obchod | Online | Automat",
            None,
            None,
            "Nákup",
            None,
            True,
            "Více odpovědí",
        ]
    )
    ws.append([None] * 10)
    ws.append(
        [
            "OBJ",
            "Jak vnímáte {object}?",
            "objektova_sada",
            "Alfa | Beta | Gama | Delta",
            0,
            5,
            "Značky",
            "Značky nápojů",
            None,
            None,
        ]
    )
    ws.append(
        ["Q4", "Co byste zlepšili?", "otevrena", None, None, None, "Závěr", None, "yes", None]
    )
    typed = save(wb)

    # No DOTAZNIK sheet: the first one is read.
    first = Workbook()
    first.active.title = "List1"
    first.active.append(HEADER.split(","))
    first.active.append(["Q1", "Znáte nás?", "vyber", "Ano|Ne", None, None, None, None, None, None])
    return {"typovany.xlsx": typed, "prvni_list.xlsx": save(first)}


def _no_sheet_xlsx() -> bytes:
    main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rels = "http://schemas.openxmlformats.org/package/2006/relationships"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, body in {
            "xl/workbook.xml": f'<workbook xmlns="{main}"><sheets/></workbook>',
            "xl/_rels/workbook.xml.rels": f'<Relationships xmlns="{rels}"/>',
        }.items():
            info = zipfile.ZipInfo(name, date_time=(2026, 9, 27, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(info, body)
    return buf.getvalue()


def inputs() -> None:
    INPUTS.mkdir(parents=True, exist_ok=True)
    files: dict[str, bytes] = {name: text.encode("utf-8") for name, text in CSV_INPUTS.items()}
    files.update(_openpyxl_inputs())
    files.update(
        {
            "bez_listu.xlsx": _no_sheet_xlsx(),
            "rozbity.xlsx": b"this is not a workbook",
            "latin2.csv": (HEADER + "\nQ1,Líbí se vám?,vyber,Ano|Ne,,,,,,\n").encode("cp1250"),
            "dotaznik.txt": b"id,otazka,typ\n",
            "jen_bom.csv": b"\xef\xbb\xbf",
            "nahrajte.csv": b"",
        }
    )
    for name, data in files.items():
        (INPUTS / name).write_bytes(data)
        print(f"{name}: {len(data)} bytes")


def template() -> None:
    sys.path.insert(0, str(ROOT / "packages" / "aia_core" / "src"))
    from aia_core.infrastructure.questionnaire_file import template_xlsx

    INPUTS.mkdir(parents=True, exist_ok=True)
    (INPUTS / "aia_sablona.xlsx").write_bytes(template_xlsx())
    print("aia_sablona.xlsx written")


# -------------------------------------------------------------------- capture --


def _unit_functions() -> tuple[Any, dict[str, str]]:
    """The unit's import function, compiled from its own source, and the sources' hashes."""
    sys.dont_write_bytecode = True  # the unit's tree is frozen: no __pycache__ in it
    sys.path.insert(0, str(UNIT))
    import research_project

    source = UNIT_SERVER.read_text(encoding="utf-8")
    tree = ast.parse(source)
    wanted = {"_xlsx_rows", "import_questionnaire_payload"}
    segments = {
        n.name: ast.get_source_segment(source, n)
        for n in tree.body
        if isinstance(n, ast.FunctionDef) and n.name in wanted
    }
    assert set(segments) == wanted, segments.keys()
    namespace: dict[str, Any] = {
        "base64": base64,
        "Path": Path,
        "normalize_project": research_project.normalize_project,
        "empty_project": research_project.empty_project,
    }
    for name in ("_xlsx_rows", "import_questionnaire_payload"):
        segment = segments[name]
        assert segment is not None
        # The unit's own functions, compiled from its own source: nothing is rewritten.
        exec(compile(segment, str(UNIT_SERVER), "exec"), namespace)
    hashes = {
        f"ui_server.py::{name}": hashlib.sha256(str(segments[name]).encode()).hexdigest()
        for name in sorted(wanted)
    }
    for module in ("research_project.py", "study_contract.py", "product_policy.py"):
        hashes[module] = hashlib.sha256((UNIT / module).read_bytes()).hexdigest()
    hashes["PRODUCT_POLICY.json"] = hashlib.sha256(
        (UNIT / "PRODUCT_POLICY.json").read_bytes()
    ).hexdigest()
    return namespace["import_questionnaire_payload"], hashes


def _run(function: Any, name: str, data: bytes) -> dict[str, Any]:
    filename = name.split("/")[-1]
    try:
        result = function({"filename": filename, "data_b64": base64.b64encode(data).decode()})
    except Exception as exc:
        return {"ok": False, "error_type": type(exc).__name__, "error": str(exc)}
    return {
        "ok": True,
        "sections": result["project"]["sections"],
        "summary": result["summary"],
        "filename": result["filename"],
        "plan_status": result["project"]["research_plan"]["status"],
    }


def capture() -> None:
    function, sources = _unit_functions()
    files = {p.name: p.read_bytes() for p in sorted(INPUTS.iterdir()) if p.is_file()}
    files[UNIT_TEMPLATE_KEY] = UNIT_TEMPLATE.read_bytes()
    expected = {name: _run(function, name, data) for name, data in sorted(files.items())}
    (OUT / "expected.json").write_text(
        json.dumps(expected, ensure_ascii=False, indent=1, sort_keys=True) + "\n", encoding="utf-8"
    )
    index = {
        "unit_sources": sources,
        "inputs": {name: hashlib.sha256(data).hexdigest() for name, data in sorted(files.items())},
        "expected_sha256": hashlib.sha256((OUT / "expected.json").read_bytes()).hexdigest(),
        "python": sys.version.split()[0],
    }
    (OUT / "index.json").write_text(json.dumps(index, indent=1, sort_keys=True) + "\n")
    for name, outcome in expected.items():
        print(
            f"{name}: {'ok' if outcome['ok'] else outcome['error_type'] + ': ' + outcome['error']}"
        )


def main() -> None:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("command", choices=["inputs", "template", "capture"])
    args = parser.parse_args()
    {"inputs": inputs, "template": template, "capture": capture}[args.command]()


if __name__ == "__main__":
    main()
