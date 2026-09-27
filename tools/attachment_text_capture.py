#!/usr/bin/env python3
"""Capture attachment-text fixtures from the vendored 18.6.6 unit (ADR 0018, phase-out chunk 3).

AIA reads an attached file's text with ``aia_core.infrastructure.document_text``, a port
of the unit's ``data_library._extract_text``. The fixtures pin the unit's own answer on
fictional documents, so the port is compared with the function it replaces.

Two steps, both in an environment with the unit's document libraries (openpyxl 3.1.5,
python-docx, pypdf -- ``pip install "openpyxl==3.1.5" "python-docx>=1.1,<2"
"pypdf>=6.1,<7"``), because the unit reads with them and so does AIA:

    python tools/attachment_text_capture.py inputs
        write the fictional input documents to
        packages/aia_core/tests/fixtures/attachment_text/inputs/

    python tools/attachment_text_capture.py capture
        run the unit's own, unmodified ``_extract_text`` (taken from the vendored
        ``data_library.py`` by its source, never imported) on every input; write
        expected.json (the unit's output per file) and index.json (the SHA256 of every
        input, of the function's source, and the library versions)

The Word and Excel inputs carry their writer's timestamps, so ``inputs`` writes new
bytes each time; index.json pins whatever ``capture`` last read, and the test checks
the pins. Nothing here imports AIA or writes under ``legacy/``. No network, no real
data.
"""

from __future__ import annotations

import argparse
import ast
import datetime
import hashlib
import io
import json
import sys
import zipfile
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
UNIT_SOURCE = ROOT / "legacy" / "npc-panel-18.6.6" / "app" / "data_library.py"
OUT = ROOT / "packages" / "aia_core" / "tests" / "fixtures" / "attachment_text"
INPUTS = OUT / "inputs"


# --------------------------------------------------------------------- inputs --


def _pdf(pages: list[list[str]]) -> bytes:
    """A minimal PDF, one Helvetica text line per string (Latin-1, octal-escaped)."""

    def literal(text: str) -> str:
        out = []
        for byte in text.encode("latin-1"):
            ch = chr(byte)
            if ch in "()\\":
                out.append("\\" + ch)
            elif 32 <= byte < 127:
                out.append(ch)
            else:
                out.append(f"\\{byte:03o}")
        return "".join(out)

    objects: list[bytes] = []
    font_id = 3 + 2 * len(pages)
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    for i, lines in enumerate(pages):
        body = ["BT", "/F1 14 Tf", "72 760 Td"]
        for j, line in enumerate(lines):
            if j:
                body.append("0 -22 Td")
            body.append(f"({literal(line)}) Tj")
        body.append("ET")
        stream = "\n".join(body).encode("latin-1")
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {4 + 2 * i} 0 R >>".encode()
        )
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream")
    objects.append(
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"
    )
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for n, obj in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{n} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def _docx() -> bytes:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.oxml.ns import qn
    from docx.oxml.parser import OxmlElement

    doc = Document()
    doc.add_heading("Zadání výzkumu: fiktivní ranní nápoj", level=1)
    doc.add_paragraph("Klient zvažuje uvedení nového nápoje na trh v Česku.")
    p = doc.add_paragraph("Cíl:\tzjistit zájem")
    p.add_run().add_break()
    p.add_run("a cenovou citlivost.")
    p.add_run().add_break(WD_BREAK.PAGE)
    p.add_run("Po zalomení stránky.")
    doc.add_paragraph("")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Tabulka se nečte"
    table.cell(1, 1).text = "ani tato buňka"
    linked = doc.add_paragraph("Zdroj: ")
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), "zdroj")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "fiktivní studie 2026"
    run.append(text)
    link.append(run)
    linked._p.append(link)
    doc.add_paragraph("Konec zadání – bez-zalomení")  # noqa: RUF001 - Czech typography, on purpose
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Přehled"
    ws.append(["Značka", "Podíl", "Cena", "Nový", "Uvedeno", "Čas"])
    ws.append(["Alfa", 0.25, 39, True, datetime.datetime(2026, 3, 1, 8, 30), datetime.time(7, 15)])
    ws.append(["Beta", 0.1, 42.5, False, datetime.date(2025, 11, 20), None])
    ws.append([])
    ws.cell(row=5, column=1, value="Gama")
    ws.cell(row=5, column=3, value=1e-05)
    ws.cell(row=6, column=1, value=12345678901234567890)
    ws.append([f"S{i}" for i in range(1, 36)])
    long = wb.create_sheet("Dlouhý")
    for i in range(1, 306):
        long.append([i, f"řádek {i}"])
    formulas = wb.create_sheet("Vzorce")
    formulas["A1"] = 2
    formulas["B1"] = "=A1*2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_UNSIZED_SHEET = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
    '<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>Bez rozměru</t></is></c>'
    '<c r="B1"><v>7</v></c></row></sheetData></worksheet>'
)


def _unsized_xlsx() -> bytes:
    """A workbook whose sheet declares no <dimension>: the unit keeps no text of it."""
    xml = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    package = "http://schemas.openxmlformats.org/package/2006"
    office = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    sheetml = "application/vnd.openxmlformats-officedocument.spreadsheetml"
    rels = "application/vnd.openxmlformats-package.relationships+xml"
    files = {
        "[Content_Types].xml": (
            f'{xml}<Types xmlns="{package}/content-types">'
            f'<Default Extension="rels" ContentType="{rels}"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            f'<Override PartName="/xl/workbook.xml" ContentType="{sheetml}.sheet.main+xml"/>'
            '<Override PartName="/xl/worksheets/sheet1.xml" '
            f'ContentType="{sheetml}.worksheet+xml"/>'
            "</Types>"
        ),
        "_rels/.rels": (
            f'{xml}<Relationships xmlns="{package}/relationships">'
            f'<Relationship Id="rId1" Type="{office}/officeDocument" Target="xl/workbook.xml"/>'
            "</Relationships>"
        ),
        "xl/workbook.xml": (
            f'{xml}<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            f'xmlns:r="{office}">'
            '<sheets><sheet name="List1" sheetId="1" r:id="rId1"/></sheets></workbook>'
        ),
        "xl/_rels/workbook.xml.rels": (
            f'{xml}<Relationships xmlns="{package}/relationships">'
            f'<Relationship Id="rId1" Type="{office}/worksheet" Target="worksheets/sheet1.xml"/>'
            "</Relationships>"
        ),
        "xl/worksheets/sheet1.xml": _UNSIZED_SHEET,
    }
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, body in files.items():
            info = zipfile.ZipInfo(name, date_time=(2026, 9, 27, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(info, body)
    return buf.getvalue()


# One PNG pixel: a kind of file with no text reader.
_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010806000000"
    "1f15c4890000000d49444154789c6360f8cfc0000003010100c9fe92ef0000000049454e44ae426082"
)


def inputs() -> None:
    INPUTS.mkdir(parents=True, exist_ok=True)
    xlsx = _xlsx()
    documents = {
        "zadani.txt": "Fiktivní zadání: ranní nápoj, 18–65 let.\nDruhý řádek.\n".encode(),  # noqa: RUF001
        "neplatne.txt": b"Neplatn\xe9 UTF-8 \xff\xfe konec",
        "poznamky.md": "# Poznámky\n\n- bod *jedna*\n- bod **dva**\n".encode(),
        "data.csv": b"znacka;podil\nAlfa;0,25\nBeta;0,10\n",
        "zadani.docx": _docx(),
        "tabulka.xlsx": xlsx,
        "kopie.xlsm": xlsx,
        "bez_rozmeru.xlsx": _unsized_xlsx(),
        "prezentace.pdf": _pdf(
            [
                ["Zadání výzkumu", "Cílová skupina: lidé 18-65 let"],
                ["Strana 2", "Rozsah (orientace): 250 000 CZK"],
            ]
        ),
        "poskozeny.docx": b"this is not a zip archive",
        "poskozeny.pdf": b"%PDF-1.4 truncated",
        "obrazek.png": _PNG,
    }
    for name, data in documents.items():
        (INPUTS / name).write_bytes(data)
        print(f"{name}: {len(data)} bytes")


# -------------------------------------------------------------------- capture --


def _unit_function() -> tuple[Any, str]:
    """The unit's ``_extract_text``, compiled from its own source, and that source's hash."""
    source = UNIT_SOURCE.read_text(encoding="utf-8")
    tree = ast.parse(source)
    node = next(
        n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == "_extract_text"
    )
    segment = ast.get_source_segment(source, node)
    assert segment is not None
    namespace: dict[str, Any] = {"io": io, "Path": Path}
    # The unit's own function, compiled from its own source: nothing of it is rewritten.
    exec(compile(segment, str(UNIT_SOURCE), "exec"), namespace)
    return namespace["_extract_text"], hashlib.sha256(segment.encode("utf-8")).hexdigest()


def capture() -> None:
    import docx
    import openpyxl
    import pypdf

    extract, source_sha = _unit_function()
    names = sorted(p.name for p in INPUTS.iterdir() if p.is_file())
    expected = {name: extract((INPUTS / name).read_bytes(), name) for name in names}
    (OUT / "expected.json").write_text(
        json.dumps(expected, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    index = {
        "unit_function": "legacy/npc-panel-18.6.6/app/data_library.py::_extract_text",
        "unit_function_sha256": source_sha,
        "libraries": {
            "openpyxl": openpyxl.__version__,
            "pypdf": pypdf.__version__,
            "python-docx": getattr(docx, "__version__", "unknown"),
            "python": sys.version.split()[0],
        },
        "inputs": {
            name: hashlib.sha256((INPUTS / name).read_bytes()).hexdigest() for name in names
        },
        "expected_sha256": hashlib.sha256((OUT / "expected.json").read_bytes()).hexdigest(),
    }
    (OUT / "index.json").write_text(json.dumps(index, indent=2, sort_keys=True) + "\n")
    for name in names:
        text = expected[name]
        failed = " (failed)" if text.startswith("[TEXT_EXTRACTION_FAILED:") else ""
        print(f"{name}: {len(text)} characters{failed}")


def main() -> None:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("command", choices=["inputs", "capture"])
    args = parser.parse_args()
    {"inputs": inputs, "capture": capture}[args.command]()


if __name__ == "__main__":
    main()
