"""Read a client's source into segments with locators (lifecycle K2a).

:func:`read_source` turns one original's bytes into a
:class:`~aia_core.domain.client_sources.SourceReading`: every readable part with the
place it stands in the original, and an outcome that never claims more than happened.
It is a pure function of the bytes and the name, so where it runs (the API at upload
today, a worker later) changes no contract.

The libraries are the ``documents`` extra (pypdf, pdfplumber, python-docx, olefile), imported
only when a file of their kind is read. PPTX and ODT are read from their XML parts with
the stdlib, a part declaring a DTD refused; a Word 97-2003 DOC from its piece table, by
AIA's own code over the compound file olefile opens. A PDF is read
with :func:`document_text.read_pdf`, the Deep Research reader: its 400-page refusal, its
password and malformed-file refusals and its bounded table extraction. An Office file is
checked against :func:`document_text.check_zip_bounds` before anything in it is inflated.

A reader not in AIA yet reads as ``UNSUPPORTED`` (``reader_pending``), an image or a
video as ``UNREADABLE``: both are kept, and neither poses as read.
"""

from __future__ import annotations

import io
import struct
import xml.etree.ElementTree as ET
import zipfile
from typing import Final

from ..domain.attachments import extension_of
from ..domain.client_datasets import DatasetPreview, preview_sheet
from ..domain.client_sources import (
    Locator,
    ReadingNote,
    ReadingOutcome,
    Segment,
    SourceFamily,
    SourceReading,
    family_of,
    reading_of_segments,
)
from ..domain.deep_research import documents as grid_bounds
from ..domain.deep_research.documents import DocumentRefused, GridContent
from .document_text import check_zip_bounds, grid_from_rows, read_csv, read_pdf, read_xlsx

__all__ = ["READERS", "XML_PART_MAX_BYTES", "piece_text", "read_source"]

#: The reader each readable format goes through, and its version. A new version is a
#: new name: a reading records which reader made it.
READERS: Final[dict[str, str]] = {
    ".pdf": "pdf/1",
    ".docx": "docx/1",
    ".doc": "doc/1",
    ".odt": "odt/1",
    ".pptx": "pptx/1",
    ".txt": "text/1",
    ".md": "text/1",
    ".csv": "csv/1",
    ".xlsx": "xlsx/1",
    ".xls": "xls/1",
    ".ods": "ods/1",
}

#: The most one XML part of an Office or OpenDocument file may hold once inflated.
XML_PART_MAX_BYTES: Final = 50 * 1024 * 1024


def _cells(values: tuple[str, ...] | list[str]) -> str:
    return " | ".join(v.strip() for v in values)


def _pdf(data: bytes) -> tuple[list[Segment], list[ReadingNote]]:
    content = read_pdf(data)
    segments = [
        Segment(locator=Locator(page=number), text=text)
        for number, text in enumerate(content.pages, start=1)
    ]
    notes: list[ReadingNote] = []
    # A page without text is said only when others have it: a PDF with none at all is
    # a scan, and the reading says that once (``no_text``), not once a page.
    if any(text.strip() for text in content.pages):
        notes.extend(
            ReadingNote(code="page_without_text", at=Locator(page=number))
            for number, text in enumerate(content.pages, start=1)
            if not text.strip()
        )
    for index, table in enumerate(content.tables, start=1):
        segments.extend(
            Segment(
                locator=Locator(page=table.page, table=index, row=row.number),
                text=_cells(row.values),
            )
            for row in table.rows
        )
    if content.tables_truncated:
        notes.append(ReadingNote(code="tables_bound", count=len(content.tables)))
    return segments, notes


def _docx(data: bytes) -> tuple[list[Segment], list[ReadingNote]]:
    check_zip_bounds(data, "DOCX")
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentRefused(
            f"the DOCX cannot be read ({type(exc).__name__})", reason="document_malformed"
        ) from exc
    segments: list[Segment] = []
    paragraph = table = 0
    # The body in its own order: a paragraph, then a table, as a reader meets them.
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph += 1
            text = Paragraph(child, document).text
            segments.append(Segment(locator=Locator(paragraph=paragraph), text=text))
        elif tag == "tbl":
            table += 1
            for number, row in enumerate(Table(child, document).rows, start=1):
                segments.append(
                    Segment(
                        locator=Locator(table=table, row=number),
                        text=_cells([cell.text for cell in row.cells]),
                    )
                )
    return segments, []


def _text(data: bytes) -> tuple[list[Segment], list[ReadingNote]]:
    notes: list[ReadingNote] = []
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("utf-8-sig", errors="replace")
        notes.append(ReadingNote(code="not_utf8"))
    # A paragraph is a run of lines; its locator is the line it starts on.
    segments: list[Segment] = []
    start: int | None = None
    lines: list[str] = []
    for number, line in enumerate(text.splitlines(), start=1):
        if line.strip():
            if start is None:
                start = number
            lines.append(line)
        elif start is not None:
            segments.append(Segment(locator=Locator(line=start), text="\n".join(lines)))
            start, lines = None, []
    if start is not None:
        segments.append(Segment(locator=Locator(line=start), text="\n".join(lines)))
    return segments, notes


# --------------------------------------------------------------------------- #
# XML parts of a ZIP: PPTX and ODT
# --------------------------------------------------------------------------- #

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_PKG_REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"
_TEXT = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
_TABLE = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
_OFFICE = "{urn:oasis:names:tc:opendocument:xmlns:office:1.0}"


def _part(archive: zipfile.ZipFile, name: str, kind: str) -> ET.Element:
    """One XML part, parsed; a part with a DTD is refused before it is parsed.

    Neither OOXML nor OpenDocument ever carries a document type declaration, so one
    is refused outright: no entity can be declared, and none expanded.
    """
    try:
        info = archive.getinfo(name)
    except KeyError as exc:
        raise DocumentRefused(f"the {kind} has no {name}", reason="document_malformed") from exc
    if info.file_size > XML_PART_MAX_BYTES:
        raise DocumentRefused(f"the {kind}'s {name} is too large", reason="document_zip_bounds")
    data = archive.read(name)
    if b"<!DOCTYPE" in data or b"<!ENTITY" in data:
        raise DocumentRefused(f"the {kind}'s {name} declares a DTD", reason="document_malformed")
    try:
        return ET.fromstring(data)
    except ET.ParseError as exc:
        raise DocumentRefused(
            f"the {kind}'s {name} is not XML", reason="document_malformed"
        ) from exc


def _pptx(data: bytes) -> tuple[list[Segment], list[ReadingNote]]:
    """Each slide's text in the presentation's own order, a paragraph a line."""
    check_zip_bounds(data, "PPTX")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        presentation = _part(archive, "ppt/presentation.xml", "PPTX")
        rels = _part(archive, "ppt/_rels/presentation.xml.rels", "PPTX")
        targets = {r.get("Id"): r.get("Target", "") for r in rels.iter(f"{_PKG_REL}Relationship")}
        segments: list[Segment] = []
        for number, slide in enumerate(presentation.iter(f"{_P}sldId"), start=1):
            target = targets.get(slide.get(f"{_R}id"), "")
            name = "ppt/" + target.lstrip("/").removeprefix("ppt/")
            root = _part(archive, name, "PPTX")
            lines = ["".join(t.text or "" for t in p.iter(f"{_A}t")) for p in root.iter(f"{_A}p")]
            text = "\n".join(line for line in lines if line.strip())
            segments.append(Segment(locator=Locator(slide=number), text=text))
    return segments, []


def _odf_text(element: ET.Element) -> str:
    """A paragraph's text as written: spaces (``text:s``), tabs and line breaks kept."""
    out: list[str] = [element.text or ""]
    for child in element:
        tag = child.tag
        if tag == f"{_TEXT}s":
            out.append(" " * int(child.get(f"{_TEXT}c", "1") or 1))
        elif tag == f"{_TEXT}tab":
            out.append("\t")
        elif tag == f"{_TEXT}line-break":
            out.append("\n")
        elif tag in (f"{_TEXT}note", f"{_OFFICE}annotation"):
            pass  # a footnote or a comment is not the paragraph's text
        else:
            out.append(_odf_text(child))
        out.append(child.tail or "")
    return "".join(out)


def _odt(data: bytes) -> tuple[list[Segment], list[ReadingNote]]:
    """Headings and paragraphs numbered in body order, tables by row."""
    check_zip_bounds(data, "ODT")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        root = _part(archive, "content.xml", "ODT")
    body = root.find(f"{_OFFICE}body/{_OFFICE}text")
    if body is None:
        raise DocumentRefused("the ODT has no text body", reason="document_malformed")
    segments: list[Segment] = []
    counts = {"paragraph": 0, "table": 0}

    def walk(element: ET.Element) -> None:
        for child in element:
            if child.tag in (f"{_TEXT}p", f"{_TEXT}h"):
                counts["paragraph"] += 1
                segments.append(
                    Segment(locator=Locator(paragraph=counts["paragraph"]), text=_odf_text(child))
                )
            elif child.tag == f"{_TABLE}table":
                counts["table"] += 1
                rows = list(child.iter(f"{_TABLE}table-row"))
                for number, row in enumerate(rows, start=1):
                    cells = [
                        "\n".join(_odf_text(p) for p in cell.iter(f"{_TEXT}p"))
                        for cell in row.iter(f"{_TABLE}table-cell")
                    ]
                    segments.append(
                        Segment(
                            locator=Locator(table=counts["table"], row=number), text=_cells(cells)
                        )
                    )
            elif child.tag in (f"{_TEXT}list", f"{_TEXT}list-item", f"{_TEXT}section"):
                walk(child)

    walk(body)
    return segments, []


# --------------------------------------------------------------------------- #
# Word 97-2003 (.doc)
# --------------------------------------------------------------------------- #

#: The FIB's identifier, its flags and where the counts and the piece table are
#: ([MS-DOC] 2.5.1-2.5.6): fixed for every Word 97 and later file (nFib >= 0x00C1).
_WORD_IDENT: Final = 0xA5EC
_NFIB_WORD97: Final = 0x00C1
_F_ENCRYPTED: Final = 0x0100
_F_WHICH_TABLE: Final = 0x0200
_OFF_FLAGS: Final = 0x000A
_OFF_CCP_TEXT: Final = 0x004C
_OFF_FC_CLX: Final = 0x01A2
# Characters a paragraph ends on: the paragraph mark, and the cell or row-end mark.
_PARAGRAPH_ENDS: Final = frozenset("\r\x07")
# What a character means in the main text that is not itself text ([MS-DOC] 2.8.25).
_DROPPED: Final = frozenset("\x01\x02\x03\x04\x05\x08\x1f")
_REPLACED: Final = {"\x0b": "\n", "\x0c": "\n", "\x1e": "-", "\x0e": "\n"}


def _doc_text(data: bytes) -> str:
    """The main document's characters, from the compound file's two streams."""
    import olefile

    if not olefile.isOleFile(data):
        raise DocumentRefused("the DOC is not a compound file", reason="document_malformed")
    with olefile.OleFileIO(data) as ole:
        if not ole.exists("WordDocument"):
            raise DocumentRefused("the DOC has no WordDocument stream", reason="document_malformed")
        word = ole.openstream("WordDocument").read()
        table_name = "1Table" if _fib_flags(word) & _F_WHICH_TABLE else "0Table"
        if not ole.exists(table_name):
            raise DocumentRefused(
                f"the DOC has no {table_name} stream", reason="document_malformed"
            )
        table = ole.openstream(table_name).read()
    return piece_text(word, table)


def _fib_flags(word: bytes) -> int:
    """The FIB's flags, once its identity and version are checked."""
    if len(word) < _OFF_FC_CLX + 8:
        raise DocumentRefused("the DOC's FIB is cut short", reason="document_malformed")
    ident, nfib = struct.unpack_from("<HH", word, 0)
    if ident != _WORD_IDENT:
        raise DocumentRefused("the DOC is not a Word document", reason="document_malformed")
    if nfib < _NFIB_WORD97:
        raise DocumentRefused("the DOC predates Word 97", reason="document_too_old")
    (flags,) = struct.unpack_from("<H", word, _OFF_FLAGS)
    if flags & _F_ENCRYPTED:
        raise DocumentRefused("the DOC is encrypted", reason="document_encrypted")
    return int(flags)


def piece_text(word: bytes, table: bytes) -> str:
    """The main text from a WordDocument stream and its table stream ([MS-DOC] 2.9.38).

    The FIB gives the main text's length (``ccpText``) and where the CLX is; its piece
    table maps character positions to bytes: compressed pieces are one cp1252 byte a
    character at ``fc / 2``, the others UTF-16LE at ``fc``. Anything out of bounds is
    ``document_malformed``.
    """
    _fib_flags(word)
    (ccp_text,) = struct.unpack_from("<i", word, _OFF_CCP_TEXT)
    fc_clx, lcb_clx = struct.unpack_from("<II", word, _OFF_FC_CLX)
    if ccp_text < 0 or fc_clx + lcb_clx > len(table):
        raise DocumentRefused("the DOC's piece table is out of bounds", reason="document_malformed")
    clx = table[fc_clx : fc_clx + lcb_clx]
    at = 0
    while at + 3 <= len(clx) and clx[at] == 0x01:  # Prc: property modifiers, skipped
        (cb,) = struct.unpack_from("<H", clx, at + 1)
        at += 3 + cb
    if at + 5 > len(clx) or clx[at] != 0x02:
        raise DocumentRefused("the DOC has no piece table", reason="document_malformed")
    (lcb,) = struct.unpack_from("<I", clx, at + 1)
    plc = clx[at + 5 : at + 5 + lcb]
    count = (len(plc) - 4) // 12
    if count <= 0 or len(plc) < 4 * (count + 1) + 8 * count:
        return ""
    cps = struct.unpack_from(f"<{count + 1}I", plc, 0)
    out: list[str] = []
    wanted = ccp_text
    for i in range(count):
        if wanted <= 0:
            break
        (fc_raw,) = struct.unpack_from("<I", plc, 4 * (count + 1) + 8 * i + 2)
        chars = min(cps[i + 1] - cps[i], wanted)
        if chars <= 0:
            continue
        if fc_raw & 0x40000000:
            start = (fc_raw & 0x3FFFFFFF) // 2
            text = word[start : start + chars].decode("cp1252", errors="replace")
        else:
            start = fc_raw & 0x3FFFFFFF
            text = word[start : start + 2 * chars].decode("utf-16-le", errors="replace")
        if len(text) < chars:
            raise DocumentRefused("a piece of the DOC is cut short", reason="document_malformed")
        out.append(text)
        wanted -= chars
    return "".join(out)


def _doc(data: bytes) -> tuple[list[Segment], list[ReadingNote]]:
    """Paragraphs as Word counts them: a table's cell and row-end marks end one too.

    A field's code (``\x13`` … ``\x14``) is dropped and its result kept, so a link
    reads as its words. A table's rows are not given: Word tells an empty cell from a
    row's end only in the paragraph's properties, so every cell is the paragraph it
    is in Word, numbered as such, and no row number is guessed.
    """
    segments: list[Segment] = []
    fields: list[bool] = []  # one per open field: True while in its code
    current: list[str] = []
    number = 0
    for ch in _doc_text(data):
        if ch == "\x13":
            fields.append(True)
            continue
        if ch == "\x14":
            if fields:
                fields[-1] = False
            continue
        if ch == "\x15":
            if fields:
                fields.pop()
            continue
        if any(fields):
            continue
        if ch in _PARAGRAPH_ENDS:
            number += 1
            segments.append(Segment(locator=Locator(paragraph=number), text="".join(current)))
            current = []
        elif ch in _DROPPED:
            continue
        else:
            current.append(_REPLACED.get(ch, ch))
    if current:
        number += 1
        segments.append(Segment(locator=Locator(paragraph=number), text="".join(current)))
    return segments, []


# --------------------------------------------------------------------------- #
# Datasets: CSV, XLSX, XLS, ODS (K2c)
# --------------------------------------------------------------------------- #

_TABLE_NS = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"


def _xls(data: bytes) -> tuple[tuple[GridContent, ...], bool]:
    """A Word-era workbook's sheets, through xlrd, as the other readers' grids."""
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=data, on_demand=True)
    except xlrd.biffh.XLRDError as exc:
        if "encrypted" in str(exc).lower():
            raise DocumentRefused("the XLS is encrypted", reason="document_encrypted") from exc
        raise DocumentRefused(
            f"the XLS cannot be read ({exc})", reason="document_malformed"
        ) from exc
    try:
        budget = [grid_bounds.MAX_GRID_CELLS]
        grids: list[GridContent] = []
        names = book.sheet_names()
        for index in range(min(len(names), grid_bounds.MAX_GRID_SHEETS)):
            sheet = book.sheet_by_index(index)
            cut = (
                sheet.nrows > grid_bounds.MAX_GRID_ROWS
                or sheet.ncols > grid_bounds.MAX_GRID_COLUMNS
            )
            rows_wanted = min(sheet.nrows, grid_bounds.MAX_GRID_ROWS)
            cols_wanted = min(sheet.ncols, grid_bounds.MAX_GRID_COLUMNS)

            def values(
                r: int, sheet: object = sheet, cols: int = cols_wanted
            ) -> tuple[object, ...]:
                out: list[object] = []
                for c in range(cols):
                    cell = sheet.cell(r, c)  # type: ignore[attr-defined]
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        out.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        out.append(bool(cell.value))
                    elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                        out.append(None)
                    else:
                        out.append(cell.value)
                return tuple(out)

            grids.append(
                grid_from_rows(names[index], (values(r) for r in range(rows_wanted)), budget, cut)
            )
            book.unload_sheet(index)
            if budget[0] <= 0:
                break
        return tuple(grids), len(names) > len(grids)
    finally:
        book.release_resources()


def _ods_cell(cell: ET.Element) -> object:
    kind = cell.get("{urn:oasis:names:tc:opendocument:xmlns:office:1.0}value-type")
    office = "{urn:oasis:names:tc:opendocument:xmlns:office:1.0}"
    if kind in ("float", "percentage", "currency"):
        raw = cell.get(f"{office}value")
        try:
            return float(raw) if raw is not None else None
        except ValueError:
            return raw
    if kind == "date":
        return cell.get(f"{office}date-value")
    if kind == "time":
        return cell.get(f"{office}time-value")
    if kind == "boolean":
        return cell.get(f"{office}boolean-value") == "true"
    text = "\n".join(_odf_text(p) for p in cell.iter(f"{_TEXT}p"))
    return text or None


def _ods(data: bytes) -> tuple[tuple[GridContent, ...], bool]:
    """An OpenDocument spreadsheet's sheets, repeats expanded only up to the bounds.

    A row or cell repeated a million times (how a spreadsheet stores the empty rest of a
    sheet) is never materialized past the bounds: an empty repeat only moves the count.
    """
    check_zip_bounds(data, "ODS")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        root = _part(archive, "content.xml", "ODS")
    body = root.find(f"{_OFFICE}body/{_OFFICE}spreadsheet")
    if body is None:
        raise DocumentRefused("the ODS has no spreadsheet body", reason="document_malformed")
    tables = body.findall(f"{_TABLE_NS}table")
    budget = [grid_bounds.MAX_GRID_CELLS]
    grids: list[GridContent] = []
    for table in tables[: grid_bounds.MAX_GRID_SHEETS]:
        cut = False

        def rows(table: ET.Element = table) -> list[tuple[object, ...]]:
            nonlocal cut
            out: list[tuple[object, ...]] = []
            for row in table.iter(f"{_TABLE_NS}table-row"):
                values: list[object] = []
                for cell in row:
                    if cell.tag not in (f"{_TABLE_NS}table-cell", f"{_TABLE_NS}covered-table-cell"):
                        continue
                    repeat = int(cell.get(f"{_TABLE_NS}number-columns-repeated", "1") or 1)
                    value = _ods_cell(cell)
                    room = grid_bounds.MAX_GRID_COLUMNS + 1 - len(values)
                    if value is None and repeat > room:
                        repeat = 0  # the empty rest of the row: nothing to keep
                    if repeat > room:
                        cut = True
                    values.extend([value] * max(0, min(repeat, room)))
                while values and values[-1] is None:
                    values.pop()
                repeat_rows = int(row.get(f"{_TABLE_NS}number-rows-repeated", "1") or 1)
                if not values:
                    # Empty rows only move the numbering, however many.
                    out.extend([()] * min(repeat_rows, grid_bounds.MAX_GRID_ROWS + 1 - len(out)))
                else:
                    room = grid_bounds.MAX_GRID_ROWS + 1 - len(out)
                    if repeat_rows > room:
                        cut = True
                    out.extend([tuple(values)] * max(0, min(repeat_rows, room)))
                if len(out) > grid_bounds.MAX_GRID_ROWS:
                    cut = True
                    break
            return out

        read_rows = rows()
        name = table.get(f"{_TABLE_NS}name", "") or ""
        grids.append(grid_from_rows(name, read_rows, budget, cut))
        if budget[0] <= 0:
            break
    return tuple(grids), len(tables) > len(grids)


def _xlsx(data: bytes) -> tuple[tuple[GridContent, ...], bool]:
    _, grids, more = read_xlsx(data)
    return grids, more


def _csv(data: bytes) -> tuple[tuple[GridContent, ...], bool]:
    return (read_csv(data),), False


def _trimmed(values: tuple[str, ...]) -> tuple[str, ...]:
    end = len(values)
    while end and not values[end - 1].strip():
        end -= 1
    return values[:end]


_DATASETS = {".csv": _csv, ".xlsx": _xlsx, ".xls": _xls, ".ods": _ods}


def _dataset_reading(reader: str, grids: tuple[GridContent, ...], more: bool) -> SourceReading:
    """A dataset's rows as segments (sheet and row) and its preview, from its grids."""
    segments: list[Segment] = []
    notes: list[ReadingNote] = []
    sheets = []
    for grid in grids:
        name = grid.name or None
        # A sheet that declares more columns than it fills pads every row with empties:
        # a row ends at its last filled cell.
        rows = [(row.number, _trimmed(row.values)) for row in grid.rows]
        sheets.append(preview_sheet(name, rows, truncated=grid.truncated))
        segments.extend(
            Segment(locator=Locator(sheet=name, row=number), text=_cells(list(values)))
            for number, values in rows
        )
        if grid.truncated:
            notes.append(ReadingNote(code="rows_bound", at=Locator(sheet=name) if name else None))
    if more:
        notes.append(ReadingNote(code="sheets_bound", count=len(grids)))
    preview = DatasetPreview(sheets=tuple(sheets), more_sheets=more) if segments else None
    return reading_of_segments(reader, segments, notes=notes, dataset=preview)


_READ = {
    ".pdf": _pdf,
    ".docx": _docx,
    ".doc": _doc,
    ".odt": _odt,
    ".pptx": _pptx,
    ".txt": _text,
    ".md": _text,
}


def read_source(data: bytes, filename: str) -> SourceReading:
    """What reading these bytes under this name produces; never raises on a bad file.

    A missing reader library is not a file AIA cannot read: the ``documents`` extra is
    part of the deployment, and its absence raises ``ImportError``.
    """
    suffix = extension_of(filename)
    family = family_of(filename)
    if family in (SourceFamily.IMAGE, SourceFamily.VIDEO):
        return SourceReading(outcome=ReadingOutcome.UNREADABLE, reason="not_text_media")
    read = _READ.get(suffix)
    dataset = _DATASETS.get(suffix)
    if read is None and dataset is None:
        return SourceReading(outcome=ReadingOutcome.UNSUPPORTED, reason="reader_pending")
    try:
        if dataset is not None:
            grids, more = dataset(data)
            return _dataset_reading(READERS[suffix], grids, more)
        assert read is not None
        segments, notes = read(data)
    except ImportError:
        raise
    except DocumentRefused as exc:
        return SourceReading(outcome=ReadingOutcome.REFUSED, reason=exc.reason)
    except Exception:  # a hostile or broken file is refused, never half-read
        return SourceReading(outcome=ReadingOutcome.REFUSED, reason="document_malformed")
    return reading_of_segments(READERS[suffix], segments, notes=notes)
