"""The text of a brief attachment, read as 18.6.6 read it (ADR 0018).

A port of the unit's ``data_library._extract_text`` (``data_library.py:146-169``) that
uses the libraries the unit used -- pypdf, python-docx and openpyxl, the ``documents``
extra, imported here only and only when a file of that kind is read -- so the text an
analysis reads of an attachment is the text 18.6.6 read. ``test_document_text.py``
compares it with the unit's own function on fictional documents
(``tests/fixtures/attachment_text``, captured by ``tools/attachment_text_capture.py``).

The unit's contract is kept: a file AIA cannot read yields ``""`` (the unit's failure
marker, which its attachment code turned into ``""``), so its record says
``text_extracted = False`` and never poses as read. That includes the unit's own
limits -- the first 80 PDF pages, the top-level paragraphs of a Word document (not its
tables), the first 12 sheets, 300 rows and 30 columns of a workbook, and no text from a
sheet that does not declare its size.

One thing is added. An Office file is a ZIP, and the unit decompressed whatever it was
given inside the request; here a ZIP whose members declare more than
``ZIP_MAX_UNCOMPRESSED`` bytes, or more than ``ZIP_MAX_MEMBERS`` members, is not
opened, and keeps no text. ``zipfile`` never inflates a member past its declared size,
so the declaration bounds the work.

**Web documents** (Deep Research, plan ``deep-research-web-search.md`` chunk 6) are
read here too, by the same libraries, with their own contract: :func:`read_web_document`
keeps a PDF's pages apart and an XLSX's or a CSV's cells in their grid, and hands
them to ``domain.deep_research.documents``, which makes the snapshot text and its
layout. A web document AIA will not keep is never ``""``: it raises
:class:`~aia_core.domain.deep_research.documents.DocumentRefused` with a stable
reason -- ``document_malformed``, ``document_encrypted``,
``document_too_many_pages``, ``document_zip_bounds``, ``document_no_text`` -- and
the fetch is refused with it. Its bounds are the domain's (``PDF_MAX_PAGES`` there
is 400 pages, refused past it, not the attachment's first 80; the grid bounds
truncate). The attachment reader above is unchanged.
"""

from __future__ import annotations

import codecs
import contextlib
import csv
import io
import zipfile
from collections.abc import Callable, Iterable
from datetime import date, datetime, time
from typing import Any, Final

from ..domain.attachments import TEXT_MAX_CHARS, extension_of
from ..domain.deep_research import documents as web_documents
from ..domain.deep_research.documents import (
    CapturedDocument,
    DocumentRefused,
    GridContent,
    GridRow,
    PdfBookmark,
    PdfContent,
    grid_document,
    pdf_document,
)
from ..domain.deep_research.web import CSV_MEDIA_TYPE, PDF_MEDIA_TYPE, XLSX_MEDIA_TYPE

__all__ = [
    "PDF_MAX_PAGES",
    "TEXT_SUFFIXES",
    "ZIP_MAX_MEMBERS",
    "ZIP_MAX_UNCOMPRESSED",
    "extract_text",
    "read_csv",
    "read_pdf",
    "read_web_document",
    "read_xlsx",
]

#: Read as UTF-8 text, undecodable bytes replaced (the unit's list).
TEXT_SUFFIXES: Final = frozenset({".txt", ".md", ".csv", ".json", ".tsv"})
PDF_MAX_PAGES: Final = 80
XLSX_MAX_SHEETS: Final = 12
XLSX_MAX_ROWS: Final = 300
XLSX_MAX_COLUMNS: Final = 30
ZIP_MAX_UNCOMPRESSED: Final = 200 * 1024 * 1024
ZIP_MAX_MEMBERS: Final = 5000
_ZIP_SUFFIXES: Final = frozenset({".docx", ".xlsx", ".xlsm"})


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages[:PDF_MAX_PAGES]:
        try:
            parts.append(page.extract_text() or "")
        except Exception:  # the unit skips a page it cannot read
            continue
    return "\n".join(parts)


def _docx(data: bytes) -> str:
    from docx import Document

    return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        out: list[str] = []
        for sheet in workbook.worksheets[:XLSX_MAX_SHEETS]:
            out.append(f"### {sheet.title}")
            # A sheet that declares no size has max_row None; min() raises, and the
            # workbook keeps no text -- as in the unit.
            rows = sheet.iter_rows(
                min_row=1, max_row=min(sheet.max_row, XLSX_MAX_ROWS), values_only=True
            )
            for row in rows:
                out.append(" | ".join("" if v is None else str(v) for v in row[:XLSX_MAX_COLUMNS]))
        return "\n".join(out)
    finally:
        workbook.close()


_READERS: Final[dict[str, Callable[[bytes], str]]] = {
    ".pdf": _pdf,
    ".docx": _docx,
    ".xlsx": _xlsx,
    ".xlsm": _xlsx,
}


def _zip_within_bounds(data: bytes) -> bool:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
    except Exception:  # not a ZIP at all: nothing to read
        return False
    return len(members) <= ZIP_MAX_MEMBERS and (
        sum(m.file_size for m in members) <= ZIP_MAX_UNCOMPRESSED
    )


def extract_text(data: bytes, filename: str) -> str:
    """The text of an attached file, at most ``TEXT_MAX_CHARS`` characters; ``""`` if none.

    The kind of file is its extension, as in the unit. A kind with no reader (an
    image, a video) and a file its reader cannot parse both yield ``""``. A missing
    reader library is not a file AIA cannot read: the ``documents`` extra is part of
    the deployment, and its absence raises.
    """
    suffix = extension_of(filename)
    if suffix in TEXT_SUFFIXES:
        return data.decode("utf-8", errors="replace")[:TEXT_MAX_CHARS]
    reader = _READERS.get(suffix)
    if reader is None:
        return ""
    if suffix in _ZIP_SUFFIXES and not _zip_within_bounds(data):
        return ""
    try:
        text = reader(data)
    except ImportError:
        raise
    except Exception:
        # A malformed or hostile document: kept, with no text (the unit's contract).
        return ""
    return text[:TEXT_MAX_CHARS]


# --------------------------------------------------------------------------- #
# Web documents (Deep Research)
# --------------------------------------------------------------------------- #

#: Nested bookmark levels followed in a PDF's outline.
_MAX_BOOKMARK_DEPTH: Final = 10
#: How much of a CSV the dialect is guessed from.
_CSV_SAMPLE_CHARS: Final = 65_536


def _malformed(kind: str, exc: BaseException) -> DocumentRefused:
    return DocumentRefused(
        f"the {kind} cannot be read ({type(exc).__name__})", reason="document_malformed"
    )


def _bookmarks(reader: Any) -> tuple[PdfBookmark, ...]:
    """The PDF's own outline, depth-first, as far as it can be read."""
    found: list[PdfBookmark] = []

    def walk(items: Iterable[Any], level: int) -> None:
        for item in items:
            if len(found) >= web_documents.MAX_BOOKMARKS:
                return
            if isinstance(item, list):
                if level + 1 < _MAX_BOOKMARK_DEPTH:
                    walk(item, level + 1)
                continue
            try:
                number = reader.get_destination_page_number(item)
            except Exception:  # a destination pypdf cannot follow: no page, still a title
                number = None
            page = number + 1 if isinstance(number, int) and number >= 0 else None
            found.append(PdfBookmark(str(getattr(item, "title", "") or ""), level, page))

    # A malformed outline keeps what was read before it.
    with contextlib.suppress(Exception):
        walk(reader.outline, 0)
    return tuple(found)


def read_pdf(data: bytes) -> PdfContent:
    """A PDF's pages' text, its title and its outline; or :class:`DocumentRefused`.

    Encrypted with an empty user password (a "restricted" PDF a reader opens without
    asking), it is read; with any other, refused. pypdf bounds what one stream may
    inflate to (``pypdf.filters.ZLIB_MAX_OUTPUT_LENGTH``), so a compressed bomb
    fails as malformed. A page whose text cannot be extracted yields no text.
    """
    if b"%PDF-" not in data[:1024]:
        raise DocumentRefused("the body is not a PDF", reason="document_malformed")
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                opened = bool(reader.decrypt(""))
            except Exception:  # an unsupported cipher is a password we do not have
                opened = False
            if not opened:
                raise DocumentRefused("the PDF needs a password", reason="document_encrypted")
        count = len(reader.pages)
    except DocumentRefused:
        raise
    except Exception as exc:
        raise _malformed("PDF", exc) from exc
    if count > web_documents.PDF_MAX_PAGES:
        raise DocumentRefused(
            f"the PDF has {count} pages; at most {web_documents.PDF_MAX_PAGES} are kept",
            reason="document_too_many_pages",
        )
    pages: list[str] = []
    for index in range(count):
        try:
            pages.append(reader.pages[index].extract_text() or "")
        except Exception:  # one unreadable page is a page without text
            pages.append("")
    try:
        title = str(getattr(reader.metadata, "title", None) or "")
    except Exception:
        title = ""
    return PdfContent(pages=tuple(pages), title=title, bookmarks=_bookmarks(reader))


def _cell_text(value: object) -> str:
    """A cell's stored value as text: what the file holds, never a display format."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(" ")
    if isinstance(value, date | time):
        return value.isoformat()
    return str(value)


def _grid(
    name: str, rows: Iterable[tuple[object, ...]], budget: list[int], cut: bool
) -> GridContent:
    """Rows as read (numbered from 1), within the row, column and cell bounds."""
    kept: list[GridRow] = []
    truncated = cut
    for number, row in enumerate(rows, start=1):
        if number > web_documents.MAX_GRID_ROWS:
            truncated = True
            break
        values = [_cell_text(v) for v in row]
        if any(values[web_documents.MAX_GRID_COLUMNS :]):
            truncated = True
        values = values[: web_documents.MAX_GRID_COLUMNS]
        filled = sum(1 for v in values if v.strip())
        if not filled:
            continue
        if filled > budget[0]:
            truncated = True
            budget[0] = 0  # the document's cells are spent: no later sheet is read
            break
        budget[0] -= filled
        kept.append(GridRow(number=number, values=tuple(values)))
    return GridContent(name=name, rows=tuple(kept), truncated=truncated)


def _zip_within_web_bounds(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
    except Exception as exc:
        raise _malformed("XLSX", exc) from exc
    declared = sum(m.file_size for m in members)
    if (
        len(members) > ZIP_MAX_MEMBERS
        or declared > ZIP_MAX_UNCOMPRESSED
        or declared > web_documents.ZIP_MAX_RATIO * max(len(data), 1)
    ):
        raise DocumentRefused(
            f"the XLSX declares {len(members)} members and {declared} bytes unpacked",
            reason="document_zip_bounds",
        )


def read_xlsx(data: bytes) -> tuple[str, tuple[GridContent, ...], bool]:
    """A workbook's title, its sheets as grids, and whether a sheet bound cut it."""
    if not data.startswith(b"PK\x03\x04"):
        raise DocumentRefused("the body is not an XLSX", reason="document_malformed")
    _zip_within_web_bounds(data)
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise _malformed("XLSX", exc) from exc
    try:
        title = str(workbook.properties.title or "")
        sheets = workbook.worksheets
        budget = [web_documents.MAX_GRID_CELLS]
        grids: list[GridContent] = []
        for sheet in sheets[: web_documents.MAX_GRID_SHEETS]:
            # A declared size bounds the padding openpyxl adds; an undeclared one
            # yields the rows as stored.
            max_row, max_col = sheet.max_row, sheet.max_column
            cut = bool(
                (max_row and max_row > web_documents.MAX_GRID_ROWS)
                or (max_col and max_col > web_documents.MAX_GRID_COLUMNS)
            )
            rows = sheet.iter_rows(
                min_row=1,
                min_col=1,
                max_row=min(max_row, web_documents.MAX_GRID_ROWS) if max_row else None,
                max_col=min(max_col, web_documents.MAX_GRID_COLUMNS) if max_col else None,
                values_only=True,
            )
            grids.append(_grid(str(sheet.title), rows, budget, cut))
            if budget[0] <= 0:
                break
        more = len(sheets) > len(grids)
    except DocumentRefused:
        raise
    except Exception as exc:
        raise _malformed("XLSX", exc) from exc
    finally:
        workbook.close()
    return title, tuple(grids), more


def _encoding(charset: str | None) -> str:
    try:
        name = codecs.lookup(charset or "utf-8").name
    except LookupError:
        name = "utf-8"
    # A UTF-8 file often starts with a byte-order mark; it is not part of the first cell.
    return "utf-8-sig" if name == "utf-8" else name


def read_csv(data: bytes, *, charset: str | None = None) -> GridContent:
    """A CSV as one grid: its declared charset (UTF-8 if none), its dialect sniffed.

    Undecodable bytes are replaced, as a page's are. A delimiter is one of comma,
    semicolon or tab (a Czech CSV is often semicolon-separated); when none can be
    told, a comma. A row the reader cannot parse refuses the file.
    """
    text = data.decode(_encoding(charset), errors="replace")
    if "\x00" in text:
        # Text in another encoding than declared (UTF-16 from a spreadsheet, say), or
        # not text at all: its cells would be noise.
        raise DocumentRefused(
            "the CSV holds NUL characters: not text in its declared charset",
            reason="document_malformed",
        )
    sample = text[:_CSV_SAMPLE_CHARS]
    if len(text) > len(sample) and "\n" in sample:
        sample = sample[: sample.rindex("\n")]
    try:
        dialect: Any = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    try:
        rows = list(_bounded_rows(csv.reader(io.StringIO(text, newline=""), dialect, strict=True)))
    except csv.Error as exc:
        raise _malformed("CSV", exc) from exc
    return _grid("", rows, [web_documents.MAX_GRID_CELLS], False)


def _bounded_rows(reader: Iterable[list[str]]) -> Iterable[tuple[object, ...]]:
    # One row past the bound, so the grid can say it was cut.
    for number, row in enumerate(reader, start=1):
        yield tuple(row)
        if number > web_documents.MAX_GRID_ROWS:
            return


def read_web_document(
    data: bytes, media_type: str, *, charset: str | None = None
) -> CapturedDocument:
    """A fetched document's snapshot text and layout; or :class:`DocumentRefused`.

    A missing reader library is a defect of the deployment, not of the document,
    and raises ``ImportError``.
    """
    if media_type == PDF_MEDIA_TYPE:
        return pdf_document(read_pdf(data))
    if media_type == XLSX_MEDIA_TYPE:
        title, sheets, more = read_xlsx(data)
        return grid_document("xlsx", sheets, title=title, truncated=more)
    if media_type == CSV_MEDIA_TYPE:
        return grid_document("csv", (read_csv(data, charset=charset),))
    raise ValueError(f"{media_type!r} is not a document type")
