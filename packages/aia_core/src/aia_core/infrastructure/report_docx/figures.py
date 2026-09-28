"""Figures: a chart or a Sociomap, captioned, sourced and described for a screen reader.

"Graf N — title (n = …; grade)" above, the picture (SVG with a PNG fallback,
``wp:docPr/@descr`` = the alt text), the source line below — including how many
suppressed points were left out.

**A Sociomap enters a client report only through ``require_client_facing``**,
which fails closed; validation refuses such a document first, and the renderer
checks again rather than trust that nothing reached it another way. An internal
report may show an INTERNAL_ONLY map, and its caption says it is not approved.
"""

from __future__ import annotations

from io import BytesIO

from docx.shared import Mm
from PIL import Image

from aia_core.domain.report import numbers
from aia_core.domain.report.copy import t
from aia_core.domain.report.model import Figure, SociomapFigure
from aia_core.domain.research_sociomap import MethodologyStatus, require_client_facing
from aia_core.infrastructure.report_docx import layout
from aia_core.infrastructure.report_docx.charts import TEXT_WIDTH_MM, draw_chart
from aia_core.infrastructure.report_docx.context import Container, RenderContext
from aia_core.infrastructure.report_docx.images import add_vector_image
from aia_core.infrastructure.report_docx.styles import S
from aia_core.infrastructure.report_docx.tables import caption, source_line, uniform_grade


def render_figure(ctx: RenderContext, container: Container, block: Figure) -> None:
    number = ctx.outline.item_numbers[ctx.position]
    if block.landscape:
        section = layout.end_section(ctx, layout.last_paragraph(ctx), landscape=True)
        layout.running(ctx, section, wide=True)
    refs = [
        r
        for s in block.chart.series
        for r in s.refs
        if r is not None and not ctx.ledger.is_suppressed(r)
    ]
    caption(
        ctx,
        container,
        word=t("figure"),
        number=number,
        title=block.title,
        anchor=block.id,
        base_ref=block.base_ref,
        grade=uniform_grade(ctx, refs),
    )
    drawn = draw_chart(ctx, block.chart, landscape=block.landscape)
    p = container.add_paragraph(style=S.FIGURE)
    add_vector_image(
        p,
        ctx.svgs,
        svg=drawn.svg,
        png=drawn.png,
        width_mm=drawn.width_mm,
        height_mm=drawn.height_mm,
        alt=block.alt,
        name=f"{t('figure')} {number}",
    )
    source_line(container, block.source, drawn.omitted, block.notes)
    if block.landscape:
        layout.running(ctx, layout.end_section(ctx, layout.last_paragraph(ctx)))


def render_sociomap(ctx: RenderContext, container: Container, block: SociomapFigure) -> None:
    if ctx.client:
        require_client_facing({"methodology_status": block.methodology_status})
    number = ctx.outline.item_numbers[ctx.position]
    stress = f"{t('stress')} = {numbers.number(round(block.stress_1, 3), 3)}"
    approved = block.methodology_status == MethodologyStatus.CLIENT_FACING.value
    caption(
        ctx,
        container,
        word=t("figure"),
        number=number,
        title=block.title,
        anchor=block.id,
        base_ref=None,
        grade=None,
        extra=stress if approved else f"{stress}; {t('sociomap_internal')}",
    )
    with Image.open(BytesIO(block.image_png)) as im:
        px_w, px_h = im.size
    width = TEXT_WIDTH_MM
    p = container.add_paragraph(style=S.FIGURE)
    shape = p.add_run().add_picture(
        BytesIO(block.image_png), width=Mm(width), height=Mm(width * px_h / px_w)
    )
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", block.alt)
    doc_pr.set("name", f"{t('figure')} {number}")
    for side in ("distT", "distB", "distL", "distR"):
        shape._inline.set(side, "0")
    source_line(container, block.source, 0, ())
