"""One renderer per block of the document model.

Each function takes the render context, the container it writes into (the body,
or a table cell) and one block, and writes paragraphs that carry a named style
and nothing else: no run or paragraph here sets a font, a size, a colour or an
indent. How a block looks is the style sheet's (``styles.py``).
"""

from __future__ import annotations

from collections.abc import Callable
from typing import Any

from docx.enum.text import WD_BREAK
from docx.text.paragraph import Paragraph as DocxParagraph

from aia_core.domain.report.copy import method_status_text, t
from aia_core.domain.report.model import (
    Block,
    Callout,
    CalloutKind,
    CrossRef,
    Footnote,
    Heading,
    Inline,
    Link,
    PageBreak,
    Paragraph,
    ParagraphRole,
    Section,
    Text,
    Value,
)
from aia_core.domain.report.outline import OutlineEntry
from aia_core.infrastructure.report_docx.context import Container, RenderContext
from aia_core.infrastructure.report_docx.layout import heading_label
from aia_core.infrastructure.report_docx.ooxml import add_external_link, add_internal_link
from aia_core.infrastructure.report_docx.styles import S

# ---------------------------------------------------------------------- inline

_EMPHASIS = {"none": None, "em": S.EMPHASIS, "strong": S.STRONG}


def write_inlines(ctx: RenderContext, p: DocxParagraph, content: tuple[Inline, ...]) -> None:
    for inline in content:
        if isinstance(inline, Text):
            p.add_run(inline.text, style=_EMPHASIS[inline.emphasis])
        elif isinstance(inline, Value):
            write_value(ctx, p, inline.ref, with_interval=inline.with_interval)
        elif isinstance(inline, Footnote):
            ctx.footnotes.add(p, inline.text)
        elif isinstance(inline, CrossRef):
            add_internal_link(
                p,
                ctx.outline.labels[inline.target],
                ctx.bookmarks.name(inline.target),
                style=S.LINK,
            )
        elif isinstance(inline, Link):
            add_external_link(p, inline.text, inline.url, style=S.LINK)


def write_value(ctx: RenderContext, p: DocxParagraph, ref: str, *, with_interval: bool) -> None:
    raise NotImplementedError("evidenced values render with the report components (R5)")


# ---------------------------------------------------------------------- headings


def chapter_heading(ctx: RenderContext, index: int, section: Section) -> DocxParagraph:
    """``Heading 1`` (or the appendix style): the number, then the title; bookmarked."""
    entry = _chapter_entry(ctx, index)
    style = S.APPENDIX if section.appendix else S.H1
    p = ctx.document.add_paragraph(style=style)
    p.add_run(heading_label(entry, appendix=section.appendix), style=S.HEADING_NUMBER)
    p.add_run(f" {section.title}")
    ctx.bookmarks.wrap(p, entry.anchor)
    return p


def _chapter_entry(ctx: RenderContext, index: int) -> OutlineEntry:
    return [e for e in ctx.outline.headings if e.level == 1][index]


def render_heading(ctx: RenderContext, container: Container, block: Heading) -> None:
    number = ctx.outline.heading_numbers[ctx.position]
    anchor = next(
        e.anchor for e in ctx.outline.headings if e.number == number and e.level == block.level
    )
    p = container.add_paragraph(style=S.H2 if block.level == 2 else S.H3)
    p.add_run(number, style=S.HEADING_NUMBER)
    p.add_run(f" {block.text}")
    ctx.bookmarks.wrap(p, anchor)


# ---------------------------------------------------------------------- text


def render_paragraph(ctx: RenderContext, container: Container, block: Paragraph) -> None:
    style = S.LEDE if block.role is ParagraphRole.LEDE else S.BODY
    write_inlines(ctx, container.add_paragraph(style=style), block.content)


def render_page_break(ctx: RenderContext, container: Container, block: PageBreak) -> None:
    container.add_paragraph(style=S.BODY).add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------------- callouts

_CALLOUT_TITLES = {
    CalloutKind.DECISION: "decision_answer",
    CalloutKind.METHOD: "method_status",
    CalloutKind.LIMITATION: "limitation",
    CalloutKind.PROVISIONAL: "provisional",
    CalloutKind.NOTE: "note",
}


def render_callout(ctx: RenderContext, container: Container, block: Callout) -> None:
    """A boxed statement: title (with the grades of its refs), then its text.

    The method callout always prints the document's method status in the words
    code chose (``copy.method_status_text``) before anything the template adds.
    """
    title = container.add_paragraph(style=S.CALLOUT_TITLE)
    title.add_run(block.title or t(_CALLOUT_TITLES[block.kind]))
    for ref in block.refs:
        title.add_run(" ")
        write_mark(ctx, title, ref)
    if block.kind is CalloutKind.METHOD:
        container.add_paragraph(method_status_text(ctx.report.meta.method_status), style=S.CALLOUT)
    if block.content:
        write_inlines(ctx, container.add_paragraph(style=S.CALLOUT), block.content)


def write_mark(ctx: RenderContext, p: DocxParagraph, ref: str) -> None:
    raise NotImplementedError("evidence marks render with the figures (R7)")


# ---------------------------------------------------------------------- dispatch

_RENDERERS: dict[type[Any], Callable[[RenderContext, Container, Any], None]] = {
    Heading: render_heading,
    Paragraph: render_paragraph,
    PageBreak: render_page_break,
    Callout: render_callout,
}


def render_block(ctx: RenderContext, container: Container, block: Block) -> None:
    renderer = _RENDERERS.get(type(block))
    if renderer is None:
        raise NotImplementedError(f"no renderer for {type(block).__name__} yet")
    renderer(ctx, container, block)
