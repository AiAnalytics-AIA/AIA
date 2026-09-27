"""The data table: captioned, ruled, its header repeating, suppression stated.

A table prints as a consultancy table does:

* the caption above it — "Tabulka N — title (n = 1 204; měřeno)" — with the
  number as a ``SEQ Tabulka`` field (so Word's list of tables finds it), the base
  from the ``base_ref`` row's effective n, and the grade when every number in the
  table shares one (otherwise each cell carries its own mark);
* the header row repeating on every page, the unit in the header, numbers
  right-aligned in the tabular face;
* **a row that cites a suppressed ref is removed, never greyed**, and the source
  line under the table says how many rows were removed and why;
* INDICATIVE cells marked with a dagger, explained under the table.

A landscape table gets a landscape section of its own.
"""

from __future__ import annotations

from typing import Final

from docx.table import _Cell

from aia_core.domain.report import numbers
from aia_core.domain.report.copy import t
from aia_core.domain.report.evidence import PrintGrade
from aia_core.domain.report.model import (
    Align,
    Column,
    EmptyCell,
    NumberCell,
    Table,
    TableRow,
    TextCell,
)
from aia_core.infrastructure.report_docx import layout
from aia_core.infrastructure.report_docx.blocks import base_text, shows_interval, write_mark
from aia_core.infrastructure.report_docx.context import Container, RenderContext
from aia_core.infrastructure.report_docx.marks import grade_label
from aia_core.infrastructure.report_docx.ooxml import add_field, full_width, header_row
from aia_core.infrastructure.report_docx.styles import S

INDICATIVE_MARK: Final = "\u2020"  # dagger


def _grade_phrase(grade: PrintGrade) -> str:
    label = grade_label(grade)
    return label[0].lower() + label[1:]


def caption(
    ctx: RenderContext,
    container: Container,
    *,
    word: str,
    number: int,
    title: str,
    anchor: str,
    base_ref: str | None,
    grade: PrintGrade | None,
    extra: str | None = None,
) -> None:
    """ "Tabulka 3 — title (n = 480; měřeno)": the SEQ field, then what it rests on."""
    p = container.add_paragraph(style=S.CAPTION)
    p.add_run(f"{word} ")
    add_field(p, f"SEQ {word} \\* ARABIC", str(number))
    p.add_run(f" — {title}")
    basis = [base_text(ctx, base_ref)] if base_ref else []
    if grade is not None:
        basis.append(_grade_phrase(grade))
    if extra:
        basis.append(extra)
    if basis:
        p.add_run(f" ({'; '.join(basis)})", style=S.GRADE)
    ctx.bookmarks.wrap(p, anchor)


def source_line(container: Container, source: str, removed: int, notes: tuple[str, ...]) -> None:
    """ "Zdroj: …", what was removed and why, then the notes."""
    p = container.add_paragraph(style=S.SOURCE)
    p.add_run(f"{t('source')}: ", style=S.STRONG)
    p.add_run(source)
    if removed:
        p.add_run(f"{t('separator')}{t('suppressed_rows')}: {removed}.")
    for note in notes:
        n = container.add_paragraph(style=S.SOURCE)
        n.add_run(f"{t('note')}: ", style=S.STRONG)
        n.add_run(note)


def _refs(row: TableRow) -> list[str]:
    return [c.ref for c in row.cells if isinstance(c, NumberCell)]


def uniform_grade(ctx: RenderContext, refs: list[str]) -> PrintGrade | None:
    """The one grade every ref shares, or ``None`` when they differ (or there are none)."""
    grades = {ctx.ledger.grade(r) for r in refs}
    return grades.pop() if len(grades) == 1 else None


def _head_label(column: Column) -> str:
    return f"{column.label} ({column.unit})" if column.unit else column.label


def _write_cell(
    ctx: RenderContext,
    cell: _Cell,
    column: Column,
    value: TextCell | NumberCell | EmptyCell,
    *,
    strong: bool,
    marks: bool,
) -> bool:
    """Write one cell; return whether it was INDICATIVE."""
    style = S.TABLE_NUMBER if column.align is Align.RIGHT else S.TABLE
    p = cell.paragraphs[0]
    p.style = style
    run_style = S.STRONG if strong else None
    if isinstance(value, TextCell):
        p.add_run(value.text, style=run_style)
        return False
    if isinstance(value, EmptyCell):
        p.add_run(value.reason, style=S.GRADE)
        return False
    row = ctx.ledger.row(value.ref)
    # The unit is in the header when the column names one; never twice.
    p.add_run(
        numbers.number(row.value, row.decimals)
        if column.unit
        else numbers.with_unit(row.value, row.decimals, row.unit),
        style=run_style,
    )
    indicative = ctx.ledger.is_indicative(value.ref)
    if indicative:
        p.add_run(INDICATIVE_MARK, style=S.GRADE)
    if marks:
        p.add_run(numbers.NBSP)
        write_mark(ctx, p, value.ref)
    if row.interval is not None and shows_interval(ctx, value.ref, value.with_interval):
        under = cell.add_paragraph(style=S.TABLE_NUMBER)
        under.add_run(numbers.interval(row.interval, row.decimals, row.unit), style=S.GRADE)
    return indicative


def render_table(ctx: RenderContext, container: Container, block: Table) -> None:
    number = ctx.outline.item_numbers[ctx.position]
    if block.landscape:
        layout.running(
            ctx, layout.end_section(ctx, layout.last_paragraph(ctx), landscape=True), wide=True
        )

    kept = [r for r in block.rows if not any(ctx.ledger.is_suppressed(x) for x in _refs(r))]
    removed = len(block.rows) - len(kept)
    refs = [ref for r in kept for ref in _refs(r)]
    grade = uniform_grade(ctx, refs)
    caption(
        ctx,
        container,
        word=t("table"),
        number=number,
        title=block.title,
        anchor=block.id,
        base_ref=block.base_ref,
        grade=grade,
    )

    table = ctx.document.add_table(rows=1, cols=len(block.columns))
    table.style = S.DATA_TABLE
    full_width(table)
    header_row(table.rows[0])
    for cell, column in zip(table.rows[0].cells, block.columns, strict=True):
        p = cell.paragraphs[0]
        p.style = S.TABLE_HEAD_NUMBER if column.align is Align.RIGHT else S.TABLE_HEAD
        p.add_run(_head_label(column))
    any_indicative = False
    for trow in kept:
        cells = table.add_row().cells
        for cell, column, value in zip(cells, block.columns, trow.cells, strict=True):
            any_indicative |= _write_cell(
                ctx, cell, column, value, strong=trow.emphasis, marks=grade is None
            )

    source_line(container, block.source, removed, block.notes)
    if any_indicative:
        p = container.add_paragraph(style=S.SOURCE)
        p.add_run(f"{INDICATIVE_MARK} {t('indicative_note')}")
    if block.landscape:
        layout.running(ctx, layout.end_section(ctx, layout.last_paragraph(ctx)))
