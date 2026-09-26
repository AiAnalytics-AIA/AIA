"""The report's vendored fonts, and their embedding in a DOCX package.

The faces must be the upstream files unmodified (their licences reserve the font
names), must be exactly the set the print register can ask for, and must survive
the ECMA-376 obfuscation round trip byte for byte.
"""

from __future__ import annotations

import hashlib
import io
import struct
import zipfile

import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from aia_core.domain.report import print_tokens
from aia_core.infrastructure.report_docx import embed

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _tables(data: bytes) -> dict[str, bytes]:
    """The sfnt table directory of a TrueType file, by tag (stdlib only)."""
    (num,) = struct.unpack(">H", data[4:6])
    out: dict[str, bytes] = {}
    for i in range(num):
        tag, _, offset, length = struct.unpack(">4sLLL", data[12 + 16 * i : 28 + 16 * i])
        out[tag.decode("latin-1")] = data[offset : offset + length]
    return out


def _family_name(data: bytes) -> str:
    """``name`` table ID 1 (the legacy family name), Windows Unicode record."""
    table = _tables(data)["name"]
    _, count, storage = struct.unpack(">HHH", table[:6])
    for i in range(count):
        platform, _enc, _lang, name_id, length, offset = struct.unpack(
            ">HHHHHH", table[6 + 12 * i : 18 + 12 * i]
        )
        if platform == 3 and name_id == 1:
            return table[storage + offset : storage + offset + length].decode("utf-16-be")
    raise AssertionError("no Windows family name")


def test_vendored_fonts_match_their_checksums() -> None:
    lines = (embed.FONT_DIR / "SHA256SUMS").read_text().splitlines()
    sums = dict(reversed(line.split(maxsplit=1)) for line in lines)
    files = {p.name for p in embed.FONT_DIR.glob("*.ttf")}
    assert files == set(sums)
    for name, expected in sums.items():
        assert hashlib.sha256((embed.FONT_DIR / name).read_bytes()).hexdigest() == expected, name


def test_every_print_font_is_vendored_and_nothing_else() -> None:
    assert set(embed.FACES) == set(print_tokens.FONTS)
    vendored = {f for faces in embed.FACES.values() for f in faces.values()}
    assert vendored == {p.name for p in embed.FONT_DIR.glob("*.ttf")}


@pytest.mark.parametrize("name", sorted(embed.FACES))
def test_each_face_is_the_family_word_will_ask_for(name: str) -> None:
    for filename in embed.FACES[name].values():
        data = (embed.FONT_DIR / filename).read_bytes()
        assert _family_name(data) == name, filename
        # OS/2 fsType 0: installable embedding permitted.
        (fs_type,) = struct.unpack(">H", _tables(data)["OS/2"][8:10])
        assert fs_type == 0, filename
        assert "glyf" in _tables(data), f"{filename}: TrueType outlines required"


def test_obfuscation_is_its_own_inverse_and_touches_only_32_bytes() -> None:
    data = bytes(range(256)) * 4
    key = embed.font_key(data)
    hidden = embed.obfuscate(data, key)
    assert hidden[32:] == data[32:]
    assert hidden[:32] != data[:32]
    assert embed.obfuscate(hidden, key) == data


def test_the_key_is_deterministic_per_face() -> None:
    a = (embed.FONT_DIR / "SourceSerif4-Regular.ttf").read_bytes()
    b = (embed.FONT_DIR / "IBMPlexSans-Regular.ttf").read_bytes()
    assert embed.font_key(a) == embed.font_key(a)
    assert embed.font_key(a) != embed.font_key(b)
    assert embed.font_key(a).startswith("{") and embed.font_key(a).endswith("}")


def test_obfuscation_follows_the_standard_key_order() -> None:
    # ECMA-376 Part 1 §17.8.1: the key is the GUID's bytes read from the end.
    key = "{00112233-4455-6677-8899-AABBCCDDEEFF}"
    out = embed.obfuscate(bytes(32), key)
    assert out[:16] == bytes.fromhex("FFEEDDCCBBAA99887766554433221100")
    assert out[16:32] == out[:16]


def _saved(names: list[str]) -> zipfile.ZipFile:
    doc = Document()
    doc.add_paragraph("Příliš žluťoučký kůň úpěl ďábelské ódy.")
    embed.embed_fonts(doc, names)
    buf = io.BytesIO()
    doc.save(buf)
    return zipfile.ZipFile(buf)


def test_embedded_faces_round_trip_to_the_vendored_bytes() -> None:
    z = _saved(["Source Serif 4", "IBM Plex Sans SmBld"])
    table = etree.fromstring(z.read("word/fontTable.xml"))
    rels = etree.fromstring(z.read("word/_rels/fontTable.xml.rels"))
    targets = {r.get("Id"): r.get("Target") for r in rels}
    fonts = {f.get(f"{W}name"): f for f in table.iter(f"{W}font")}
    for name in ("Source Serif 4", "IBM Plex Sans SmBld"):
        slots = embed.FACES[name]
        embeds = [e for e in fonts[name] if e.tag.startswith(f"{W}embed")]
        assert {e.tag.removeprefix(f"{W}embed") for e in embeds} == set(slots)
        for e in embeds:
            slot = e.tag.removeprefix(f"{W}embed")
            part = z.read("word/" + targets[e.get(f"{R}id")])
            original = (embed.FONT_DIR / slots[slot]).read_bytes()  # type: ignore[index]
            assert embed.obfuscate(part, e.get(f"{W}fontKey")) == original


def test_parts_are_declared_as_obfuscated_fonts() -> None:
    z = _saved(["IBM Plex Mono"])
    types = z.read("[Content_Types].xml").decode()
    assert embed.OBFUSCATED_FONT in types
    assert any(n.startswith("word/fonts/") and n.endswith(".odttf") for n in z.namelist())


def test_settings_enable_embedding_in_schema_order() -> None:
    z = _saved(["IBM Plex Sans"])
    settings = etree.fromstring(z.read("word/settings.xml"))
    tags = [child.tag.removeprefix(W) for child in settings]
    assert "embedTrueTypeFonts" in tags
    later = set(embed._AFTER_EMBED_TRUETYPE)
    first_later = next((i for i, t in enumerate(tags) if t in later), len(tags))
    assert tags.index("embedTrueTypeFonts") < first_later


def test_embedding_the_same_font_twice_adds_no_duplicate_font_entry() -> None:
    doc = Document()
    embed.embed_fonts(doc, ["IBM Plex Mono", "IBM Plex Mono"])
    buf = io.BytesIO()
    doc.save(buf)
    table = etree.fromstring(zipfile.ZipFile(buf).read("word/fontTable.xml"))
    names = [f.get(qn("w:name")) for f in table.iter(f"{W}font")]
    assert names.count("IBM Plex Mono") == 1


def test_an_unvendored_font_is_refused_not_substituted() -> None:
    with pytest.raises(KeyError, match="Aptos"):
        embed.embed_fonts(Document(), ["Aptos"])
