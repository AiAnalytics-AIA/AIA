"""A client's sources (lifecycle K2a).

The rules these tests pin:
- a reading never claims more than happened, and every segment says where it stands;
- the formats the owner chose are read, kept unread, or refused by name -- never a
  title-only card;
- a source is kept once per client by its bytes, with every occurrence;
- it is found only within its client, and read back against the hash kept for it;
- it is not knowledge: nothing a study reads carries it.
"""

from __future__ import annotations

import dataclasses
import io
import json
from pathlib import Path
from typing import Any

import pytest
from pydantic import ValidationError
from sqlalchemy import select

from aia_core.domain.client_sources import (
    NAME_MAX_CHARS,
    SOURCE_MAX_BYTES,
    Locator,
    ReadingNote,
    ReadingOutcome,
    Segment,
    SourceFamily,
    SourceReading,
    SourceRejected,
    display_name,
    family_of,
    reading_of_segments,
)
from aia_core.domain.knowledge import KnowledgeKind
from aia_core.domain.research_agents import context_snapshot
from aia_core.domain.scope import ClientContext, ClientPermission, ScopeDenied
from aia_core.infrastructure.client_knowledge_repository import ClientKnowledgeRepository
from aia_core.infrastructure.client_source_repository import ClientSourceRepository
from aia_core.infrastructure.source_reading import read_source
from aia_core.infrastructure.storage import InMemoryArtifactStore, IntegrityError
from aia_core.infrastructure.tables import AccessAuditRow

FIXTURES = Path(__file__).parent / "fixtures" / "attachment_text" / "inputs"


def fixture(name: str) -> bytes:
    return (FIXTURES / name).read_bytes()


def client_scope(scoped: Any, user: str, client: str = "primary") -> ClientContext:
    return scoped.resolver.client_context(
        scoped.principal(scoped.users[user]), client_id=scoped.clients[client].client_id
    )


@pytest.fixture
def store() -> InMemoryArtifactStore:
    return InMemoryArtifactStore()


@pytest.fixture
def sources(session: Any, store: InMemoryArtifactStore) -> ClientSourceRepository:
    return ClientSourceRepository(session, store)


# --------------------------------------------------------------------- the domain


def test_a_locator_places_its_segment() -> None:
    assert Locator(page=2).page == 2
    with pytest.raises(ValidationError):
        Locator()
    with pytest.raises(ValidationError):
        Locator(page=0)


def test_a_reading_says_why_it_is_not_complete_and_keeps_nothing_it_did_not_read() -> None:
    with pytest.raises(ValidationError):  # complete, yet explaining something
        SourceReading(outcome=ReadingOutcome.READ, reader="t/1", notes=(ReadingNote(code="x"),))
    with pytest.raises(ValidationError):  # incomplete, with no reason
        SourceReading(outcome=ReadingOutcome.PARTIAL, reader="t/1")
    with pytest.raises(ValidationError):  # nothing read, yet a segment
        SourceReading(
            outcome=ReadingOutcome.UNSUPPORTED,
            reason="reader_pending",
            segments=(Segment(locator=Locator(line=1), text="x"),),
        )


def test_reading_drops_empty_parts_bounds_the_text_and_names_a_scan() -> None:
    parts = [
        Segment(locator=Locator(page=1), text="  "),
        Segment(locator=Locator(page=2), text="abcdef"),
        Segment(locator=Locator(page=3), text="ghij"),
    ]
    whole = reading_of_segments("t/1", parts)
    assert whole.outcome is ReadingOutcome.READ
    assert [s.locator.page for s in whole.segments] == [2, 3]

    cut = reading_of_segments("t/1", parts, limit=8)
    assert cut.outcome is ReadingOutcome.PARTIAL and cut.reason == "text_bound"
    assert [s.text for s in cut.segments] == ["abcdef", "gh"]
    assert cut.characters == 8

    scan = reading_of_segments("t/1", parts[:1])
    assert (scan.outcome, scan.reason, scan.segments) == (ReadingOutcome.PARTIAL, "no_text", ())


def test_a_name_is_kept_as_written_without_a_path_or_control_characters() -> None:
    assert display_name("C:\\Data\\Zadání - 2026.pdf") == "Zadání - 2026.pdf"
    assert display_name("../../etc/passwd") == "passwd"
    assert display_name("a\x00b\n.txt") == "ab.txt"
    assert display_name("   ") == "soubor"
    long = display_name("x" * 300 + ".docx")
    assert len(long) == NAME_MAX_CHARS and long.endswith(".docx")


def test_the_owners_formats_have_a_family_and_nothing_else_does() -> None:
    for name in ("a.pdf", "a.docx", "a.doc", "a.odt", "a.txt", "a.md", "a.pptx"):
        assert family_of(name) is SourceFamily.DOCUMENT, name
    for name in ("a.csv", "a.xlsx", "a.xls", "a.ods"):
        assert family_of(name) is SourceFamily.DATASET, name
    assert family_of("a.PNG") is SourceFamily.IMAGE
    assert family_of("a.mov") is SourceFamily.VIDEO
    for name in ("a.exe", "a.zip", "a.xlsm", "a"):
        assert family_of(name) is SourceFamily.UNKNOWN, name


# --------------------------------------------------------------------- the reader


def test_a_pdf_is_read_page_by_page() -> None:
    reading = read_source(fixture("prezentace.pdf"), "prezentace.pdf")
    assert (reading.outcome, reading.reader) == (ReadingOutcome.READ, "pdf/1")
    assert [s.locator.page for s in reading.segments] == [1, 2]
    assert "Zadání výzkumu" in reading.segments[0].text


def test_a_docx_is_read_in_body_order_with_its_tables_by_row() -> None:
    from docx import Document

    document = Document()
    document.add_paragraph("Úvod")
    document.add_paragraph("")
    table = document.add_table(rows=2, cols=2)
    for r, row in enumerate([("Značka", "Podíl"), ("A", "12 %")]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    document.add_paragraph("Závěr")
    out = io.BytesIO()
    document.save(out)

    reading = read_source(out.getvalue(), "zadani.docx")
    assert reading.outcome is ReadingOutcome.READ
    placed = [(s.locator.model_dump(exclude_none=True), s.text) for s in reading.segments]
    assert placed == [
        ({"paragraph": 1}, "Úvod"),
        ({"table": 1, "row": 1}, "Značka | Podíl"),
        ({"table": 1, "row": 2}, "A | 12 %"),
        # The empty paragraph is counted, so the number is the one a reader finds.
        ({"paragraph": 3}, "Závěr"),
    ]


def test_text_is_read_by_paragraph_with_its_first_line() -> None:
    reading = read_source("Nadpis\n\nprvní\ndruhý\n\n\nposlední\n".encode(), "poznamky.md")
    placed = [(s.locator.line, s.text) for s in reading.segments]
    assert placed == [(1, "Nadpis"), (3, "první\ndruhý"), (7, "poslední")]


@pytest.mark.parametrize(
    ("name", "outcome", "reason"),
    [
        ("obrazek.png", ReadingOutcome.UNREADABLE, "not_text_media"),
        ("poskozeny.pdf", ReadingOutcome.REFUSED, "document_malformed"),
        ("poskozeny.docx", ReadingOutcome.REFUSED, "document_malformed"),
        ("neplatne.txt", ReadingOutcome.PARTIAL, "not_utf8"),
    ],
)
def test_what_cannot_be_read_says_so_by_name(
    name: str, outcome: ReadingOutcome, reason: str
) -> None:
    reading = read_source(fixture(name), name)
    assert (reading.outcome, reading.reason) == (outcome, reason)


def test_every_format_on_the_owners_list_has_a_reader() -> None:
    # The owner's list of 2026-10-13, as the families name it: nothing on it is left
    # "reader pending" once K2c lands.
    from aia_core.domain.client_sources import _FAMILIES
    from aia_core.infrastructure.source_reading import READERS

    readable = {
        s for s, f in _FAMILIES.items() if f in (SourceFamily.DOCUMENT, SourceFamily.DATASET)
    }
    assert readable == set(READERS)


# ------------------------------------------------------------------ the repository


def test_a_source_is_kept_once_by_its_bytes_with_every_occurrence(
    scoped: Any, sources: ClientSourceRepository, store: InMemoryArtifactStore
) -> None:
    lead = client_scope(scoped, "lead")
    first, created = sources.add(lead, filename="Zadání.pdf", data=fixture("prezentace.pdf"))
    assert created
    assert (first.family, first.outcome, first.reader, first.segments) == (
        SourceFamily.DOCUMENT,
        ReadingOutcome.READ,
        "pdf/1",
        2,
    )
    again, created = sources.add(lead, filename="kopie.pdf", data=fixture("prezentace.pdf"))
    assert not created and again.source_id == first.source_id
    assert again.filename == "kopie.pdf"
    assert [o.filename for o in again.occurrences] == ["kopie.pdf", "Zadání.pdf"]
    # One original, one reading, both under the client's prefix.
    prefix = f"org/{lead.organization_id}/client/{lead.client_id}/sources/{first.sha256}/"
    assert sorted(store.keys) == [prefix + "original", prefix + "reading-pdf-1.json"]

    reading = sources.reading(lead, first.source_id)
    assert [s.locator.page for s in reading.segments] == [1, 2]
    name, data = sources.original(lead, first.source_id)
    assert (name, data) == ("kopie.pdf", fixture("prezentace.pdf"))


def test_a_source_is_found_only_within_its_client(
    scoped: Any, sources: ClientSourceRepository
) -> None:
    lead = client_scope(scoped, "lead")
    kept, _ = sources.add(lead, filename="a.txt", data=b"Fakt.")
    other = client_scope(scoped, "other_lead", "other")
    assert sources.sources(other) == []
    for read in (sources.source, sources.reading, sources.original):
        with pytest.raises(ScopeDenied):
            read(other, kept.source_id)
    # The same bytes at another client are that client's own source.
    theirs, created = sources.add(other, filename="a.txt", data=b"Fakt.")
    assert created and theirs.source_id != kept.source_id


def test_a_refused_file_stores_nothing(
    scoped: Any, sources: ClientSourceRepository, store: InMemoryArtifactStore
) -> None:
    lead = client_scope(scoped, "lead")
    for name, data, reason in (
        ("a.txt", b"", "empty"),
        ("a.txt", b"x" * (SOURCE_MAX_BYTES + 1), "too_large"),
        ("a.exe", b"MZ", "unsupported_format"),
    ):
        with pytest.raises(SourceRejected) as refused:
            sources.add(lead, filename=name, data=data)
        assert refused.value.reason == reason
    assert store.keys == [] and sources.sources(lead) == []


def test_what_cannot_be_read_is_kept_and_says_so(
    scoped: Any, sources: ClientSourceRepository, store: InMemoryArtifactStore
) -> None:
    lead = client_scope(scoped, "lead")
    image, _ = sources.add(lead, filename="logo.png", data=fixture("obrazek.png"))
    assert (image.outcome, image.reason, image.segments) == (
        ReadingOutcome.UNREADABLE,
        "not_text_media",
        0,
    )
    assert sources.reading(lead, image.source_id) == SourceReading(
        outcome=ReadingOutcome.UNREADABLE, reason="not_text_media"
    )
    # Only the original is stored: there is no reading to keep.
    assert [k.rsplit("/", 1)[-1] for k in store.keys] == ["original"]


def test_a_reading_altered_in_storage_is_refused(
    scoped: Any, sources: ClientSourceRepository, store: InMemoryArtifactStore
) -> None:
    lead = client_scope(scoped, "lead")
    kept, _ = sources.add(lead, filename="a.md", data=b"Fakt.")
    [key] = [k for k in store.keys if "reading" in k]
    body = json.loads(store.get(key))
    body["segments"][0]["text"] = "Jiný fakt."
    store.put(key, json.dumps(body).encode())
    with pytest.raises(IntegrityError):
        sources.reading(lead, kept.source_id)


def test_the_same_bytes_kept_by_another_request_first_is_one_source(
    scoped: Any, sources: ClientSourceRepository
) -> None:
    lead = client_scope(scoped, "lead")
    kept, _ = sources.add(lead, filename="a.txt", data=b"Fakt.")
    # The race: a request that looked before the other's row landed tries to keep the
    # bytes too; the unique constraint refuses it and nothing is half-written.
    assert (
        sources._keep(lead, name="b.txt", data=b"Fakt.", sha256=kept.sha256, family=kept.family)
        is None
    )
    assert [s.source_id for s in sources.sources(lead)] == [kept.source_id]


def test_adding_and_reading_need_their_permissions(
    scoped: Any, sources: ClientSourceRepository
) -> None:
    lead = client_scope(scoped, "lead")
    kept, _ = sources.add(lead, filename="a.txt", data=b"Fakt.")

    def without(permission: ClientPermission) -> ClientContext:
        return dataclasses.replace(
            lead, permissions=frozenset(p for p in lead.permissions if p is not permission)
        )

    with pytest.raises(ScopeDenied):
        sources.add(without(ClientPermission.PROPOSE_CLIENT_KNOWLEDGE), filename="b.txt", data=b"x")
    with pytest.raises(ScopeDenied):
        sources.sources(without(ClientPermission.VIEW_CLIENT_KNOWLEDGE))
    with pytest.raises(ScopeDenied):
        sources.reading(without(ClientPermission.VIEW_CLIENT_KNOWLEDGE), kept.source_id)


def test_a_person_marks_a_source_internal_and_unmarks_it_with_who_when_and_audit(
    scoped: Any, sources: ClientSourceRepository, session: Any
) -> None:
    """The product owner's rule of 2026-10-19: a source marked "Interní dokument" is Class C
    for AI. Who marked it and when are kept; each change is audited; a repeat changes
    nothing; another client's source is not found."""
    lead = client_scope(scoped, "lead")
    kept, _ = sources.add(lead, filename="metodika.md", data=b"# Interni postup AIA")
    assert kept.internal is None

    marked, changed = sources.mark_internal(lead, kept.source_id, internal=True)
    assert changed and marked.internal is not None
    assert marked.internal.by == lead.actor_id and marked.internal.at is not None
    assert sources.source(lead, kept.source_id).internal == marked.internal
    assert [s.internal for s in sources.sources(lead)] == [marked.internal]

    again, changed = sources.mark_internal(lead, kept.source_id, internal=True)
    assert not changed and again.internal == marked.internal

    unmarked, changed = sources.mark_internal(lead, kept.source_id, internal=False)
    assert changed and unmarked.internal is None
    assert sources.mark_internal(lead, kept.source_id, internal=False)[1] is False

    audits = session.scalars(
        select(AccessAuditRow)
        .where(AccessAuditRow.action.like("CLIENT_SOURCE_%INTERNAL"))
        .order_by(AccessAuditRow.event_id)
    ).all()
    assert [(a.action, a.actor_id, a.payload["source_id"]) for a in audits] == [
        ("CLIENT_SOURCE_MARKED_INTERNAL", lead.actor_id, kept.source_id),
        ("CLIENT_SOURCE_UNMARKED_INTERNAL", lead.actor_id, kept.source_id),
    ]
    assert audits[1].payload["marked_by"] == lead.actor_id

    other = client_scope(scoped, "other_lead", "other")
    with pytest.raises(ScopeDenied):
        sources.mark_internal(other, kept.source_id, internal=True)


def test_marking_a_source_internal_needs_the_permission_adding_it_needs(
    scoped: Any, sources: ClientSourceRepository
) -> None:
    lead = client_scope(scoped, "lead")
    kept, _ = sources.add(lead, filename="a.txt", data=b"Fakt.")
    without = dataclasses.replace(
        lead,
        permissions=frozenset(
            p for p in lead.permissions if p is not ClientPermission.PROPOSE_CLIENT_KNOWLEDGE
        ),
    )
    with pytest.raises(ScopeDenied):
        sources.mark_internal(without, kept.source_id, internal=True)
    assert sources.source(lead, kept.source_id).internal is None


def test_a_source_is_not_knowledge_and_no_study_reads_it(
    scoped: Any, sources: ClientSourceRepository, session: Any
) -> None:
    lead = client_scope(scoped, "lead")
    # The marker has letters no generated id can hold (ids are a prefix and hex): a digit
    # marker once flaked when a fresh item id happened to contain it.
    sources.add(lead, filename="tajne.md", data="Tajný obsah zdroje qzxw.".encode())
    knowledge = ClientKnowledgeRepository(session)
    assert knowledge.items(lead) == []
    knowledge.add(lead, kind=KnowledgeKind.FACT, title="Fakt")
    consumed = knowledge.for_study(scoped.scope(user="researcher"))
    snapshot = json.dumps(context_snapshot({"title": "Studie"}, consumed), ensure_ascii=False)
    assert "qzxw" not in snapshot and "tajne" not in snapshot
