"""The internal report's design chapter and questionnaire appendix (complete-study-report B1b)."""

from __future__ import annotations

import io
import zipfile
from typing import Any

from docx import Document
from test_report_docx_render import _meta

from aia_core.application.analysis_results import Brief
from aia_core.application.report_design import NOT_STATED, compose_design_sections
from aia_core.application.report_tables import scale_notes
from aia_core.domain.report.document_model import document_bytes
from aia_core.domain.report.model import (
    BulletList,
    Callout,
    CalloutKind,
    Heading,
    Paragraph,
    ReportDocument,
    Section,
    Text,
)
from aia_core.domain.research_design import compile_design
from aia_core.infrastructure.document_text import paragraph_text
from aia_core.infrastructure.report_docx.lint import lint_docx
from aia_core.infrastructure.report_docx.renderer import DocxRenderer

DESIGN: dict[str, Any] = {
    "title": "Ranní nápoj",
    "n": 300,
    "audience": {
        "strategy": "filters",
        "source_mode": "population",
        "filters": {"vek": {"min": 25, "max": 55}, "kraj": ["Praha"]},
    },
    "sections": [
        {
            "type": "questions",
            "questions": [
                {"id": "q1", "text": "Jak často pijete kávu?", "typ": "skala", "skala": [1, 5]},
                {
                    "id": "q2",
                    "text": "Co si ráno koupíte?",
                    "typ": "vyber",
                    "kategorie": ["Kávu", "Čaj", "Nic"],
                },
                {"id": "q3", "text": "Proč?", "typ": "otevrena"},
            ],
        },
        {
            "type": "object_battery",
            "title": "Nápoje",
            "object_family": "nápoje",
            "objects": ["Káva", "Čaj", "Kakao", "Džus"],
            "object_question": "Jak hodnotíte {object}?",
            "scale": [1, 10],
        },
    ],
}


def _words(blocks: Any) -> str:
    out: list[str] = []
    for b in blocks:
        if isinstance(b, Paragraph):
            out.append("".join(i.text for i in b.content if isinstance(i, Text)))
        elif isinstance(b, BulletList):
            out.extend(
                "".join(i.text for i in item.content if isinstance(i, Text)) for item in b.items
            )
    return "\n".join(out)


def test_the_chapter_names_the_sets_the_sample_and_the_audience_the_design_stated() -> None:
    spec, problems = compile_design(DESIGN)
    assert spec is not None, problems
    design = compose_design_sections(spec, ("Co lidé ráno pijí?",))
    chapter = _words(design.chapter.blocks)
    assert "Co lidé ráno pijí?" in chapter
    assert "Nápoje" in chapter and all(o in chapter for o in ("Káva", "Čaj", "Kakao", "Džus"))
    assert "„Jak hodnotíte …?“" in chapter and "1 (" in chapter and "10 (" in chapter
    assert "300 respondenty" in chapter
    assert "zúžil podle faktorů populace" in chapter
    assert "vek: 25\u201355" in chapter and "kraj: Praha" in chapter


def test_a_design_that_states_no_sample_or_audience_says_so_rather_than_guessing() -> None:
    spec, problems = compile_design({**DESIGN, "n": None, "audience": {}})
    assert spec is not None, problems
    chapter = _words(compose_design_sections(spec, ()).chapter.blocks)
    assert "Návrh neuvádí velikost vzorku" in chapter
    assert "Cílovou skupinou je celá populace." in chapter


def test_the_appendix_lists_every_question_as_asked_and_each_set() -> None:
    spec, problems = compile_design(DESIGN)
    assert spec is not None, problems
    appendix = compose_design_sections(spec, ()).appendix
    assert appendix.appendix and appendix.id == "app-questionnaire"
    words = _words(appendix.blocks)
    assert "Jak často pijete kávu? — škála 1\u20135" in words
    assert "Co si ráno koupíte? — jedna odpověď: Kávu, Čaj, Nic" in words
    assert "Proč? — otevřená odpověď" in words
    assert "Káva, Čaj, Kakao, Džus" in words


def test_a_sets_count_agrees_with_its_words() -> None:
    # "4 položek hodnocených" broke Czech agreement in every report with a 2-4 item set
    # (cross-review); 2-4 take "položky hodnocené", 5 and more "položek hodnocených".
    spec, problems = compile_design(DESIGN)
    assert spec is not None, problems
    assert "Nápoje — 4 položky hodnocené otázkou „Jak hodnotíte …?“" in _words(
        compose_design_sections(spec, ()).chapter.blocks
    )
    battery = {**DESIGN["sections"][1], "objects": ["Káva", "Čaj", "Kakao", "Džus", "Voda"]}
    spec, problems = compile_design({**DESIGN, "sections": [DESIGN["sections"][0], battery]})
    assert spec is not None, problems
    assert "Nápoje — 5 položek hodnocených otázkou" in _words(
        compose_design_sections(spec, ()).chapter.blocks
    )


def test_a_scales_labels_are_printed_beside_its_values() -> None:
    """input-interactions 3b: what respondents were told a point means is in the appendix and
    under the point's table, the two ends or every point."""
    every = [f"bod {i}" for i in range(1, 11)]
    labelled = {
        **DESIGN,
        "sections": [
            {
                "type": "questions",
                "questions": [
                    {**DESIGN["sections"][0]["questions"][0], "popisky_skaly": ["nikdy", "denně"]},
                    *DESIGN["sections"][0]["questions"][1:],
                ],
            },
            {**DESIGN["sections"][1], "scale_labels": every},
        ],
    }
    spec, problems = compile_design(labelled)
    assert spec is not None, problems
    words = _words(compose_design_sections(spec, ()).appendix.blocks)
    assert "Jak často pijete kávu? — škála 1\u20135 (1 = nikdy, 5 = denně)" in words
    assert "škála 1\u201310 (1 = bod 1, 2 = bod 2," in words and "10 = bod 10)" in words
    chapter = _words(compose_design_sections(spec, ()).chapter.blocks)
    assert "na škále 1 (bod 1) až 10 (bod 10)" in chapter
    notes = scale_notes(spec)
    assert notes["q1"] == "Škála: 1 = nikdy, 5 = denně."
    assert notes["napoje_obj_kava"].startswith("Škála: 1 = bod 1, 2 = bod 2,")
    assert "q2" not in notes
    bare, _ = compile_design(DESIGN)
    assert bare is not None and "q1" not in scale_notes(bare)
    assert scale_notes(bare)["napoje_obj_kava"] == "Škála: 1 = vůbec, 10 = velmi."


def test_the_audiences_description_and_subpanel_are_kept_and_printed_when_given() -> None:
    """Compiler 3 keeps what a person said about the audience (artifact-chain CHN5): the
    description and the built-in subpanel, each only when given, printed before the filters."""
    stated = {
        **DESIGN,
        "audience": {
            **DESIGN["audience"],
            "description": "Pražané, kteří ráno pijí kávu",
            "builtin_subpanel": "young_18_29",
        },
    }
    spec, problems = compile_design(stated)
    assert spec is not None, problems
    assert spec.audience["description"] == "Pražané, kteří ráno pijí kávu"
    assert spec.audience["builtin_subpanel"] == "young_18_29"
    words = _words(compose_design_sections(spec, ()).chapter.blocks)
    assert "Cílová skupina podle návrhu: Pražané, kteří ráno pijí kávu." in words
    assert "vestavěného podpanelu young_18_29" in words
    assert "kraj: Praha" in words

    plain, _ = compile_design(DESIGN)
    assert plain is not None
    # Neither given: the body has neither key, so its fingerprint is what it was under this
    # compiler, and the chapter says nothing it was not told.
    assert set(plain.audience) == {"source_mode", "strategy", "has_filters"}
    blank, _ = compile_design({**DESIGN, "audience": {**DESIGN["audience"], "description": "  "}})
    assert blank is not None and blank.fingerprint() == plain.fingerprint()
    assert "podle návrhu" not in _words(compose_design_sections(plain, ()).chapter.blocks)


def test_a_subpanel_left_from_a_population_choice_is_not_an_external_audiences() -> None:
    """The audience step keeps the last built-in subpanel when the source moves outside AIA's
    population (OI-51); an external audience did not come from it, so it is neither compiled
    nor printed (Codex on #587)."""
    external = {
        **DESIGN,
        "audience": {
            **DESIGN["audience"],
            "source_mode": "customer",
            "builtin_subpanel": "young_18_29",
        },
    }
    spec, problems = compile_design(external)
    assert spec is not None, problems
    assert "builtin_subpanel" not in spec.audience
    words = _words(compose_design_sections(spec, ()).chapter.blocks)
    assert "podpanelu" not in words
    assert "mimo populaci AIA" in words


def test_the_chapter_opens_with_what_the_brief_asked_for_before_how_it_was_measured() -> None:
    """F19: the draft said "co měl výzkum zjistit" and printed only how it measured it."""
    spec, problems = compile_design(DESIGN)
    assert spec is not None, problems
    blocks = compose_design_sections(
        spec,
        ("Co lidé ráno pijí?",),
        brief=Brief(
            goal="Zjistit, zda dojíždějící koupí ranní nápoj do 60 Kč.",
            decision_use="Rozhodnutí o uvedení nápoje.",
        ),
    ).chapter.blocks
    chapter = _words(blocks)
    goal, use, question = (
        chapter.index("Zjistit, zda dojíždějící koupí ranní nápoj do 60 Kč."),
        chapter.index("Rozhodnutí o uvedení nápoje."),
        chapter.index("Co lidé ráno pijí?"),
    )
    assert goal < use < question
    headings = [b.text for b in blocks if isinstance(b, Heading)]
    assert headings[:3] == ["Cíl studie", "Rozhodnutí, které má výzkum podpořit", "Výzkumné otázky"]


def test_an_empty_brief_field_is_said_not_stated_never_left_out() -> None:
    """A reader tells "the person wrote no goal" from "the report does not show it" (review)."""
    spec, problems = compile_design(DESIGN)
    assert spec is not None, problems
    blocks = compose_design_sections(
        spec, (), brief=Brief(goal=None, decision_use=None)
    ).chapter.blocks
    headings = [b.text for b in blocks if isinstance(b, Heading)]
    assert headings[:2] == ["Cíl studie", "Rozhodnutí, které má výzkum podpořit"]
    assert _words(blocks).count(NOT_STATED) == 2
    # Without a brief (the client report) the chapter has neither heading.
    plain = compose_design_sections(spec, ()).chapter.blocks
    assert "Cíl studie" not in [b.text for b in plain if isinstance(b, Heading)]


def test_a_goal_in_two_paragraphs_prints_as_two() -> None:
    spec, problems = compile_design(DESIGN)
    assert spec is not None, problems
    blocks = compose_design_sections(
        spec, (), brief=Brief(goal="První odstavec.\n\nDruhý odstavec.")
    ).chapter.blocks
    written = [_words((b,)) for b in blocks if isinstance(b, Paragraph)]
    assert "První odstavec." in written and "Druhý odstavec." in written


def test_a_control_character_in_the_brief_never_fails_the_report(report_ledger: Any) -> None:
    """A soft line break pasted from Word (\\x0b) made python-docx refuse the draft (review)."""
    spec, problems = compile_design(DESIGN)
    assert spec is not None, problems
    chapter = compose_design_sections(
        spec, (), brief=Brief(goal="a\x0bb\x07c", decision_use="\x01")
    ).chapter
    assert "a\nbc" in _words(chapter.blocks) and NOT_STATED in _words(chapter.blocks)
    doc = ReportDocument(
        _meta(),
        (Section("Návrh", (Callout(CalloutKind.METHOD, ()), *chapter.blocks)),),
        report_ledger(),
    )
    assert lint_docx(DocxRenderer().render(doc)) == []


def test_a_control_character_a_person_typed_never_fails_the_render(report_doc: Any) -> None:
    """A research question or a question's wording carrying a character XML 1.0 cannot hold
    (a vertical tab pasted from Word, a bell, a lone surrogate) failed the whole draft in
    python-docx. Word's breaks print as a space, the rest are dropped, and the stored document
    model holds the same text."""
    design: dict[str, Any] = {**DESIGN}
    sections = list(design["sections"])
    first = dict(sections[0])
    first["questions"] = [{**first["questions"][0], "text": "Jak často\x07 pijete kávu?"}]
    design["sections"] = [first, *sections[1:]]
    spec, problems = compile_design(design)
    assert spec is not None, problems
    composed = compose_design_sections(spec, ("Co lidé\x0bráno pijí?\x1f\ud800",))
    doc = report_doc(
        *composed.chapter.blocks, *composed.appendix.blocks, title="Důvěra\x0c v banky"
    )

    out = DocxRenderer().render(doc)

    assert lint_docx(out) == []
    read = Document(io.BytesIO(out))
    # Read as AIA's own reader reads it: inside the locked controls too (RPT8a).
    cells = (p for t in read.tables for row in t.rows for c in row.cells for p in c.paragraphs)
    words = "\n".join(paragraph_text(p) for p in (*read.paragraphs, *cells))
    assert "Co lidé ráno pijí?" in words  # the research question, its break a space
    assert "Jak často pijete kávu?" in words  # the question as asked, its bell dropped
    assert "\x07" not in words and "\x1f" not in words
    assert read.core_properties.title == "Důvěra  v banky"
    assert 'Co lidé ráno pijí?"' in document_bytes(doc).decode()  # as the DOCX printed it


def test_the_designs_numbers_print_under_the_specification_it_came_from(report_doc: Any) -> None:
    """A scale's ends, a count, the planned n and an audience bound are code's words from the
    specification the run executed: each block printing one is locked under its fingerprint
    (artifact-report RPT8a, part 3, P10); what the person wrote alone (a research question, an
    item's name, the brief) stays as it is."""
    spec, problems = compile_design(DESIGN)
    assert spec is not None, problems
    tag = f"design:{spec.fingerprint()[:16]}"
    composed = compose_design_sections(
        spec, ("Co lidé ráno pijí?",), brief=Brief(goal="Zjistit, co lidé pijí.")
    )
    for block in (*composed.chapter.blocks, *composed.appendix.blocks):
        if isinstance(block, Paragraph | BulletList) and any(c.isdigit() for c in _words((block,))):
            assert block.control == tag, _words((block,))
    personal = [b for b in composed.chapter.blocks if isinstance(b, BulletList)]
    assert personal and all(b.control is None for b in personal)  # questions, items: no digits

    out = DocxRenderer().render(report_doc(*composed.chapter.blocks, *composed.appendix.blocks))
    with zipfile.ZipFile(io.BytesIO(out)) as z:
        body = z.read("word/document.xml").decode()
    assert f'w:val="{tag}"' in body and "Návrh počítal s 300 respondenty." in body
