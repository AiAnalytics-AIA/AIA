"""The report's outline numbers and the DOCX style sheet built from print tokens."""

from __future__ import annotations

from dataclasses import replace
from datetime import date
from typing import Any

from docx import Document
from docx.document import Document as WordDocument
from docx.oxml.ns import qn

from aia_core.domain.evidence.validation import METHOD_STATUS_PENDING
from aia_core.domain.report.model import (
    Chart,
    ChartKind,
    Figure,
    Heading,
    ReportDocument,
    ReportKind,
    ReportMeta,
    Section,
    Series,
)
from aia_core.domain.report.outline import appendix_letter, build_outline
from aia_core.domain.report.print_tokens import STYLES
from aia_core.infrastructure.report_docx.styles import (
    KEPT,
    PARAGRAPHS,
    SHORT_TABLE_ROWS,
    S,
    install_styles,
    kept_style,
    row_keeps,
)


def test_appendix_letters() -> None:
    assert [appendix_letter(i) for i in (0, 25, 26)] == ["A", "Z", "AA"]


def test_outline_numbers_chapters_headings_figures_and_appendices(report_ledger: Any) -> None:
    fig = Figure(
        "fig-1",
        "Důvěra",
        Chart(ChartKind.BAR, ("Praha",), (Series("x", ("trust_praha",)),), "%"),
        "AIA",
        "alt",
    )
    meta = ReportMeta(
        ReportKind.CLIENT, "T", "S", "C", "St", "STU-1", date(2026, 9, 25), 1, METHOD_STATUS_PENDING
    )
    doc = ReportDocument(
        meta,
        (
            Section("Shrnutí", (Heading("A", 2), Heading("B", 3), fig), id="ch-1"),
            Section("Metodika", (Heading("C", 2, id="h-c"),)),
            Section("Dotazník", (), appendix=True),
        ),
        report_ledger(),
    )
    o = build_outline(doc)
    assert o.chapter_numbers == ("1", "2", "A")
    assert [h.number for h in o.headings] == ["1", "1.1", "1.1.1", "2", "2.1", "A"]
    assert o.labels == {"ch-1": "kapitola 1", "fig-1": "graf 1", "h-c": "oddíl 2.1"}


def test_every_paragraph_style_comes_from_a_print_token() -> None:
    assert {spec.token for spec in PARAGRAPHS.values()} <= set(STYLES)


def test_the_style_sheet_installs_with_no_theme_fonts() -> None:
    document = Document()
    install_styles(document)
    names = {s.name for s in document.styles}
    for name in (S.BODY, S.H1, S.CAPTION, S.TOC_1, S.FOOTNOTE, S.DATA_TABLE, S.STRONG):
        assert name in names, name
    xml = document.styles.element.xml
    assert "asciiTheme" not in xml and "hAnsiTheme" not in xml
    h1 = document.styles[S.H1].element
    assert h1.find(qn("w:rPr")).find(qn("w:rFonts")).get(qn("w:ascii")) == STYLES["doc-h1"].font


def test_starter_theme_furniture_cannot_override_the_aia_template() -> None:
    from aia_core.infrastructure.report_docx.ooxml import el

    document = Document()
    props = document.styles["Title"].element.get_or_add_pPr()
    props.append(el("w:pBdr"))
    props.append(el("w:shd", fill="FF0000"))
    props.append(el("w:contextualSpacing"))
    install_styles(document)
    title = document.styles[S.TITLE].element.get_or_add_pPr()
    assert all(title.find(qn(tag)) is None for tag in ("w:pBdr", "w:shd", "w:contextualSpacing"))
    assert PARAGRAPHS[S.FINDING_TITLE].token == "doc-body"


def _on(document: WordDocument, name: str, tag: str) -> bool:
    """Whether the installed style ``name`` sets the paragraph property ``tag`` on."""
    ppr = document.styles[name].element.find(qn("w:pPr"))
    prop = None if ppr is None else ppr.find(qn(f"w:{tag}"))
    return prop is not None and prop.get(qn("w:val")) not in ("0", "false", "off")


def _installed() -> WordDocument:
    document = Document()
    install_styles(document)
    return document


def test_only_the_contents_starts_a_page_of_its_own() -> None:
    # Chapters and appendices flow (research-integrity-report D5): no heading but the
    # contents' breaks the page; the renderer's sections break it where the furniture
    # changes.
    assert [n for n, p in PARAGRAPHS.items() if p.page_break_before] == [S.FRONT_HEADING_BREAK]
    document = _installed()
    for name in PARAGRAPHS:
        assert _on(document, name, "pageBreakBefore") is (name == S.FRONT_HEADING_BREAK), name


def test_headings_captions_and_figures_keep_with_what_follows() -> None:
    kept = {
        S.H1,
        S.H2,
        S.H3,
        S.APPENDIX,
        S.FRONT_HEADING,
        S.FRONT_HEADING_BREAK,
        S.H2_FRONT,
        S.CAPTION,
        S.CALLOUT_TITLE,
        S.FINDING_TITLE,
        S.FINDING_LABEL,
        S.LEAD_IN,
    }
    document = _installed()
    for name in kept | {S.FIGURE}:  # caption -> picture -> the first source line
        assert _on(document, name, "keepNext"), name
    for name in kept:  # a heading, caption or title is never split across a page
        assert _on(document, name, "keepLines"), name
    assert not _on(document, S.SOURCE, "keepNext")  # the chain ends at the source line


def test_each_kept_twin_is_its_style_keeping_with_the_next() -> None:
    # A table row and a note line keep with what follows them through a style: the lint
    # forbids a paragraph's own keepNext (research-integrity-report D5, review documents#4).
    assert set(KEPT) == {
        S.TABLE,
        S.TABLE_NUMBER,
        S.TABLE_HEAD,
        S.TABLE_HEAD_NUMBER,
        S.MONO,
        S.SOURCE,
    }
    document = _installed()
    for style, twin in KEPT.items():
        assert PARAGRAPHS[twin] == replace(PARAGRAPHS[style], keep_next=True), twin
        assert _on(document, twin, "keepNext") and not _on(document, style, "keepNext"), twin
        assert kept_style(style, True) == twin and kept_style(style, False) == style
    # Under a figure: the picture keeps with "Zdroj", each line with the next, the last ends it.
    assert _on(document, S.SOURCE_KEPT, "keepNext") and not _on(document, S.SOURCE, "keepNext")


def test_a_short_table_keeps_every_row_and_a_long_one_its_ends() -> None:
    n = SHORT_TABLE_ROWS
    assert row_keeps(0, notes=True) == []
    assert row_keeps(1, notes=True) == [True] and row_keeps(1, notes=False) == [False]
    assert row_keeps(n, notes=True) == [True] * n  # whole, the last row with the notes row
    assert row_keeps(n, notes=False) == [True] * (n - 1) + [False]
    long = n + 4
    # never broken after its first row or before its last; the last keeps with the notes
    assert row_keeps(long, notes=True) == [True] + [False] * (long - 3) + [True, True]
    assert row_keeps(long, notes=False) == [True] + [False] * (long - 3) + [True, False]


def test_the_data_table_style_draws_no_rule_under_its_notes_row() -> None:
    style = _installed().styles[S.DATA_TABLE].element
    (last,) = [p for p in style.findall(qn("w:tblStylePr")) if p.get(qn("w:type")) == "lastRow"]
    bottom = last.find(f"{qn('w:tcPr')}/{qn('w:tcBorders')}/{qn('w:bottom')}")
    assert bottom is not None and bottom.get(qn("w:val")) == "nil"


def test_every_paragraph_style_avoids_widows_and_orphans() -> None:
    document = _installed()
    for name in PARAGRAPHS:
        assert _on(document, name, "widowControl"), name


def test_a_chapter_heading_has_more_space_above_it_than_a_section_heading() -> None:
    # A chapter no longer opens a page, so the air above it is what tells it from a
    # section. The rule, not the number: a token edit cannot bring back 0 pt.
    assert STYLES["doc-h1"].before_pt > STYLES["doc-h2"].before_pt
    document = _installed()
    h2 = document.styles[S.H2].paragraph_format.space_before
    for name in (S.H1, S.APPENDIX):
        assert document.styles[name].paragraph_format.space_before > h2, name
