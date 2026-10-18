"""The approved export of a stored report version: its footers' status runs, and nothing else.

On real rendered sample reports (``report_samples``): a stamp leaves every part of the package
but the footers byte for byte as stored, changes each footer only where its status run's words
were, is deterministic and passes the lint. A document whose renderer's status marks are
anywhere else -- every version stored before this layout -- is not stampable, never partly
stamped. The study's own words are not a status: a questionnaire item on a loan not approved
stamps, and each template's own words are held here to state no status outside the footers.
"""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from dataclasses import replace
from datetime import date
from typing import Any

import pytest
import report_samples
from docx import Document

from aia_core.domain.report.model import Paragraph, text
from aia_core.domain.report.review import ReviewMark, ReviewState, approved_words
from aia_core.infrastructure.report_docx.lint import lint_docx
from aia_core.infrastructure.report_docx.renderer import DocxRenderer, write_package
from aia_core.infrastructure.report_docx.review_stamp import NotStampable, stamp_review

DRAFT = "KONCEPT \u2014 NESCHVÁLENO"
COVER = f"{DRAFT}. Tato verze dosud nebyla schválena. Nepředávejte ji jako finální."


def _mark(data: bytes, *, artifact_id: str = "ART-0123456789abcdef") -> ReviewMark:
    sha = hashlib.sha256(data).hexdigest()
    return ReviewMark(ReviewState.APPROVED, artifact_id, sha, date(2026, 10, 18))


def _infos(data: bytes) -> list[zipfile.ZipInfo]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return z.infolist()


def _parts(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return {n: z.read(n) for n in z.namelist()}


def _footer(name: str) -> bool:
    return name.startswith("word/footer") and name.endswith(".xml")


#: The rendered samples, once per module: rendering is deterministic, and the slow part.
_RENDERED: dict[str, bytes] = {}


@pytest.fixture
def samples(report_ledger: Any) -> dict[str, bytes]:
    if not _RENDERED:
        for kind in report_samples.KINDS:
            _RENDERED[kind] = DocxRenderer().render(report_samples.build(kind, report_ledger))
    return _RENDERED


#: A status in words: the draft words' first one as the renderer writes it, or a word beginning
#: "neschválen" in any case. The stamp never reads these (a study's words may say them); the
#: templates' own words are held to them here.
_STATUS_WORDS = re.compile(r"KONCEPT|(?<!\w)(?i:neschválen)")


# Every template stamps, and none of its own words outside the footers states a status.
@pytest.mark.parametrize("kind", report_samples.KINDS)
def test_a_stamp_changes_only_the_status_runs_of_the_footers(
    samples: dict[str, bytes], kind: str
) -> None:
    stored = samples[kind]
    mark = _mark(stored)
    stamped = stamp_review(stored, mark)
    words = approved_words(mark)

    before, after = _infos(stored), _infos(stamped)
    assert [i.filename for i in after] == [i.filename for i in before]  # the stored order
    old, new = _parts(stored), _parts(stamped)
    footers = [n for n in old if _footer(n)]
    assert len(footers) >= 3  # cover, front matter, body
    for name, data in old.items():
        if name.endswith(".xml") and not _footer(name):
            assert not _STATUS_WORDS.search(data.decode()), name
    for b, a in zip(before, after, strict=True):
        assert a.date_time == b.date_time == (1980, 1, 1, 0, 0, 0)
        if _footer(b.filename):
            continue
        # Every other part byte for byte, compressed alike.
        assert new[b.filename] == old[b.filename], b.filename
        assert (a.CRC, a.compress_size) == (b.CRC, b.compress_size), b.filename
    for name in footers:
        # Each footer changed exactly where its one status run's words were.
        said = COVER if COVER.encode() in old[name] else DRAFT
        assert old[name].count(said.encode()) == 1, name
        assert new[name] == old[name].replace(said.encode(), words.encode()), name
        assert DRAFT.encode() not in new[name]

    assert stamp_review(stored, mark) == stamped  # deterministic
    assert lint_docx(stamped) == []
    Document(io.BytesIO(stamped))  # Word's reader opens it


def test_an_approved_footer_reads_its_date_and_version_beside_the_classification(
    samples: dict[str, bytes],
) -> None:
    stored = samples["client"]
    stamped = stamp_review(stored, _mark(stored))
    document = Document(io.BytesIO(stamped))
    texts = [s.footer.paragraphs[0].text for s in document.sections]
    nb = "\u00a0"
    approved = f"SCHVÁLENO 18.{nb}10.{nb}2026 \u00b7 verze ART-0123456789abcdef"
    assert texts[0] == approved  # the cover: the status alone
    assert all(t.startswith(f"{approved} \u00b7 Důvěrné") for t in texts[1:])


def test_only_an_approved_mark_for_exactly_these_bytes_stamps(samples: dict[str, bytes]) -> None:
    stored = samples["client"]
    draft = ReviewMark(ReviewState.DRAFT, "ART-0123456789abcdef", _mark(stored).sha256)
    with pytest.raises(ValueError, match="only an approved"):
        stamp_review(stored, draft)
    with pytest.raises(ValueError, match="not version"):
        stamp_review(stored + b"\0", _mark(stored))


def _rewritten(data: bytes, edit: dict[str, Any]) -> bytes:
    """The package with some parts' bytes rewritten by ``edit[name](bytes)``."""
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        parts = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    return write_package((n, edit[n](b) if n in edit else b) for n, b in parts)


def test_a_version_stored_before_this_layout_is_not_stampable(samples: dict[str, bytes]) -> None:
    # The layout before research-integrity-report D2: the status style was "AIA Draft", and
    # the cover notice and the document control said draft in the body too.
    stored = samples["client"]
    footers = [n for n in _parts(stored) if _footer(n)]
    legacy = _rewritten(
        stored,
        {
            "word/styles.xml": lambda b: b.replace(b"AIA Status", b"AIA Draft").replace(
                b"AIAStatus", b"AIADraft"
            ),
            **{n: (lambda b: b.replace(b"AIAStatus", b"AIADraft")) for n in footers},
            "word/document.xml": lambda b: b.replace(
                "Stav této verze uvádí zápatí".encode(), DRAFT.encode()
            ),
        },
    )
    with pytest.raises(NotStampable):
        stamp_review(legacy, _mark(legacy))


@pytest.mark.parametrize(
    "edit",
    [
        # The status style is gone: nothing says which runs are the status.
        lambda parts: {"word/styles.xml": lambda b: b.replace(b"AIA Status", b"AIA Other")},
        # The body says draft, as the old cover notice did.
        lambda parts: {
            "word/document.xml": lambda b: b.replace(
                "Stav této verze uvádí zápatí".encode(), DRAFT.encode()
            )
        },
        # A footer with no status run.
        lambda parts: {
            next(n for n in parts if _footer(n)): lambda b: b.replace(b"AIAStatus", b"AIAOther")
        },
        # A footer with two: the stamp would have to choose.
        lambda parts: {
            next(n for n in parts if _footer(n)): lambda b: b.replace(
                b"</w:p>",
                b'<w:r><w:rPr><w:rStyle w:val="AIAStatus"/></w:rPr><w:t>x</w:t></w:r></w:p>',
            )
        },
    ],
    ids=["no-status-style", "status-in-body", "footer-without-status", "footer-with-two"],
)
def test_a_document_that_states_its_status_other_than_once_per_footer_is_not_stampable(
    samples: dict[str, bytes], edit: Any
) -> None:
    stored = samples["client"]
    broken = _rewritten(stored, edit(_parts(stored)))
    with pytest.raises(NotStampable):
        stamp_review(broken, _mark(broken))


def _paragraph(*runs: str) -> bytes:
    return (
        "<w:p>"
        + "".join(f'<w:r><w:t xml:space="preserve">{r}</w:t></w:r>' for r in runs)
        + "</w:p>"
    ).encode()


def _appended(data: bytes, body: bytes) -> bytes:
    """The package with ``body`` at the end of the document, before its section properties."""

    def add(b: bytes) -> bytes:
        at = b.rindex(b"<w:sectPr")
        return b[:at] + body + b[at:]

    return _rewritten(data, {"word/document.xml": add})


def _with_item(report_ledger: Any, words: str) -> bytes:
    """The internal sample rendered with one more item, ``words``, in its questionnaire."""
    doc = report_samples.build("internal", report_ledger)
    sections = tuple(
        replace(s, blocks=(*s.blocks, Paragraph(text(words)))) if s.id == "app-questionnaire" else s
        for s in doc.sections
    )
    assert sections != doc.sections
    return DocxRenderer().render(replace(doc, sections=sections))


def test_the_study_s_own_words_never_hold_back_the_stamp(
    samples: dict[str, bytes], report_ledger: Any
) -> None:
    """A bank's questionnaire asks why a loan was not approved, a quoted source counts the
    applications not approved, a study is named after them: each says "neschválen…", and none
    states the version's status. Read as status, an approved version holding one was served
    with its draft footers while its status said it was stamped (cross-review). The stamp
    reads only the renderer's own status marks and leaves the study's words as stored; the
    templates' own words are held above. The renderer's draft words are still refused, even
    split across the runs of one paragraph."""
    stored = samples["client"]
    heads = {n for n, b in _parts(stored).items() if n.startswith("word/header") and b"AIA" in b}
    question = "Schválení úvěru proběhlo rychle?"
    said = [
        _with_item(report_ledger, "Q7. Jaký byl důvod neschválení vaší žádosti o hypotéku?"),
        _appended(stored, _paragraph("„Podíl neschválených žádostí o úvěr meziročně vzrostl.“")),
        _rewritten(
            stored,
            {
                n: lambda b: b.replace("Důvěra 2026".encode(), "Neschválené úvěry 2026".encode())
                for n in heads
            },
        ),
        # A yes/no item's options above a question beginning "Schválení" ("…Ne" + "Schválení…"):
        # in two paragraphs, two table cells, either side of a line break.
        _appended(stored, _paragraph("Jedna odpověď: Ano, Ne") + _paragraph(question)),
        _appended(
            stored,
            b"<w:tbl><w:tr><w:tc>"
            + _paragraph("Ne")
            + b"</w:tc><w:tc>"
            + _paragraph(question)
            + b"</w:tc></w:tr></w:tbl>",
        ),
        _appended(
            stored,
            (
                '<w:p><w:r><w:t xml:space="preserve">Ano, Ne</w:t><w:br/>'
                f'<w:t xml:space="preserve">{question}</w:t></w:r></w:p>'
            ).encode(),
        ),
    ]
    assert heads
    for document in said:
        mark = _mark(document)
        old, new = _parts(document), _parts(stamped := stamp_review(document, mark))
        for name in old:
            if _footer(name):
                assert approved_words(mark).encode() in new[name], name
                assert DRAFT.encode() not in new[name], name
            else:
                assert new[name] == old[name], name  # the study's words as stored
        assert lint_docx(stamped) == []
    refused = _appended(stored, _paragraph("Stav: KONCEPT \u2014 NESCHVÁ", "LENO"))
    with pytest.raises(NotStampable):
        stamp_review(refused, _mark(refused))
