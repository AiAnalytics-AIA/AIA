"""A questionnaire file imports into AIA exactly as it imported into 18.6.6 (ADR 0018).

``fixtures/questionnaire_import`` holds fictional files and the unit's own
``import_questionnaire_payload`` answer on each -- the normalized sections, the summary,
or the error -- captured by ``tools/questionnaire_import_capture.py``. The unit's own
template is read in place from the vendored unit; AIA's template is one of the inputs,
so the capture also shows the unit accepts what AIA offers for download.

A rule the unit had is compared word for word. A file the unit could not read at all
failed with a Python exception's text; AIA refuses the same files as ``unreadable``.
"""

from __future__ import annotations

import hashlib
import io
import json
from pathlib import Path
from typing import Any

import pytest

from aia_core.domain.questionnaire_import import (
    QuestionnaireImportRejected,
    clean_text,
    import_rows,
    slugify,
    unique_id,
)
from aia_core.infrastructure import questionnaire_file
from aia_core.infrastructure.questionnaire_file import (
    TEMPLATE_EXAMPLES,
    TEMPLATE_HEADER,
    import_questionnaire,
    template_xlsx,
)

FIXTURES = Path(__file__).parent / "fixtures" / "questionnaire_import"
UNIT = Path(__file__).resolve().parents[3] / "legacy" / "npc-panel-18.6.6" / "app"
INDEX = json.loads((FIXTURES / "index.json").read_text())
EXPECTED: dict[str, dict[str, Any]] = json.loads(
    (FIXTURES / "expected.json").read_text(encoding="utf-8")
)
UNIT_TEMPLATE = "@unit/NPC_QUESTIONNAIRE_TEMPLATE_v17_6_0.xlsx"


def _input(name: str) -> bytes:
    if name.startswith("@unit/"):
        return (UNIT / name.removeprefix("@unit/")).read_bytes()
    return (FIXTURES / "inputs" / name).read_bytes()


def test_the_fixtures_are_the_ones_captured() -> None:
    assert sorted(INDEX["inputs"]) == sorted(EXPECTED)
    for name, digest in INDEX["inputs"].items():
        assert hashlib.sha256(_input(name)).hexdigest() == digest, name
    assert (
        hashlib.sha256((FIXTURES / "expected.json").read_bytes()).hexdigest()
        == INDEX["expected_sha256"]
    )


def test_the_unit_sources_the_fixtures_came_from_are_unchanged() -> None:
    """If any of these changed, the fixtures no longer describe the unit: recapture."""
    import ast

    source = (UNIT / "ui_server.py").read_text(encoding="utf-8")
    functions = {
        n.name: ast.get_source_segment(source, n)
        for n in ast.parse(source).body
        if isinstance(n, ast.FunctionDef)
    }
    for key, digest in INDEX["unit_sources"].items():
        if "::" in key:
            segment = functions[key.split("::")[1]]
            assert hashlib.sha256(str(segment).encode()).hexdigest() == digest, key
        else:
            assert hashlib.sha256((UNIT / key).read_bytes()).hexdigest() == digest, key


@pytest.mark.parametrize("name", sorted(EXPECTED))
def test_aia_imports_what_the_unit_imported(name: str) -> None:
    unit = EXPECTED[name]
    filename = name.split("/")[-1]
    if unit["ok"]:
        got = import_questionnaire(_input(name), filename)
        assert got.sections == unit["sections"]
        assert got.summary.model_dump() == unit["summary"]
        assert got.filename == unit["filename"]
        assert unit["plan_status"] == "questionnaire_ready"  # what the stage sets
        return
    with pytest.raises(QuestionnaireImportRejected) as refused:
        import_questionnaire(_input(name), filename)
    if unit["error_type"] == "ValueError":
        # One of the unit's own rules: the same words.
        assert str(refused.value) == unit["error"]
        assert refused.value.reason != "unreadable"
    else:
        # The unit could not read the file at all (its message was the exception's).
        assert refused.value.reason == "unreadable", unit


def test_the_fixtures_cover_every_rule_and_every_kind_of_row() -> None:
    refused = {n: e["error"] for n, e in EXPECTED.items() if not e["ok"]}
    assert len(refused) >= 15
    sections = [s for e in EXPECTED.values() if e["ok"] for s in e["sections"]]
    kinds = {q["typ"] for s in sections for q in s.get("questions", [])}
    assert kinds == {"vyber", "multi", "skala", "otevrena"}
    batteries = [s for s in sections if s["type"] == "object_battery"]
    # normalize_project drops a tracked set's scale, and StudySpec writes into its metadata.
    assert batteries and all("scale" not in b for b in batteries)
    assert all(b["metadata"]["familiarity_required"] is False for b in batteries)
    assert any("design_warning" in b["metadata"] for b in batteries)
    assert EXPECTED["aia_sablona.xlsx"]["ok"] and EXPECTED[UNIT_TEMPLATE]["ok"]


def test_aias_template_imports_as_the_units_does() -> None:
    """Same examples, so the same questionnaire -- in the unit and in AIA."""
    assert EXPECTED["aia_sablona.xlsx"]["sections"] == EXPECTED[UNIT_TEMPLATE]["sections"]
    assert EXPECTED["aia_sablona.xlsx"]["summary"] == EXPECTED[UNIT_TEMPLATE]["summary"]


def test_the_template_is_aias_own_deterministic_workbook() -> None:
    first, second = template_xlsx(), template_xlsx()
    assert first == second
    # The fixture the unit imported is the template AIA serves today.
    assert first == _input("aia_sablona.xlsx")
    unit_bytes = (UNIT / "NPC_QUESTIONNAIRE_TEMPLATE_v17_6_0.xlsx").read_bytes()
    assert first != unit_bytes  # written by AIA, not a copy of the unit's file
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(first), read_only=True)
    assert workbook.sheetnames == ["DOTAZNIK", "NAVOD"]
    rows = list(workbook["DOTAZNIK"].iter_rows(values_only=True))
    assert rows[0] == TEMPLATE_HEADER
    assert rows[1:] == [tuple(r) for r in TEMPLATE_EXAMPLES]
    guide = dict(workbook["NAVOD"].iter_rows(values_only=True))
    assert "1\u201310" in guide["skala_min / skala_max"]
    workbook.close()


def test_an_office_file_that_declares_too_much_is_not_opened(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    template = template_xlsx()
    assert import_questionnaire(template, "sablona.xlsx").summary.question_count == 2
    monkeypatch.setattr(questionnaire_file, "ZIP_MAX_UNCOMPRESSED", 100)
    with pytest.raises(QuestionnaireImportRejected) as refused:
        import_questionnaire(template, "sablona.xlsx")
    assert refused.value.reason == "unreadable"


def test_the_units_text_and_id_rules() -> None:
    assert clean_text("  <b>Ahoj</b>\t světe ") == "Ahoj světe"
    assert clean_text("&amp; <item>x</item>") == "& x"
    assert clean_text("\n".join("abcdef")) == "abcdef"  # one character per line
    assert clean_text(None, "fallback") == "fallback"
    assert slugify("Značka Č. 1") == "znacka_c_1"
    assert slugify("***") == "x"
    used: set[str] = set()
    assert [unique_id(x, used) for x in ("Q1", "Q1", "1", "")] == ["q1", "q1_2", "q_1", "x"]
    assert unique_id("1", set(), "sec") == "sec_1"


def test_rows_are_read_by_their_header_whatever_the_column_order() -> None:
    rows = [["TYP", " Otazka ", "ID", "moznosti"], ["vyber", "Ano, nebo ne?", "A", "Ano|Ne"]]
    got = import_rows(rows, filename="x.csv")
    assert got.sections[0]["questions"][0]["kategorie"] == ["Ano", "Ne"]
    assert got.sections[0]["title"] == "Dotazník"
