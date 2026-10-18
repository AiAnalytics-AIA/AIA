"""The DOCX renderer's furniture: sections, running heads, front matter, the package.

Structural tests on the package XML. They pin what a reader relies on and a
previewer cannot show: which section pages how, what every footer says, that the
contents are real fields with cached entries, and that the package is one Word
opens (schema order) and the same content always gives the same bytes.
"""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from dataclasses import replace
from datetime import date
from typing import Any

import pytest
from docx import Document
from lxml import etree

from aia_core.domain.evidence import ClaimSurface
from aia_core.domain.evidence.validation import METHOD_STATUS_PENDING
from aia_core.domain.report.model import (
    Callout,
    CalloutKind,
    Classification,
    CrossRef,
    Footnote,
    Heading,
    Link,
    Paragraph,
    ReportDocument,
    ReportKind,
    ReportMeta,
    Section,
    Text,
    text,
)
from aia_core.domain.report.validation import ReportInvalid
from aia_core.infrastructure.report_docx.lint import lint_docx
from aia_core.infrastructure.report_docx.ooxml import SECTPR_ORDER, SETTINGS_ORDER, in_schema_order
from aia_core.infrastructure.report_docx.renderer import DocxRenderer

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
DRAFT = "KONCEPT \u2014 NESCHVÁLENO"


def _meta(**overrides: Any) -> ReportMeta:
    base = ReportMeta(
        ReportKind.CLIENT,
        "Důvěra v banky 2026",
        "Co rozhoduje o přechodu",
        "Klient a.s.",
        "Důvěra 2026",
        "STU-1",
        date(2026, 9, 25),
        2,
        METHOD_STATUS_PENDING,
    )
    return replace(base, **overrides)


def _doc(ledger: Any, **meta: Any) -> ReportDocument:
    method = Callout(CalloutKind.METHOD, text("Výsledky jsou modelované."))
    return ReportDocument(
        _meta(**meta),
        (
            Section(
                "Shrnutí",
                (
                    method,
                    Paragraph(
                        (
                            Text("Hlavní zjištění "),
                            Text("je jasné", "strong"),
                            Footnote("Poznámka pod čarou."),
                            Text(", viz "),
                            CrossRef("ch-method"),
                            Text("."),
                        )
                    ),
                    Heading("Kontext", 2, id="h-context"),
                    Paragraph((Link("Zdroj dat", "https://example.org/data"),)),
                    Heading("Detail", 3),
                ),
            ),
            Section("Metodika", (Paragraph(text("Jak jsme postupovali.")),), id="ch-method"),
            Section(
                "Dotazník",
                (
                    Paragraph(
                        text("Otázky."),
                    ),
                ),
                appendix=True,
            ),
        ),
        ledger,
    )


def _parts(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return {n: z.read(n) for n in z.namelist()}


def _xml(parts: dict[str, bytes], name: str) -> etree._Element:
    return etree.fromstring(parts[name])


def _text(root: etree._Element) -> str:
    return "".join(t.text or "" for t in root.iter(f"{W}t"))


def _instr(root: etree._Element) -> list[str]:
    return [(t.text or "").strip() for t in root.iter(f"{W}instrText")]


@pytest.fixture
def rendered(report_ledger: Any) -> dict[str, bytes]:
    return _parts(DocxRenderer().render(_doc(report_ledger())))


def _sections(parts: dict[str, bytes]) -> list[etree._Element]:
    return list(_xml(parts, "word/document.xml").iter(f"{W}sectPr"))


def test_an_invalid_document_is_refused_before_anything_is_written(report_ledger: Any) -> None:
    doc = replace(_doc(report_ledger()), sections=())
    with pytest.raises(ReportInvalid):
        DocxRenderer().render(doc)


def test_the_package_opens_and_renders_identically_twice(report_ledger: Any) -> None:
    doc = _doc(report_ledger())
    first, second = DocxRenderer().render(doc), DocxRenderer().render(doc)
    assert first == second
    Document(io.BytesIO(first))  # python-docx reads it back


def test_cover_front_matter_and_body_are_three_sections_and_a_portrait_appendix_flows_on(
    rendered: dict[str, bytes],
) -> None:
    # Was four sections: a portrait appendix opened a section, and so a page, of its
    # own. Chapters and appendices now flow (research-integrity-report D5).
    cover, front, body = _sections(rendered)
    assert cover.find(f"{W}headerReference") is not None  # artwork only, no running text
    front_pg = front.find(f"{W}pgNumType")
    body_pg = body.find(f"{W}pgNumType")
    assert front_pg is not None and front_pg.get(f"{W}fmt") == "lowerRoman"
    assert front_pg.get(f"{W}start") == "1"
    assert body_pg is not None and body_pg.get(f"{W}fmt") == "decimal"
    assert body_pg.get(f"{W}start") == "1"
    for sect in (front, body):  # each opens a page: its numbering restarts
        start = sect.find(f"{W}type")
        assert start is not None and start.get(f"{W}val") == "nextPage"
    for sect in (cover, front, body):
        assert in_schema_order(sect, SECTPR_ORDER)
        size = sect.find(f"{W}pgSz")
        assert size is not None and (size.get(f"{W}w"), size.get(f"{W}h")) == ("11906", "16838")
    # Every section the document ends before a paragraph closes the cover or the
    # front matter: the chapters and the appendix after them share the body's.
    root = _xml(rendered, "word/document.xml")
    paragraphs = list(root.iter(f"{W}p"))
    ends = [i for i, p in enumerate(paragraphs) if p.find(f"{W}pPr/{W}sectPr") is not None]
    levels = {
        i: style.get(f"{W}val")
        for i, p in enumerate(paragraphs)
        if (style := p.find(f"{W}pPr/{W}pStyle")) is not None
        and style.get(f"{W}val") in ("Heading1", "AIAAppendixHeading")
    }
    assert [levels[i] for i in sorted(levels)] == ["Heading1", "Heading1", "AIAAppendixHeading"]
    assert len(ends) == 2 and max(ends) < min(levels)


def _header_footer_text(parts: dict[str, bytes], kind: str) -> list[tuple[str, list[str]]]:
    return [
        (_text(_xml(parts, n)), _instr(_xml(parts, n)))
        for n in sorted(parts)
        if n.startswith(f"word/{kind}") and n.endswith(".xml")
    ]


def test_running_heads_match_the_supplied_study_and_client_template(
    rendered: dict[str, bytes],
) -> None:
    headers = _header_footer_text(rendered, "header")
    running = [text for text, _ in headers if text]
    assert len(running) == 2  # front matter and body: the appendix flows on in the body
    assert all(text == "AIA · Důvěra 2026Klient a.s." for text in running)
    assert all("STYLEREF" not in " ".join(fields) for _, fields in headers)
    for n in rendered:
        if n.startswith("word/header") and n.endswith(".xml") and _text(_xml(rendered, n)):
            assert (
                _xml(rendered, "word/styles.xml")
                .find(f"{W}style[@{W}styleId='Header']/{W}pPr/{W}pBdr/{W}bottom")
                .get(f"{W}sz")
                == "6"
            )


def _status_runs(root: etree._Element) -> list[str]:
    return [
        "".join(t.text or "" for t in r.iter(f"{W}t"))
        for r in root.iter(f"{W}r")
        if (st := r.find(f"{W}rPr/{W}rStyle")) is not None and st.get(f"{W}val") == "AIAStatus"
    ]


def test_a_draft_says_its_status_only_in_its_footers(rendered: dict[str, bytes]) -> None:
    # Was: the cover notice, the cover's "Schválení: čeká na podpis" and the document control
    # said draft too, so an approved export could not differ from its version in the footers
    # alone (research-integrity-report D2). Now every footer part says it in one run.
    footers = _header_footer_text(rendered, "footer")
    assert len(footers) == 3  # cover, front matter, body (the appendix flows on in it)
    assert all(DRAFT in text for text, _ in footers)
    assert sum("PAGE" in instr for _, instr in footers) == 2
    assert all("Důvěrné" in text for text, instr in footers if "PAGE" in instr)
    runs = {
        n: _status_runs(_xml(rendered, n))
        for n in rendered
        if n.startswith("word/footer") and n.endswith(".xml")
    }
    assert all(len(r) == 1 for r in runs.values()), runs
    assert sorted(r for (r,) in runs.values()) == [
        DRAFT,
        DRAFT,
        f"{DRAFT}. Tato verze dosud nebyla schválena. Nepředávejte ji jako finální.",  # the cover
    ]
    body = _xml(rendered, "word/document.xml")
    assert _status_runs(body) == []
    assert DRAFT not in _text(body) and "schválena" not in _text(body)
    assert "Stav této verze uvádí zápatí; záznam o schválení vede AIA." in _text(body)


def _wrapped_lines(text: str, width_mm: float, size_pt: float, font: Any) -> int:
    """How many lines ``text`` takes in a column ``width_mm`` wide, broken greedily where a
    renderer may break it: at a space and after a hyphen."""
    pieces = [p for p in re.split(r"(?<=[ -])", text) if p]
    lines, line = 1, ""
    for piece in pieces:
        if line and font.getlength((line + piece).rstrip()) / 1000 * size_pt * 25.4 / 72 > width_mm:
            lines, line = lines + 1, piece
        else:
            line += piece
    return lines


def _printed_lines(paragraph: etree._Element) -> list[str]:
    """A paragraph's text as its line breaks part it."""
    lines = [""]
    for node in paragraph.iter(f"{W}t", f"{W}br"):
        if node.tag == f"{W}br":
            lines.append("")
        else:
            lines[-1] += node.text or ""
    return lines


def test_the_cover_status_wraps_short_of_the_cover_field(rendered: dict[str, bytes]) -> None:
    # Was: the cover's footer in the running-foot style, its long draft sentence (and, once
    # approved, "SCHVÁLENO … · verze ART-…") printed over the cover field's dots across the
    # full text width (review documents#3).
    from PIL import ImageFont

    from aia_core.domain.report.print_tokens import PAGE, STYLES
    from aia_core.domain.report.review import ReviewMark, ReviewState
    from aia_core.infrastructure.report_docx.cover import ASSET
    from aia_core.infrastructure.report_docx.renderer import write_package
    from aia_core.infrastructure.report_docx.review_stamp import stamp_review
    from aia_core.infrastructure.report_docx.styles import S

    cover_text = f"{DRAFT}. Tato verze dosud nebyla schválena. Nepředávejte ji jako finální."
    (name,) = [
        n
        for n in rendered
        if n.startswith("word/footer") and cover_text in _text(_xml(rendered, n))
    ]
    (paragraph,) = _xml(rendered, name).iter(f"{W}p")
    assert paragraph.find(f"{W}pPr/{W}pStyle").get(f"{W}val") == "AIACoverFooter"
    assert _status_runs(paragraph) == [cover_text]  # still the one run a stamp replaces

    style = _xml(rendered, "word/styles.xml").find(f"{W}style[@{W}styleId='AIACoverFooter']")
    assert style.find(f"{W}name").get(f"{W}val") == S.COVER_FOOTER
    right_indent_mm = int(style.find(f"{W}pPr/{W}ind").get(f"{W}right")) / 1440 * 25.4
    cover = _sections(rendered)[0].find(f"{W}pgMar")
    left_mm = int(cover.get(f"{W}left")) / 1440 * 25.4
    right_mm = int(cover.get(f"{W}right")) / 1440 * 25.4
    width_mm = PAGE.width_mm - left_mm - right_mm - right_indent_mm
    text_ends_mm = left_mm + width_mm

    # The field's dots in the band the footer can occupy (its distance plus four lines).
    field = etree.parse(str(ASSET)).getroot()
    _, _, vb_w, vb_h = (float(v) for v in field.get("viewBox").split())
    band_top_mm = PAGE.height_mm - PAGE.footer_mm - 4 * STYLES["doc-running"].leading_pt * 25.4 / 72
    dots = [
        (float(c.get("cx")) - float(c.get("r"))) * PAGE.width_mm / vb_w
        for c in field
        if etree.QName(c).localname == "circle"
        and float(c.get("cy")) * PAGE.height_mm / vb_h >= band_top_mm
    ]
    assert dots and text_ends_mm + 3 <= min(dots), (text_ends_mm, min(dots))

    # The draft's sentence wraps to two lines (measured in the status face), and the cover
    # leaves those two lines room above its bottom edge.
    face = STYLES["doc-table-head"]
    font = ImageFont.truetype(
        str(ASSET.parent.parent / "fonts" / "IBMPlexSans-SemiBold.ttf"), size=1000
    )
    assert _wrapped_lines(cover_text, width_mm, face.size_pt, font) == 2, cover_text
    placement = next(_xml(rendered, "word/document.xml").iter(f"{W}tbl"))
    row_mm = int(placement.find(f"{W}tr/{W}trPr/{W}trHeight").get(f"{W}val")) / 1440 * 25.4
    top_mm = int(cover.get(f"{W}top")) / 1440 * 25.4
    # the footer's distance, its rule (6 pt above it, 0.75 pt) and two lines; the section's
    # closing 1 pt paragraph under the placement table
    footer_mm = PAGE.footer_mm + (6 + 0.75 + 2 * STYLES["doc-running"].leading_pt) * 25.4 / 72
    assert top_mm + row_mm + 25.4 / 72 + footer_mm <= PAGE.height_mm - 3, row_mm

    # An approved cover prints its stamp as two lines, each whole. Was: the stamp's one line,
    # wrapped by the column after "verze ART-", so the id read from page 1 lost its hyphen
    # ("ART0123…", an id that does not exist; cross-review). The widest id the repository
    # issues ("ART-" and 16 hex characters), on a date with two-digit day and month.
    widest = max("0123456789abcdef", key=font.getlength)
    stored = write_package(rendered.items())
    mark = ReviewMark(
        ReviewState.APPROVED,
        f"ART-{widest * 16}",
        hashlib.sha256(stored).hexdigest(),
        date(2026, 12, 28),
    )
    (stamped,) = _xml(_parts(stamp_review(stored, mark)), name).iter(f"{W}p")
    lines = _printed_lines(stamped)
    nb = "\u00a0"
    assert lines == [f"SCHVÁLENO 28.{nb}12.{nb}2026", f"verze {mark.artifact_id}"]
    for line in lines:
        assert _wrapped_lines(line, width_mm, face.size_pt, font) == 1, line


def test_chapters_and_appendices_flow_and_only_furniture_breaks_the_page(
    rendered: dict[str, bytes],
) -> None:
    root = _xml(rendered, "word/document.xml")
    assert not [br for br in root.iter(f"{W}br") if br.get(f"{W}type") == "page"]
    _cover, *rest = _sections(rendered)
    assert len(rest) == 2  # the front matter and the body, each after a change of furniture
    assert all(s.find(f"{W}type").get(f"{W}val") == "nextPage" for s in rest)
    # The one style that breaks the page is the contents' heading, used once.
    styles = _xml(rendered, "word/styles.xml")
    breaking = {
        s.get(f"{W}styleId")
        for s in styles.iter(f"{W}style")
        if (b := s.find(f"{W}pPr/{W}pageBreakBefore")) is not None
        and b.get(f"{W}val") not in ("0", "false", "off")
    }
    assert breaking == {"AIAFrontHeadingPage"}
    opened = [
        _text(p)
        for p in root.iter(f"{W}p")
        if (st := p.find(f"{W}pPr/{W}pStyle")) is not None and st.get(f"{W}val") in breaking
    ]
    assert opened == ["Obsah"]


def test_the_contents_lists_the_outline_with_a_pageref_per_entry_and_no_toc_field(
    rendered: dict[str, bytes],
) -> None:
    root = _xml(rendered, "word/document.xml")
    # No TOC field: headless LibreOffice leaves its cached page numbers empty
    # (research-integrity-report F3); each entry carries its own PAGEREF instead.
    assert not any(i.startswith("TOC") for i in _instr(root))
    entries = [
        _text(p)
        for p in root.iter(f"{W}p")
        if (s := p.find(f"{W}pPr/{W}pStyle")) is not None and s.get(f"{W}val").startswith("toc")
    ]
    assert entries == [
        "1 Shrnutí",
        "1.1 Kontext",
        "1.1.1 Detail",
        "2 Metodika",
        "Příloha A Dotazník",
    ]
    anchors = {b.get(f"{W}name") for b in root.iter(f"{W}bookmarkStart")}
    for link in root.iter(f"{W}hyperlink"):
        if link.get(f"{W}anchor"):
            assert link.get(f"{W}anchor") in anchors  # every entry and cross-ref resolves
    for p in root.iter(f"{W}p"):
        style = p.find(f"{W}pPr/{W}pStyle")
        if style is not None and style.get(f"{W}val").startswith("toc"):
            refs = [i.split()[1] for i in _instr(p) if i.startswith("PAGEREF")]
            assert len(refs) == 1 and refs[0] in anchors  # numbered in Word and in the PDF


def test_headings_print_their_number_and_cross_references_their_label(
    rendered: dict[str, bytes],
) -> None:
    root = _xml(rendered, "word/document.xml")
    styled = {
        _text(p): p.find(f"{W}pPr/{W}pStyle").get(f"{W}val")
        for p in root.iter(f"{W}p")
        if p.find(f"{W}pPr/{W}pStyle") is not None
    }
    assert styled["1 Shrnutí"] == "Heading1"
    assert styled["1.1 Kontext"] == "Heading2"
    assert styled["Příloha A Dotazník"] == "AIAAppendixHeading"
    assert "viz kapitola 2." in _text(root)


def test_document_control_with_no_identifiers_in_a_client_report(
    rendered: dict[str, bytes],
) -> None:
    body = _text(_xml(rendered, "word/document.xml"))
    assert "Řízení dokumentu" in body
    assert "STU-1" not in body  # the study id is an identifier: internal only
    assert "Syntetický a modelovaný výzkum" in body  # the method status, in words


def test_the_revision_a_report_prints_is_named_as_its_designs(
    rendered: dict[str, bytes],
) -> None:
    # Review 4 of D2: an approved file printed "Revize 1" on its cover and in its document
    # control, and "verze ART-…" in its footers, two numbers either of which read as the
    # version. The first is the Design Revision the report was composed from, and says so.
    body = _text(_xml(rendered, "word/document.xml"))
    assert body.count("Revize návrhu") == 2  # the cover and the document control


def test_an_internal_report_prints_its_identifiers(report_ledger: Any) -> None:
    from aia_core.domain.evidence.claims import ClaimSurface

    doc = _doc(
        report_ledger(ClaimSurface.INTERNAL),
        kind=ReportKind.INTERNAL,
        classification=Classification.INTERNAL,
        identifiers=(("run", "run-7f3a"),),
    )
    body = _text(_xml(_parts(DocxRenderer().render(doc)), "word/document.xml"))
    assert "run-7f3a" in body and "STU-1" in body
    assert "Interní \u2014 nepředávat klientovi" in body


def test_footnotes_links_properties_fonts_and_settings(rendered: dict[str, bytes]) -> None:
    notes = _xml(rendered, "word/footnotes.xml")
    assert "Poznámka pod čarou." in _text(notes)
    rels = rendered["word/_rels/document.xml.rels"].decode()
    assert "https://example.org/data" in rels and 'TargetMode="External"' in rels
    core = rendered["docProps/core.xml"].decode()
    assert "<dc:language>cs-CZ</dc:language>" in core
    assert "Důvěra v banky 2026" in core
    assert sum(n.endswith(".odttf") for n in rendered) == 12
    settings = _xml(rendered, "word/settings.xml")
    assert in_schema_order(settings, SETTINGS_ORDER)
    for tag in ("embedTrueTypeFonts", "mirrorMargins", "updateFields", "footnotePr"):
        assert settings.find(f"{W}{tag}") is not None, tag
    assert settings.find(f"{W}rsids") is None


def test_nothing_is_formatted_directly(report_ledger: Any) -> None:
    assert lint_docx(DocxRenderer().render(_doc(report_ledger()))) == []


def test_the_lint_finds_direct_formatting() -> None:
    document = Document()
    document.add_paragraph().add_run("tučně").bold = True
    out = io.BytesIO()
    document.save(out)
    problems = lint_docx(out.getvalue())
    assert any("<b>" in p and "tučně" in p for p in problems), problems


def test_a_method_status_with_no_printed_wording_is_refused(report_ledger: Any) -> None:
    with pytest.raises(ReportInvalid, match="no printed wording"):
        DocxRenderer().render(_doc(report_ledger(), method_status="validated, trust me"))


def test_the_cover_names_the_study_and_anchors_the_original_field(report_ledger: Any) -> None:
    internal = _doc(
        report_ledger(ClaimSurface.INTERNAL),
        kind=ReportKind.INTERNAL,
        classification=Classification.INTERNAL,
    )
    parts = _parts(DocxRenderer().render(internal))
    body = _xml(parts, "word/document.xml")
    wp = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
    headers = [
        etree.fromstring(v)
        for k, v in parts.items()
        if k.startswith("word/header") and k.endswith(".xml")
    ]
    anchors = [a for h in headers for a in h.iter(wp + "anchor")]
    assert len(anchors) == 2
    assert {a.get("behindDoc") for a in anchors} == {"0", "1"}
    assert any(a.find(wp + "docPr").get("name") == "AIA report lockup" for a in anchors)
    cover = Document(io.BytesIO(DocxRenderer().render(internal)))
    cell = cover.tables[0].cell(0, 0)
    assert (
        next(p.text for p in cell.paragraphs if p.style.name == "Title") == internal.meta.study_name
    )
    assert "čeká na podpis" not in _text(body)  # the status is the footers' (D2)
    assert cover.tables[0].columns[0].width.mm == pytest.approx(120, abs=0.02)
    from aia_core.infrastructure.report_docx.cover import ASSETS

    assert (ASSETS / "aia-lockup.svg").read_bytes() in parts.values()
    # Header SVG relationships must resolve in their own part, not document.xml.
    ns = "{http://schemas.microsoft.com/office/drawing/2016/SVG/main}"
    rel_ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    r = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    for name, data in parts.items():
        if (
            name.startswith("word/header")
            and name.endswith(".xml")
            and list(etree.fromstring(data).iter(ns + "svgBlip"))
        ):
            rels = etree.fromstring(parts["word/_rels/" + name.split("/")[-1] + ".rels"])
            targets = {e.get("Id"): e.get("Target") for e in rels.iter(rel_ns + "Relationship")}
            for image in etree.fromstring(data).iter(ns + "svgBlip"):
                assert "word/" + targets[image.get(r + "embed")] in parts


def test_the_evidence_appendix_is_wide_and_audit_returns_to_portrait(report_ledger: Any) -> None:
    from aia_core.domain.report.model import AuditBlock, EvidenceAppendix

    sample = _doc(report_ledger())
    document = replace(
        sample,
        sections=(
            *sample.sections[:2],
            Section("Evidence", (EvidenceAppendix(),), appendix=True),
            Section("Audit", (AuditBlock((("Run", "FIC-1"),)),), appendix=True),
        ),
        meta=replace(sample.meta, kind=ReportKind.INTERNAL, classification=Classification.INTERNAL),
        ledger=report_ledger(ClaimSurface.INTERNAL),
    )
    parts = _parts(DocxRenderer().render(document))
    sections = _sections(parts)
    assert sections[-2].find(W + "pgSz").get(W + "orient") == "landscape"
    assert sections[-1].find(W + "pgSz").get(W + "orient") != "landscape"
