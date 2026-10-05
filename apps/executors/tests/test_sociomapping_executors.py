"""The experimental Sociomapping and its report under the real worker loop (plan I1, I2).

A run started with the switch gains two steps after fieldwork; one started without it is
the research graph as before. On the workbench's fictional fieldwork the whole run
completes: the result is stored EXPERIMENTAL_AIA and never client-facing, and the report is
an INTERNAL, unapproved DOCX that says what the method is and is not.
"""

from __future__ import annotations

import io
from collections.abc import Callable
from typing import Any

import pytest
from aia_core.application.research import ResearchRuns, research_artifacts
from aia_core.application.sociomapping_report import SOCIOMAPPING_REPORT_ARTIFACT_TYPE
from aia_core.domain.fieldwork import FieldworkSource
from aia_core.domain.workflow import StepRunStatus, WorkflowRunStatus
from aia_core.infrastructure.build_identity import BuildIdentity
from aia_core.infrastructure.report_docx.renderer import DocxRenderer
from aia_core.infrastructure.storage import InMemoryArtifactStore
from aia_core.infrastructure.study_design_repository import StudyDesignRepository
from aia_executors import workbench
from aia_executors.registry import registry_for
from aia_executors.research import SOCIOMAPPING
from aia_worker.executor import StepExecutor
from aia_worker.settings import WorkerSettings
from aia_worker.worker import Worker
from docx import Document
from sqlalchemy.orm import Session, sessionmaker

DESIGN = {
    "title": "Ranní nápoj",
    "n": 60,
    "sections": [
        {
            "type": "object_battery",
            "title": "Nápoje",
            "object_family": "nápoje",
            "objects": ["Káva", "Čaj", "Kakao", "Džus", "Voda"],
            "object_question": "Jak hodnotíte {object}?",
            "scale": [1, 10],
        },
    ],
}


@pytest.fixture
def worker_with(
    sessions: sessionmaker[Session], database_url: str
) -> Callable[[dict[str, StepExecutor]], Worker]:
    def build(executors: dict[str, StepExecutor]) -> Worker:
        return Worker(
            session_factory=sessions,
            executors=executors,
            settings=WorkerSettings(
                database_url=database_url,
                executors="aia_executors.registry:build_registry",
                worker_id="sociomapping-worker",
                lease_seconds=60,
                heartbeat_seconds=0.1,
                poll_seconds=0.05,
                maintenance_seconds=0.2,
            ),
        )

    return build


def _start(world: Any, *, sociomapping: bool) -> str:
    with world.sessions() as session:
        scope = world.lead_scope(session)
        revision, _ = StudyDesignRepository(session, scope).submit(
            content=DESIGN, source_stage="run"
        )
        started = ResearchRuns(session, scope).start(
            design_revision_id=revision.revision_id,
            fieldwork_source=FieldworkSource.SYNTHETIC_FIXTURE,
            sociomapping_enabled=sociomapping,
        )
        session.commit()
        return started.run_id


def _drain(worker: Worker) -> list[str]:
    endings: list[str] = []
    while (result := worker.run_once()) is not None:
        endings.append(result.ending)
    return endings


def _run(world: Any, run_id: str) -> dict[str, Any]:
    with world.sessions() as session:
        return ResearchRuns(session, world.lead_scope(session)).get(run_id)


def test_without_the_switch_the_graph_is_unchanged(world: Any) -> None:
    run = _run(world, _start(world, sociomapping=False))
    assert [s["node_key"] for s in run["steps"]] == [
        "compile",
        "preflight",
        "run",
        "aggregate",
        "sociomap",
    ]
    assert "sociomapping_enabled" not in run["metadata"]


def test_the_fictional_run_maps_reports_and_stays_experimental(
    world: Any,
    store: InMemoryArtifactStore,
    build: BuildIdentity,
    worker_with: Callable[[dict[str, StepExecutor]], Worker],
) -> None:
    run_id = _start(world, sociomapping=True)
    run = _run(world, run_id)
    assert run["metadata"]["sociomapping_enabled"] is True
    assert [s["node_key"] for s in run["steps"]][-2:] == ["sociomapping", "sociomapping_report"]

    worker = worker_with(workbench.workbench_registry_for(store=store, build=build))
    assert _drain(worker) == ["completed"] * 7
    run = _run(world, run_id)
    assert run["status"] is WorkflowRunStatus.COMPLETED
    steps = {s["node_key"]: s for s in run["steps"]}

    out = steps["sociomapping"]["output"]
    assert steps["sociomapping"]["status"] is StepRunStatus.SUCCEEDED
    assert out["artifact_type"] == SOCIOMAPPING
    assert out["method_status"] == "EXPERIMENTAL_AIA"
    assert out["data_origin"] == "SYNTHETIC_FIXTURE"
    with world.sessions() as session:
        research = research_artifacts(session, world.lead_scope(session), store)
        artifact = research.get(out["artifact_id"])
        result = research.read_json(out["artifact_id"])["sociomapping"]
        deps = {a.artifact_id for a in research.dependencies(out["artifact_id"])}
    assert artifact.metadata["client_facing"] is False
    assert artifact.metadata["method_status"] == "EXPERIMENTAL_AIA"
    assert deps == {
        steps["compile"]["output"]["artifact_id"],
        steps["run"]["output"]["artifact_id"],
    }
    assert result["client_facing"] is False and result["synthetic_data"] is True
    assert result["inputs"]["dataset_artifact_id"] == steps["run"]["output"]["artifact_id"]
    (battery,) = result["batteries"]
    assert battery["status"] == "MAPPED"
    assert sorted(battery["layout"]["element_ids"]) == ["caj", "dzus", "kakao", "kava", "voda"]
    assert battery["layout"]["accuracy"]["overall"] is not None

    report_out = steps["sociomapping_report"]["output"]
    assert report_out["artifact_type"] == SOCIOMAPPING_REPORT_ARTIFACT_TYPE
    assert report_out["review_state"] == "DRAFT_UNAPPROVED"
    with world.sessions() as session:
        research = research_artifacts(session, world.lead_scope(session), store)
        report = research.get(report_out["artifact_id"])
        data = research.read(report_out["artifact_id"])
        report_deps = {a.artifact_id for a in research.dependencies(report_out["artifact_id"])}
    assert report.content_type == DocxRenderer.media_type
    assert report.metadata["client_facing"] is False and not report.is_approved
    assert report_deps == {out["artifact_id"]}
    document = Document(io.BytesIO(data))
    body = "\n".join(p.text for p in document.paragraphs)
    cells = "\n".join(c.text for t in document.tables for row in t.rows for c in row.cells)
    assert "experimentálního H-Modelu AIA, nikoli ověřené rekonstrukce SOMECS" in body
    assert "Fiktivní data" in body or "fiktivní" in body
    assert "aia_hmodel_candidate_v1" in body + cells
    assert "Celková přesnost (Spearman)" in cells
    assert "Původ výpočtu" in body + cells
    assert len(document.inline_shapes) >= 1  # the map


def test_production_parks_at_fieldwork_and_maps_nothing(
    world: Any,
    store: InMemoryArtifactStore,
    build: BuildIdentity,
    worker_with: Callable[[dict[str, StepExecutor]], Worker],
) -> None:
    run_id = _start(world, sociomapping=True)
    _drain(worker_with(registry_for(store=store, build=build)))
    steps = {s["node_key"]: s for s in _run(world, run_id)["steps"]}
    assert steps["sociomapping"]["status"] is not StepRunStatus.SUCCEEDED
    assert steps["sociomapping"].get("output") in (None, {})
    assert steps["sociomapping_report"]["status"] is not StepRunStatus.SUCCEEDED
