"""Zpráva in the report editor: the signed config and the Document Server's fetch of a version
(artifact-report RPT8c, part 2)."""

from __future__ import annotations

import hashlib
import io
import json
import os
import time
import zipfile
from datetime import datetime
from typing import Any
from unittest.mock import ANY
from urllib.parse import parse_qs, urlparse

import jwt
import pytest
from aia_core.application.report import (
    INTERNAL_REPORT_ARTIFACT_TYPE,
    INTERNAL_REPORT_MEDIA_TYPE,
    INTERNAL_REPORT_PDF_ARTIFACT_TYPE,
)
from aia_core.application.research import research_artifacts
from aia_core.application.scope import AuthenticatedPrincipal, ScopeResolver
from aia_core.domain.report.saved_changes import read_changes_payload
from aia_core.infrastructure.db import create_session_factory
from aia_core.infrastructure.report_editor_repository import ReportEditorRepository
from aia_core.infrastructure.workflow_repository import WorkflowRepository
from fastapi import FastAPI
from fastapi.testclient import TestClient

from aia_api.config import Environment, Settings
from aia_api.editor_tokens import EditorTokens


def _now() -> Any:
    from datetime import UTC, datetime

    return datetime.now(UTC)


API = "/api/v1"
EDITOR = "https://docs.aia.example.test"
PUBLIC = "https://aia.example.test"
SECRET = "k" * 48
#: A fictional api.js integrity: one SHA-384 digest as the attribute takes it (RPT8e).
INTEGRITY = "sha384-" + "a" * 64
DESIGN = {
    "title": "Ranní nápoj",
    "goal": "Zjistit, zda nový nápoj dává smysl dojíždějícím.",
    "n": 450,
    "sections": [
        {
            "type": "questions",
            "questions": [
                {"id": "q1", "text": "Jak často pijete kávu?", "typ": "skala", "skala": [1, 5]},
            ],
        }
    ],
}


@pytest.fixture
def settings() -> Settings:
    return Settings(
        env=Environment.TEST,
        database_url=os.environ.get("DATABASE_URL", "sqlite+pysqlite:///:memory:"),
        identity_provider="test",
        log_level="WARNING",
        log_format="console",
        onlyoffice_url=EDITOR,
        onlyoffice_jwt_secret=SECRET,
        onlyoffice_callback_base=PUBLIC,
        onlyoffice_api_integrity=INTEGRITY,
        # A run has a report step only with its analysis (no model is called here).
        ai_analysis_enabled=True,
    )


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _docx(value: str = "42,5 %", prose: str = "Bance důvěřuje") -> bytes:
    """A report body with one locked number, ``value`` under ``trust_total``."""
    body = (
        f'<w:document xmlns:w="{W}"><w:body><w:p><w:r><w:t>{prose} </w:t></w:r>'
        '<w:sdt><w:sdtPr><w:tag w:val="trust_total"/></w:sdtPr><w:sdtContent>'
        f"<w:r><w:t>{value}</w:t></w:r></w:sdtContent></w:sdt></w:p></w:body></w:document>"
    )
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        # A fixed time: the same content is the same bytes whenever it is built.
        z.writestr(zipfile.ZipInfo("word/document.xml", (2026, 10, 23, 0, 0, 0)), body)
    return out.getvalue()


def _reported_run(
    app: FastAPI,
    researcher: TestClient,
    world: Any,
    *,
    title: str = "Ranní nápoj",
    with_pdf: bool = False,
) -> tuple[str, str]:
    """A run of the primary study whose report step stored an internal draft: (run, draft);
    another ``title`` is another design, so another run. ``with_pdf``: the step rendered its
    PDF too, as it does with the converter on."""
    revision = world.design(researcher, {**DESIGN, "title": title}).json()["revision_id"]
    runs = f"{API}/studies/{world.study_id()}/research/runs"
    started = researcher.post(runs, json={"design_revision_id": revision})
    assert started.status_code == 201, started.text
    run_id = started.json()["run_id"]
    factory = create_session_factory(app.state.engine)
    with factory() as session:
        principal = AuthenticatedPrincipal(
            user_id=world.users["researcher"], organization_id=world.organization_id
        )
        scope = ScopeResolver(session).study_context(principal, study_id=world.study_id())
        workflow = WorkflowRepository(session, scope)
        draft = ""
        while (work := workflow.claim_next(worker_id="editor-fixture")) is not None:
            output: dict[str, Any] = {}
            if work.node_key == "report":
                artifact, _ = research_artifacts(session, scope, app.state.artifact_store).put(
                    project_id=work.project_id,
                    revision=work.project_revision,
                    stage_type="REPORT",
                    artifact_type=INTERNAL_REPORT_ARTIFACT_TYPE,
                    data=_docx(),
                    content_type=INTERNAL_REPORT_MEDIA_TYPE,
                    metadata={"run_id": run_id, "report_kind": "internal"},
                    reuse=False,
                )
                draft = artifact.artifact_id
                output = {"artifact_id": draft}
                if with_pdf:
                    pdf, _ = research_artifacts(session, scope, app.state.artifact_store).put(
                        project_id=work.project_id,
                        revision=work.project_revision,
                        stage_type="REPORT",
                        artifact_type=INTERNAL_REPORT_PDF_ARTIFACT_TYPE,
                        data=b"%PDF-1.7 the step's rendering",
                        content_type="application/pdf",
                        metadata={"run_id": run_id, "rendering_of": draft},
                        reuse=False,
                    )
                    output["pdf"] = {"state": "READY", "artifact_id": pdf.artifact_id}
            workflow.complete_attempt(work.attempt_id, worker_id="editor-fixture", output=output)
        session.commit()
    assert draft
    return run_id, draft


def _editor(world: Any, run_id: str) -> str:
    return f"{API}/studies/{world.study_id()}/research/runs/{run_id}/report/editor"


def _fetch(client: TestClient, url: str, *, signed: bool = True) -> Any:
    """``url`` fetched as the Document Server does: with its JWT unless ``signed`` is false."""
    parsed = urlparse(url)
    assert f"{parsed.scheme}://{parsed.netloc}" == PUBLIC
    headers = {"Authorization": f"Bearer {jwt.encode({'url': url}, SECRET)}"} if signed else {}
    return client.get(parsed.path, params=parse_qs(parsed.query), headers=headers)


def test_the_config_is_signed_names_no_e_mail_and_fetches_its_version(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any
) -> None:
    run_id, draft = _reported_run(app, researcher, world)
    response = researcher.post(_editor(world, run_id), json={"variant": "internal"})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["document_server"] == EDITOR
    assert body["api_integrity"] == INTEGRITY
    assert body["line_id"] == draft and body["koncept"] == 1
    config = body["config"]
    signed = jwt.decode(config.pop("token"), SECRET, algorithms=["HS256"])
    assert isinstance(signed.pop("exp"), int)  # a config's token expires
    assert signed == config  # the Document Server checks the config as signed
    user = config["editorConfig"]["user"]
    assert user["id"] == world.users["researcher"] and "@" not in user["name"]
    assert "@" not in response.text  # no e-mail anywhere in what the browser gets
    assert config["editorConfig"]["customization"]["forcesave"] is True
    # The Document Server's routes, outside any study: the token's session names it.
    assert config["editorConfig"]["callbackUrl"].startswith(f"{PUBLIC}{API}/report-editor/")
    assert config["document"]["url"].startswith(f"{PUBLIC}{API}/report-editor/document?")
    fetched = _fetch(client, config["document"]["url"])
    assert fetched.status_code == 200 and fetched.content == _docx()


def test_a_second_researcher_gets_the_same_key(
    app: FastAPI, researcher: TestClient, lead: TestClient, world: Any
) -> None:
    run_id, _ = _reported_run(app, researcher, world)
    first = researcher.post(_editor(world, run_id), json={}).json()["config"]
    second = lead.post(_editor(world, run_id), json={}).json()["config"]
    assert first["document"]["key"] == second["document"]["key"]
    assert second["editorConfig"]["user"]["id"] == world.users["lead"]


def _signed_for(path: str, query: dict[str, Any]) -> dict[str, str]:
    """The Document Server's JWT for a fetch of ``path`` with ``query``."""
    token = query["token"][0] if isinstance(query["token"], list) else query["token"]
    url = f"{PUBLIC}{path}?token={token}"
    return {"Authorization": f"Bearer {jwt.encode({'url': url}, SECRET)}"}


def test_the_document_is_served_only_to_the_document_server_for_its_url(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any
) -> None:
    run_id, _ = _reported_run(app, researcher, world)
    body = researcher.post(_editor(world, run_id), json={}).json()
    config, url = body["config"], body["config"]["document"]["url"]
    parsed = urlparse(url)
    query = parse_qs(parsed.query)
    assert _fetch(client, url, signed=False).status_code == 403  # no Document Server JWT
    # The browser's own JWT is under the same secret: a config's token is not the Server's.
    as_browser = {"Authorization": f"Bearer {config['token']}"}
    assert client.get(parsed.path, params=query, headers=as_browser).status_code == 403
    # A Document Server JWT for another URL.
    elsewhere = {"Authorization": f"Bearer {jwt.encode({'url': f'{PUBLIC}/x'}, SECRET)}"}
    assert client.get(parsed.path, params=query, headers=elsewhere).status_code == 403
    # The editor's own secret does not sign AIA's URL tokens: they are under a derived key.
    forged = {"token": jwt.encode({"kind": "document", "key": "x", "exp": 9e9}, SECRET)}
    assert (
        client.get(parsed.path, params=forged, headers=_signed_for(parsed.path, forged))
    ).status_code == 403
    # A key no session has: nothing to serve.
    stray = EditorTokens(SECRET).url_token(
        "document", document_key="ART-0.x.1", artifact_id=body["version_artifact_id"], now=_now()
    )
    query = {"token": stray}
    assert (
        client.get(parsed.path, params=query, headers=_signed_for(parsed.path, query)).status_code
        == 404
    )


def test_an_expired_a_callbacks_or_another_lines_token_does_not_fetch(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any
) -> None:
    from datetime import UTC, datetime, timedelta

    from aia_api.editor_tokens import DOCUMENT_LIFETIME, EditorTokens

    run_id, _ = _reported_run(app, researcher, world)
    body = researcher.post(_editor(world, run_id), json={}).json()
    key = body["config"]["document"]["key"]
    path = urlparse(body["config"]["document"]["url"]).path
    tokens = EditorTokens(SECRET)
    past = datetime.now(UTC) - DOCUMENT_LIFETIME - timedelta(seconds=1)
    expired = tokens.url_token(
        "document", document_key=key, artifact_id=body["version_artifact_id"], now=past
    )
    callback = tokens.url_token("callback", document_key=key, now=datetime.now(UTC))
    for refused in (expired, callback):
        query = {"token": refused}
        assert client.get(path, params=query, headers=_signed_for(path, query)).status_code == 403
    # A version of another line: the run's second draft, which this session's line is not.
    other_run, other_draft = _reported_run(app, researcher, world, title="Večerní nápoj")
    researcher.post(_editor(world, other_run), json={})
    stray = tokens.url_token(
        "document", document_key=key, artifact_id=other_draft, now=datetime.now(UTC)
    )
    query = {"token": stray}
    assert client.get(path, params=query, headers=_signed_for(path, query)).status_code == 404


def test_the_client_report_is_refused_by_name(
    app: FastAPI, researcher: TestClient, world: Any
) -> None:
    run_id, _ = _reported_run(app, researcher, world)
    refused = researcher.post(_editor(world, run_id), json={"variant": "client"})
    assert refused.status_code == 409
    assert refused.json()["code"] == "client_report_not_editable"


def test_with_the_editor_off_every_route_is_404(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any
) -> None:
    run_id, _ = _reported_run(app, researcher, world)
    url = researcher.post(_editor(world, run_id), json={}).json()["config"]["document"]["url"]
    app.state.settings = app.state.settings.model_copy(update={"onlyoffice_callback_base": ""})
    assert researcher.post(_editor(world, run_id), json={}).status_code == 404
    assert _fetch(client, url).status_code == 404
    # Without a token too: the editor's state answers first, never a 422 (MVP-ACCEPT AC01's
    # sweep of every study route as another organization).
    assert client.get(urlparse(url).path).status_code == 404


def test_a_fetch_without_a_token_is_refused_not_invalid(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any
) -> None:
    run_id, _ = _reported_run(app, researcher, world)
    url = researcher.post(_editor(world, run_id), json={}).json()["config"]["document"]["url"]
    path = urlparse(url).path
    assert client.get(path).status_code == 403
    assert client.get(path, params={"token": "x" * 3000}).status_code == 403
    app.state.settings = app.state.settings.model_copy(update={"onlyoffice_callback_base": ""})
    assert client.get(path, params={"token": "x" * 3000}).status_code == 404  # off: never 422


@pytest.mark.parametrize(
    "base",
    [
        "http://aia.example.test",
        "https://aia.example.test/api",
        "https://u@aia.example.test",
        "https://aia.example.test:8443",
    ],
)
def test_the_public_origin_is_an_https_origin_and_nothing_else(base: str) -> None:
    with pytest.raises(ValueError):
        Settings(env=Environment.TEST, onlyoffice_callback_base=base)


def test_the_editor_is_on_only_with_its_origins_its_secret_and_its_integrity() -> None:
    on = Settings(
        env=Environment.TEST,
        onlyoffice_url=EDITOR,
        onlyoffice_jwt_secret=SECRET,
        onlyoffice_callback_base=PUBLIC,
        onlyoffice_api_integrity=INTEGRITY,
    )
    assert on.report_editor_on
    for missing in ("onlyoffice_url", "onlyoffice_callback_base", "onlyoffice_api_integrity"):
        assert not on.model_copy(update={missing: ""}).report_editor_on


@pytest.mark.parametrize(
    "value",
    [
        "sha256-" + "A" * 43 + "=",  # another digest
        "sha384-" + "A" * 63,  # one character short
        "sha384-" + "A" * 64 + " sha384-" + "B" * 64,  # two hashes: the page takes exactly one
        "A" * 64,  # no prefix
    ],
)
def test_the_api_integrity_is_one_sha384_digest_or_nothing(value: str) -> None:
    """RPT8e: the integrity the page's api.js must match is refused at start unless it is one
    SHA-384 digest as an integrity attribute takes it."""
    with pytest.raises(ValueError):
        Settings(env=Environment.TEST, onlyoffice_api_integrity=value)
    assert Settings(
        env=Environment.TEST, onlyoffice_api_integrity=INTEGRITY
    ).onlyoffice_api_integrity


# -- the Document Server's routes are outside the study tree: their scope, tested here ---------
# (MVP-ACCEPT AC01 sweeps only routes naming {study_id}; these are scoped by the editor session the
# signed URL token names, resolved through ScopeResolver for the person who opened it.)


def _document_url(researcher: TestClient, world: Any, run_id: str) -> tuple[str, dict[str, Any]]:
    body = researcher.post(_editor(world, run_id), json={}).json()
    return body["config"]["document"]["url"], body


def test_a_removed_opener_s_session_serves_nothing(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any
) -> None:
    from datetime import UTC, datetime

    from aia_core.infrastructure.scope_repository import ScopeRepository

    run_id, _ = _reported_run(app, researcher, world)
    url, _body = _document_url(researcher, world, run_id)
    assert _fetch(client, url).status_code == 200
    factory = create_session_factory(app.state.engine)
    with factory() as session:
        owner = AuthenticatedPrincipal(
            user_id=world.users["owner"], organization_id=world.organization_id
        )
        admin = ScopeResolver(session).organization_context(owner)
        ScopeRepository(session).deactivate_member(
            admin, user_id=world.users["researcher"], now=datetime.now(UTC)
        )
        session.commit()
    removed = _fetch(client, url)
    assert removed.status_code == 404 and removed.json()["code"] == "editor_session_not_found"


def test_a_token_pairing_a_session_with_another_organizations_version_serves_nothing(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any
) -> None:
    """A URL token can only be minted under the editor's secret, but even so: a session key of
    this organization naming another organization's artifact finds no version of its line."""
    from aia_core.infrastructure.scope_repository import ScopeRepository
    from aia_core.infrastructure.study_design_repository import StudyDesignRepository

    run_id, _ = _reported_run(app, researcher, world)
    url, body = _document_url(researcher, world, run_id)
    factory = create_session_factory(app.state.engine)
    with factory() as session:
        repo, resolver = ScopeRepository(session), ScopeResolver(session)
        org, owner = repo.create_organization(
            slug="other-org", name="Other", owner_email="o@other.invalid", owner_name="O"
        )
        principal = AuthenticatedPrincipal(
            user_id=owner.user_id, organization_id=org.organization_id
        )
        admin = resolver.organization_context(principal)
        foreign_client = repo.create_client(admin, slug="c", name="C")
        study = repo.create_study(admin, client_id=foreign_client.client_id, slug="s", name="S")
        scope = resolver.study_context(principal, study_id=study.study_id)
        designs = StudyDesignRepository(session, scope)
        revision, _ = designs.submit_as_working_copy(content=DESIGN, source_stage="run")
        foreign, _ = research_artifacts(session, scope, app.state.artifact_store).put(
            project_id=str(designs.project_id()),
            revision=revision.revision,
            stage_type="REPORT",
            artifact_type=INTERNAL_REPORT_ARTIFACT_TYPE,
            data=_docx("Cizí"),
            content_type=INTERNAL_REPORT_MEDIA_TYPE,
            metadata={"run_id": run_id, "report_kind": "internal"},
            reuse=False,
        )
        session.commit()
    stray = EditorTokens(SECRET).url_token(
        "document",
        document_key=body["config"]["document"]["key"],
        artifact_id=foreign.artifact_id,
        now=_now(),
    )
    path = urlparse(url).path
    query = {"token": stray}
    answer = client.get(path, params=query, headers=_signed_for(path, query))
    assert answer.status_code == 404 and answer.json()["code"] == "editor_version_not_found"


# -- the callback (part 3) ---------------------------------------------------------


class Saves:
    """The Document Server's files, as a fetcher serves them, and every URL asked for."""

    def __init__(self) -> None:
        self.files: dict[str, bytes] = {}
        self.asked: list[str] = []
        self.refuse: str | None = None

    def put(self, data: bytes, name: str = "output.docx") -> str:
        url = f"{EDITOR}/cache/files/data/{len(self.files)}/{name}?md5=x&expires=1"
        self.files[url] = data
        return url

    def __call__(self, url: str) -> bytes:
        from aia_core.infrastructure.editor_fetch import SaveFetchRefused

        self.asked.append(url)
        if self.refuse:
            raise SaveFetchRefused(self.refuse)
        return self.files[url]


@pytest.fixture
def saves(app: FastAPI) -> Saves:
    app.state.editor_fetcher = fetched = Saves()
    return fetched


def _onlyoffice(claims: dict[str, Any]) -> str:
    """``claims`` signed as ONLYOFFICE signs (jsonwebtoken with ``expiresIn``): ``iat`` and
    ``exp`` at the top level, beside whatever it signs."""
    now = int(time.time())
    return jwt.encode({**claims, "iat": now, "exp": now + 300}, SECRET)


def _callback(
    client: TestClient, config: dict[str, Any], *, signed: str = "body", **body: Any
) -> Any:
    """``body`` posted to the config's callback URL as ONLYOFFICE posts it: signed in the body
    (``inBody``), or in the header as its ``payload``, or, for the test, both."""
    body = {"key": config["document"]["key"], **body}
    url = urlparse(config["editorConfig"]["callbackUrl"])
    sent = dict(body)
    headers: dict[str, str] = {}
    if signed in ("body", "both"):
        sent["token"] = _onlyoffice(body)
    if signed in ("header", "both"):
        headers["Authorization"] = f"Bearer {_onlyoffice({'payload': body})}"
    return client.post(url.path, params=parse_qs(url.query), json=sent, headers=headers)


def test_a_force_save_stores_the_next_koncept_with_its_changes_and_keeps_the_key(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any, saves: Saves
) -> None:
    run_id, draft = _reported_run(app, researcher, world)
    config = researcher.post(_editor(world, run_id), json={}).json()["config"]
    me = world.users["researcher"]
    url = saves.put(_docx("50 %"))
    answer = _callback(client, config, status=6, url=url, users=[me])
    assert answer.status_code == 200 and answer.json() == {"error": 0}, answer.text
    again = researcher.post(_editor(world, run_id), json={}).json()
    assert again["config"]["document"]["key"] == config["document"]["key"]
    assert again["koncept"] == 2 and again["line_id"] == draft
    assert again["problem"] is None
    factory = create_session_factory(app.state.engine)
    with factory() as session:
        principal = AuthenticatedPrincipal(user_id=me, organization_id=world.organization_id)
        scope = ScopeResolver(session).study_context(principal, study_id=world.study_id())
        repo = research_artifacts(session, scope, app.state.artifact_store)
        version = repo.get(again["version_artifact_id"])
        assert version.artifact_type == "research_internal_docx_edit"
        assert version.produced_by_user_id == me
        assert [d.artifact_id for d in repo.dependencies(version.artifact_id)] == [draft]
        assert version.metadata["koncept"] == 2 and version.metadata["authors"] == [me]
        stored = ReportEditorRepository(session, scope).version(version.artifact_id)
        assert stored is not None and stored.changes_artifact_id is not None
        changes = read_changes_payload(json.loads(repo.read(stored.changes_artifact_id)))
        assert [(c.kind.value, c.tag, c.was, c.now) for c in changes.changes] == [
            ("changed", "trust_total", "42,5 %", "50 %")
        ]
    # The same bytes again (a retried callback) store nothing new.
    assert _callback(client, config, status=6, url=saves.put(_docx("50 %")), users=[me]).json() == {
        "error": 0
    }
    assert researcher.post(_editor(world, run_id), json={}).json()["koncept"] == 2


@pytest.mark.parametrize("signed", ["header", "both"])
def test_a_save_signed_as_the_document_server_is_set_to_sign_it_is_stored(
    app: FastAPI,
    researcher: TestClient,
    client: TestClient,
    world: Any,
    saves: Saves,
    signed: str,
) -> None:
    """ONLYOFFICE signs a callback in the body or, by default, in the header as its payload;
    either is a save, and both together must agree."""
    run_id, _ = _reported_run(app, researcher, world)
    config = researcher.post(_editor(world, run_id), json={}).json()["config"]
    url = saves.put(_docx("50 %"))
    me = world.users["researcher"]
    answer = _callback(client, config, signed=signed, status=6, url=url, users=[me])
    assert answer.json() == {"error": 0}, answer.text
    assert researcher.post(_editor(world, run_id), json={}).json()["koncept"] == 2


def test_a_body_and_header_that_disagree_are_refused_and_only_signed_fields_are_read(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any, saves: Saves
) -> None:
    run_id, _ = _reported_run(app, researcher, world)
    config = researcher.post(_editor(world, run_id), json={}).json()["config"]
    me = world.users["researcher"]
    key = config["document"]["key"]
    callback = urlparse(config["editorConfig"]["callbackUrl"])
    good, evil = saves.put(_docx("50 %")), saves.put(_docx("99 %"))
    signed = {"key": key, "status": 6, "url": good, "users": [me]}
    answer = client.post(
        callback.path,
        params=parse_qs(callback.query),
        json={**signed, "token": _onlyoffice(signed)},
        headers={"Authorization": f"Bearer {_onlyoffice({'payload': {**signed, 'url': evil}})}"},
    )
    assert answer.status_code == 403 and saves.asked == []
    # The header signs a status-1; the body's own fields say "save this": nothing is fetched.
    answer = client.post(
        callback.path,
        params=parse_qs(callback.query),
        json={**signed, "status": 6, "url": evil},
        headers={"Authorization": f"Bearer {_onlyoffice({'payload': {'key': key, 'status': 1}})}"},
    )
    assert answer.json() == {"error": 0} and saves.asked == []


def test_a_save_for_a_session_already_closed_is_not_stored(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any, saves: Saves
) -> None:
    """A status-2 closed the session and a new one saved since: a late or replayed force-save
    for the old key (its callback URL lasts a week) must not become the newest koncept."""
    run_id, _ = _reported_run(app, researcher, world)
    config = researcher.post(_editor(world, run_id), json={}).json()["config"]
    me = world.users["researcher"]
    first, stale = saves.put(_docx("50 %")), saves.put(_docx("10 %"))
    assert _callback(client, config, status=2, url=first, users=[me]).json() == {"error": 0}
    later = researcher.post(_editor(world, run_id), json={}).json()["config"]
    newer = saves.put(_docx("60 %"))
    assert _callback(client, later, status=6, url=newer, users=[me]).json() == {"error": 0}
    asked = len(saves.asked)
    for status in (6, 2):
        replay = _callback(client, config, status=status, url=stale, users=[me])
        assert replay.json() == {"error": 1}
    assert _callback(client, config, status=4).json() == {"error": 0}  # nothing to do
    assert len(saves.asked) == asked
    opened = researcher.post(_editor(world, run_id), json={}).json()
    assert (
        opened["koncept"] == 3 and opened["config"]["document"]["key"] == later["document"]["key"]
    )


def test_everyone_leaving_closes_the_session_and_the_next_opens_a_new_key(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any, saves: Saves
) -> None:
    run_id, _ = _reported_run(app, researcher, world)
    config = researcher.post(_editor(world, run_id), json={}).json()["config"]
    me = world.users["researcher"]
    assert _callback(client, config, status=1, users=[me]).json() == {"error": 0}
    url = saves.put(_docx("50 %"))
    assert _callback(client, config, status=2, url=url, users=[me]).json() == {"error": 0}
    reopened = researcher.post(_editor(world, run_id), json={}).json()
    assert reopened["config"]["document"]["key"] != config["document"]["key"]
    assert reopened["koncept"] == 2
    closed = reopened["config"]
    assert _callback(client, closed, status=4).json() == {"error": 0}
    assert researcher.post(_editor(world, run_id), json={}).json()["config"]["document"][
        "key"
    ] not in (
        config["document"]["key"],
        closed["document"]["key"],
    )


def test_a_callback_naming_another_host_is_refused_unfetched(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any, saves: Saves
) -> None:
    run_id, _ = _reported_run(app, researcher, world)
    config = researcher.post(_editor(world, run_id), json={}).json()["config"]
    me = world.users["researcher"]
    for url in (
        "https://evil.example.test/cache/files/x.docx",
        f"{EDITOR}/healthcheck",
        f"{EDITOR}/cache/files/../../etc/passwd",
        f"{EDITOR}:8443/cache/files/x.docx",
    ):
        assert _callback(client, config, status=6, url=url, users=[me]).json() == {"error": 1}
    assert saves.asked == []
    opened = researcher.post(_editor(world, run_id), json={}).json()
    assert opened["koncept"] == 1 and opened["problem"] == "not_the_editors_file"


def test_a_restarted_document_servers_save_is_not_stored_and_said(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any, saves: Saves
) -> None:
    run_id, _ = _reported_run(app, researcher, world)
    config = researcher.post(_editor(world, run_id), json={}).json()["config"]
    saves.refuse = "document_server_restarted"
    url = saves.put(_docx("50 %"))
    answer = _callback(client, config, status=6, url=url, users=[world.users["researcher"]])
    assert answer.json() == {"error": 1}
    opened = researcher.post(_editor(world, run_id), json={}).json()
    assert opened["koncept"] == 1 and opened["problem"] == "document_server_restarted"


def test_a_save_by_someone_no_config_was_signed_for_is_refused(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any, saves: Saves
) -> None:
    run_id, _ = _reported_run(app, researcher, world)
    config = researcher.post(_editor(world, run_id), json={}).json()["config"]
    url = saves.put(_docx("50 %"))
    answer = _callback(client, config, status=6, url=url, users=[world.users["lead"]])
    assert answer.json() == {"error": 1} and saves.asked == []


def test_an_unsigned_or_wrongly_signed_callback_is_refused(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any, saves: Saves
) -> None:
    run_id, _ = _reported_run(app, researcher, world)
    config = researcher.post(_editor(world, run_id), json={}).json()["config"]
    url = urlparse(config["editorConfig"]["callbackUrl"])
    body = {"key": config["document"]["key"], "status": 4}
    wrong = jwt.encode(body, "w" * 48)
    for headers, token in (({}, None), ({"Authorization": f"Bearer {wrong}"}, wrong)):
        answer = client.post(
            url.path, params=parse_qs(url.query), json={**body, "token": token}, headers=headers
        )
        assert answer.status_code == 403
    # A body signed for another key than the URL's.
    other = {"key": "ART-0.aaaa.1", "status": 4}
    answer = client.post(
        url.path,
        params=parse_qs(url.query),
        json={**other, "token": jwt.encode(other, SECRET)},
        headers={"Authorization": f"Bearer {jwt.encode({'payload': other}, SECRET)}"},
    )
    assert answer.status_code == 403


def test_a_save_is_stored_as_its_authors_when_a_config_was_signed_for_them(
    app: FastAPI,
    researcher: TestClient,
    lead: TestClient,
    client: TestClient,
    world: Any,
    saves: Saves,
) -> None:
    run_id, _ = _reported_run(app, researcher, world)
    config = researcher.post(_editor(world, run_id), json={}).json()["config"]
    lead.post(_editor(world, run_id), json={})  # the lead joins the same session
    author = world.users["lead"]
    url = saves.put(_docx("50 %"))
    assert _callback(client, config, status=6, url=url, users=[author]).json() == {"error": 0}
    version_id = researcher.post(_editor(world, run_id), json={}).json()["version_artifact_id"]
    factory = create_session_factory(app.state.engine)
    with factory() as session:
        principal = AuthenticatedPrincipal(user_id=author, organization_id=world.organization_id)
        scope = ScopeResolver(session).study_context(principal, study_id=world.study_id())
        repo = research_artifacts(session, scope, app.state.artifact_store)
        assert repo.get(version_id).produced_by_user_id == author


def test_the_browsers_config_token_does_not_post_a_save(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any, saves: Saves
) -> None:
    """The browser holds a JWT under the editor's secret, its config's: it must not pass for
    the Document Server's callback, in the header or the body."""
    run_id, _ = _reported_run(app, researcher, world)
    config = researcher.post(_editor(world, run_id), json={}).json()["config"]
    url = urlparse(config["editorConfig"]["callbackUrl"])
    body = {
        "key": config["document"]["key"],
        "status": 6,
        "url": saves.put(_docx("50 %")),
        "users": [world.users["researcher"]],
        "token": config["token"],
    }
    answer = client.post(
        url.path,
        params=parse_qs(url.query),
        json=body,
        headers={"Authorization": f"Bearer {config['token']}"},
    )
    assert answer.status_code == 403 and saves.asked == []


def test_a_callback_without_a_token_or_a_body_is_refused_not_invalid(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any
) -> None:
    """Whatever the body, the route answers: 403 with the editor on, 404 off, never 422."""
    run_id, _ = _reported_run(app, researcher, world)
    callback = researcher.post(_editor(world, run_id), json={}).json()["config"]["editorConfig"]
    path = urlparse(callback["callbackUrl"]).path
    bodies: list[dict[str, Any]] = [
        {},
        {"content": b"not json", "headers": {"Content-Type": "application/json"}},
        {"json": [1]},
        {"content": b"\xff\xfe", "headers": {"Content-Type": "text/plain"}},
        {"json": {"token": 7}},
        {"content": b"[" * 100_000, "headers": {"Content-Type": "application/json"}},
    ]
    for kw in bodies:
        assert client.post(path, **kw).status_code == 403, kw
    app.state.settings = app.state.settings.model_copy(update={"onlyoffice_callback_base": ""})
    for kw in bodies:
        assert client.post(path, **kw).status_code == 404, kw


def test_a_removed_opener_s_session_takes_no_save(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any, saves: Saves
) -> None:
    """The callback is outside the study tree too: its scope is the session's, resolved through
    ScopeResolver on every callback, so a member removed since stores nothing."""
    from datetime import UTC, datetime

    from aia_core.infrastructure.scope_repository import ScopeRepository

    run_id, _ = _reported_run(app, researcher, world)
    config = researcher.post(_editor(world, run_id), json={}).json()["config"]
    factory = create_session_factory(app.state.engine)
    with factory() as session:
        owner = AuthenticatedPrincipal(
            user_id=world.users["owner"], organization_id=world.organization_id
        )
        admin = ScopeResolver(session).organization_context(owner)
        ScopeRepository(session).deactivate_member(
            admin, user_id=world.users["researcher"], now=datetime.now(UTC)
        )
        session.commit()
    url = saves.put(_docx("50 %"))
    answer = _callback(client, config, status=6, url=url, users=[world.users["researcher"]])
    assert answer.status_code == 404 and answer.json()["code"] == "editor_session_not_found"
    assert saves.asked == []


def test_a_save_naming_another_organizations_member_is_refused(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any, saves: Saves
) -> None:
    """Only the people a config was signed for may be a save's authors: a member of another
    organization (the acceptance world's other owner) is refused, and nothing is fetched."""
    from aia_core.infrastructure.scope_repository import ScopeRepository

    run_id, _ = _reported_run(app, researcher, world)
    config = researcher.post(_editor(world, run_id), json={}).json()["config"]
    factory = create_session_factory(app.state.engine)
    with factory() as session:
        _org, stranger = ScopeRepository(session).create_organization(
            slug="other-org", name="Other", owner_email="o@other.invalid", owner_name="O"
        )
        session.commit()
    url = saves.put(_docx("50 %"))
    answer = _callback(client, config, status=6, url=url, users=[stranger.user_id])
    assert answer.json() == {"error": 1} and saves.asked == []
    opened = researcher.post(_editor(world, run_id), json={}).json()
    assert opened["koncept"] == 1 and opened["problem"] == "unknown_author"


# -- an edited version's changes and its approval (part 4) -------------------------------


def test_an_edited_version_says_what_it_changed_and_is_approved_as_itself(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any, saves: Saves
) -> None:
    from aia_core.infrastructure.tables import ApprovalDecisionRow
    from sqlalchemy import func, select

    run_id, draft = _reported_run(app, researcher, world)
    other_run, _ = _reported_run(app, researcher, world, title="Večerní nápoj")
    config = researcher.post(_editor(world, run_id), json={}).json()["config"]
    me = world.users["researcher"]
    url = saves.put(_docx("50 %"))
    assert _callback(client, config, status=6, url=url, users=[me]).json() == {"error": 0}
    edited = researcher.post(_editor(world, run_id), json={}).json()["version_artifact_id"]
    versions = f"{_editor(world, run_id)}/versions"

    first = researcher.get(f"{versions}/{draft}")
    assert first.status_code == 200 and first.json()["koncept"] == 1
    assert first.json()["changes"] is None  # the composer's own changed nothing
    said = researcher.get(f"{versions}/{edited}").json()
    assert said["koncept"] == 2 and said["authors"] == [me] and said["edited_from"] == draft
    assert said["review_state"] == "DRAFT_UNAPPROVED" and said["approval"] is None
    assert said["changes"]["uncomparable"] is None
    assert said["changes"]["changes"] == [
        {"kind": "changed", "tag": "trust_total", "was": "42,5 %", "now": "50 %", "paragraph": None}
    ]
    # Never under another run of the study, nor another study's path.
    assert researcher.get(f"{_editor(world, other_run)}/versions/{edited}").status_code == 404
    sibling = f"{API}/studies/{world.study_id('sibling')}/research/runs/{run_id}/report/editor"
    assert researcher.get(f"{sibling}/versions/{edited}").status_code == 404

    runs = f"{API}/studies/{world.study_id()}/research/runs"
    wrong = researcher.post(f"{runs}/{other_run}/report/approve", json={"artifact_id": edited})
    assert wrong.status_code == 409 and wrong.json()["code"] == "report_not_approvable"
    approved = researcher.post(f"{runs}/{run_id}/report/approve", json={"artifact_id": edited})
    assert approved.status_code == 200, approved.text
    body = approved.json()
    assert body["artifact_id"] == edited and body["report_kind"] == "internal"
    assert body["artifact_type"] == "research_internal_docx_edit"
    assert body["internal_only"] is True
    again = researcher.post(f"{runs}/{run_id}/report/approve", json={"artifact_id": edited})
    assert again.status_code == 200, again.text

    # Approved as itself: one decision, frozen, said by the version and by the run's status.
    factory = create_session_factory(app.state.engine)
    with factory() as session:
        rows = session.scalar(
            select(func.count())
            .select_from(ApprovalDecisionRow)
            .where(ApprovalDecisionRow.subject_id == edited)
        )
        assert rows == 1
        principal = AuthenticatedPrincipal(user_id=me, organization_id=world.organization_id)
        scope = ScopeResolver(session).study_context(principal, study_id=world.study_id())
        frozen = research_artifacts(session, scope, app.state.artifact_store).get(edited)
        assert frozen.is_approved and frozen.is_frozen
    now = researcher.get(f"{versions}/{edited}").json()
    assert now["review_state"] == "APPROVED_INTERNAL" and now["approval"]["by_you"] is True
    status_ = researcher.get(f"{runs}/{run_id}/report").json()
    assert status_["artifact_id"] == draft and status_["review_state"] == "DRAFT_UNAPPROVED"
    (named,) = status_["approved_edits"]
    assert named["artifact_id"] == edited
    assert named["koncept"] == 2 and named["review_state"] == "APPROVED_INTERNAL"

    # Downloaded as exactly the version approved, its draft footer said in its name.
    got = researcher.get(f"{versions}/{edited}/download")
    assert got.status_code == 200 and got.content == _docx("50 %")
    assert got.headers["X-AIA-Artifact-Id"] == edited
    assert got.headers["X-AIA-Review-State"] == "APPROVED_INTERNAL"
    assert "schvaleno-v-aia" in got.headers["Content-Disposition"]

    # A later save is the next koncept; the approved version stays as it was.
    later = saves.put(_docx("55 %"))
    assert _callback(client, config, status=6, url=later, users=[me]).json() == {"error": 0}
    assert researcher.post(_editor(world, run_id), json={}).json()["koncept"] == 3
    assert researcher.get(f"{versions}/{edited}/download").content == _docx("50 %")

    # Readable with the editor off: an approved version is no editing session's.
    app.state.settings = app.state.settings.model_copy(update={"onlyoffice_callback_base": ""})
    assert researcher.get(f"{versions}/{edited}").status_code == 200
    assert researcher.get(f"{versions}/{edited}/download").status_code == 200


def _editor_saved_docx() -> bytes:
    """A DOCX as ONLYOFFICE saves one AIA composed: a real package whose font table still
    names a brand face, its embedded parts and their relationships dropped (RPT15)."""
    import docx
    from aia_core.infrastructure.report_docx.embed import embed_fonts
    from lxml import etree

    document = docx.Document()
    document.add_paragraph("Bance důvěřuje 50 %")
    embed_fonts(document, ["IBM Plex Sans"])
    built = io.BytesIO()
    document.save(built)
    source = zipfile.ZipFile(built)
    table = etree.fromstring(source.read("word/fontTable.xml"))
    for font in table.iter(f"{{{W}}}font"):
        for child in [c for c in font if c.tag.startswith(f"{{{W}}}embed")]:
            font.remove(child)
    rels = etree.fromstring(source.read("word/_rels/fontTable.xml.rels"))
    for rel in [r for r in rels if r.get("Target", "").startswith("fonts/")]:
        rels.remove(rel)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for name in source.namelist():
            if name.startswith("word/fonts/"):
                continue
            data = source.read(name)
            if name == "word/fontTable.xml":
                data = etree.tostring(table, xml_declaration=True, encoding="UTF-8")
            elif name == "word/_rels/fontTable.xml.rels":
                data = etree.tostring(rels, xml_declaration=True, encoding="UTF-8")
            z.writestr(zipfile.ZipInfo(name, (2026, 10, 26, 0, 0, 0)), data)
    return out.getvalue()


def test_an_editor_saved_version_downloads_with_its_brand_fonts_embedded_again(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any, saves: Saves
) -> None:
    """RPT15 (the owner's B6 (b)): ONLYOFFICE drops the embedded brand fonts when it saves, so
    the download embeds them again. The headers still name the stored version (its id, its
    SHA-256, its review), the download says the fonts were re-embedded, and approving the
    version changes none of it."""
    run_id, _draft = _reported_run(app, researcher, world)
    config = researcher.post(_editor(world, run_id), json={}).json()["config"]
    me = world.users["researcher"]
    stored = _editor_saved_docx()
    saved = _callback(client, config, status=6, url=saves.put(stored), users=[me])
    assert saved.json() == {"error": 0}
    edited = researcher.post(_editor(world, run_id), json={}).json()["version_artifact_id"]
    versions = f"{_editor(world, run_id)}/versions"
    sha = researcher.get(f"{versions}/{edited}").json()["sha256"]
    assert sha == hashlib.sha256(stored).hexdigest()

    def download() -> Any:
        got = researcher.get(f"{versions}/{edited}/download")
        assert got.status_code == 200, got.text
        assert got.headers["X-AIA-Artifact-Id"] == edited
        assert got.headers["X-AIA-Artifact-Sha256"] == sha
        assert got.headers["X-AIA-Fonts-Reembedded"] == "true"
        package = zipfile.ZipFile(io.BytesIO(got.content))
        assert any(n.startswith("word/fonts/") and n.endswith(".odttf") for n in package.namelist())
        assert got.content != stored
        return got

    first = download()
    assert first.headers["X-AIA-Review-State"] == "DRAFT_UNAPPROVED"
    assert download().content == first.content  # the same version, the same bytes
    approve = f"{API}/studies/{world.study_id()}/research/runs/{run_id}/report/approve"
    assert researcher.post(approve, json={"artifact_id": edited}).status_code == 200
    after = download()
    assert after.headers["X-AIA-Review-State"] == "APPROVED_INTERNAL"
    assert researcher.get(f"{versions}/{edited}").json()["sha256"] == sha  # stored as it was


def test_a_version_names_its_authors_and_its_lines_newest_version(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any, saves: Saves
) -> None:
    run_id, draft = _reported_run(app, researcher, world)
    config = researcher.post(_editor(world, run_id), json={}).json()["config"]
    me = world.users["researcher"]
    name = config["editorConfig"]["user"]["name"]
    versions = f"{_editor(world, run_id)}/versions"
    first = saves.put(_docx("50 %"))
    assert _callback(client, config, status=6, url=first, users=[me]).json() == {"error": 0}
    edited = researcher.post(_editor(world, run_id), json={}).json()["version_artifact_id"]

    said = researcher.get(f"{versions}/{edited}").json()
    assert said["author_names"] == [name] and "@" not in name
    assert said["content_at"] is None  # this save named no lastsave
    assert datetime.fromisoformat(said["server_now"]).tzinfo is not None
    assert said["newest_in_line"] == {"version_artifact_id": edited, "koncept": 2}
    assert researcher.get(f"{versions}/{draft}").json()["newest_in_line"]["koncept"] == 2

    later = saves.put(_docx("55 %"))
    assert _callback(client, config, status=6, url=later, users=[me]).json() == {"error": 0}
    moved = researcher.get(f"{versions}/{edited}").json()["newest_in_line"]
    assert moved["koncept"] == 3 and moved["version_artifact_id"] != edited


def test_a_version_asked_as_its_lines_newest_is_refused_once_a_later_save_exists(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any, saves: Saves
) -> None:
    run_id, draft = _reported_run(app, researcher, world)
    config = researcher.post(_editor(world, run_id), json={}).json()["config"]
    me = world.users["researcher"]
    saved = _callback(client, config, status=6, url=saves.put(_docx("50 %")), users=[me])
    assert saved.json() == {"error": 0}
    second = researcher.post(_editor(world, run_id), json={}).json()["version_artifact_id"]
    # The question was open while a save stored koncept 3.
    saved = _callback(client, config, status=6, url=saves.put(_docx("55 %")), users=[me])
    assert saved.json() == {"error": 0}
    third = researcher.post(_editor(world, run_id), json={}).json()["version_artifact_id"]
    approve = f"{API}/studies/{world.study_id()}/research/runs/{run_id}/report/approve"

    stale = researcher.post(approve, json={"artifact_id": second, "newest_of_line": True})
    assert stale.status_code == 409, stale.text
    assert stale.json()["code"] == "report_version_stale"
    assert stale.json()["details"] == {"newest_artifact_id": third, "koncept": 3}
    assert researcher.get(f"{_editor(world, run_id)}/versions/{second}").json()["approval"] is None
    # Koncept 1, the composed draft, is no newest once the line moved on either.
    first = researcher.post(approve, json={"artifact_id": draft, "newest_of_line": True})
    assert first.status_code == 409 and first.json()["details"]["koncept"] == 3

    newest = researcher.post(approve, json={"artifact_id": third, "newest_of_line": True})
    assert newest.status_code == 200, newest.text
    assert newest.json()["artifact_id"] == third
    # Without the flag an older version is approved as itself, as before, the newer koncept on
    # its record; the newest's own approval recorded none.
    assert newest.json()["approval"]["warnings"] in (None, [])
    older = researcher.post(approve, json={"artifact_id": second})
    assert older.status_code == 200, older.text
    assert older.json()["approval"]["warnings"] == [
        {
            "code": "superseded_version",
            "where": third,
            "message": "schválen koncept 2, novější je koncept 3",
            "koncept": 3,
        }
    ]


def test_the_reports_status_names_the_version_saved_over_its_composed_draft(
    app: FastAPI,
    researcher: TestClient,
    viewer: TestClient,
    client: TestClient,
    world: Any,
    saves: Saves,
) -> None:
    """A pane that cannot show the editor reads the composed draft's status: once a person saved
    a version over it, the status names that version and its review (``edited_line``), so the
    pane offers no approval of the draft beside it. Any member who reads the report reads it."""
    run_id, _draft = _reported_run(app, researcher, world)
    status = f"{API}/studies/{world.study_id()}/research/runs/{run_id}/report"
    assert researcher.get(status).json()["edited_line"] is None
    config = researcher.post(_editor(world, run_id), json={}).json()["config"]
    me = world.users["researcher"]
    # Opening the line (koncept 1) saves nothing.
    assert researcher.get(status).json()["edited_line"] is None
    saved = _callback(client, config, status=6, url=saves.put(_docx("50 %")), users=[me])
    assert saved.json() == {"error": 0}
    edited = researcher.post(_editor(world, run_id), json={}).json()["version_artifact_id"]
    line = viewer.get(status)
    assert line.status_code == 200, line.text
    assert line.json()["edited_line"] == {
        "artifact_id": edited,
        "koncept": 2,
        "stored_at": ANY,
        "review_state": "DRAFT_UNAPPROVED",
        "approved_on": None,
        "stamped": False,
    }
    approve = f"{status}/approve"
    assert researcher.post(approve, json={"artifact_id": edited}).status_code == 200
    assert researcher.get(status).json()["edited_line"]["review_state"] == "APPROVED_INTERNAL"


def test_a_save_keeps_the_time_of_its_last_change_as_the_document_server_said_it(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any, saves: Saves
) -> None:
    run_id, _draft = _reported_run(app, researcher, world)
    config = researcher.post(_editor(world, run_id), json={}).json()["config"]
    me = world.users["researcher"]
    saved = _callback(
        client,
        config,
        status=6,
        url=saves.put(_docx("50 %")),
        users=[me],
        lastsave="2026-10-24T00:30:12.000Z",
    )
    assert saved.json() == {"error": 0}
    edited = researcher.post(_editor(world, run_id), json={}).json()["version_artifact_id"]
    said = researcher.get(f"{_editor(world, run_id)}/versions/{edited}").json()
    assert said["content_at"] == "2026-10-24T00:30:12+00:00"
    # A time that is not an aware ISO time is kept as nothing, never guessed.
    _callback(client, config, status=6, url=saves.put(_docx("51 %")), users=[me], lastsave="today")
    later = researcher.post(_editor(world, run_id), json={}).json()["version_artifact_id"]
    assert researcher.get(f"{_editor(world, run_id)}/versions/{later}").json()["content_at"] is None


def test_an_author_is_named_only_as_a_member_and_never_by_an_e_mail(
    app: FastAPI, world: Any
) -> None:
    from aia_core.application.report_editor import ReportEditor
    from aia_core.infrastructure.tables import UserRow

    me, lead = world.users["researcher"], world.users["lead"]
    factory = create_session_factory(app.state.engine)
    with factory() as session:
        session.get(UserRow, me).display_name = "Jana Výzkumná"
        session.get(UserRow, lead).display_name = "lead@art-chain.io"
        session.add(UserRow(user_id="USR-stranger", email="x@else.test", display_name="Cizí"))
        session.commit()
        principal = AuthenticatedPrincipal(user_id=me, organization_id=world.organization_id)
        scope = ScopeResolver(session).study_context(principal, study_id=world.study_id())
        names = ReportEditor(session, scope).names([me, lead, "USR-stranger"])
    assert names == {me: "Jana Výzkumná"}


def test_no_author_of_an_edited_version_approves_it_where_self_approval_is_off(
    app: FastAPI,
    researcher: TestClient,
    lead: TestClient,
    owner: TestClient,
    outsider: TestClient,
    client: TestClient,
    world: Any,
    saves: Saves,
) -> None:
    """An edited version records one producer, its save's first author; where independent
    review is on, everyone who saved the line up to it is refused, not only that one."""
    forbid = owner.put(
        f"{API}/self-approval", json={"allowed": False, "client_id": world.client_id()}
    )
    assert forbid.status_code == 200, forbid.text
    run_id, _ = _reported_run(app, researcher, world)
    config = researcher.post(_editor(world, run_id), json={}).json()["config"]
    lead.post(_editor(world, run_id), json={})  # the lead joins the same session
    me, colleague = world.users["researcher"], world.users["lead"]
    second = saves.put(_docx("50 %"))
    assert _callback(client, config, status=6, url=second, users=[colleague]).json() == {"error": 0}
    third = saves.put(_docx("60 %"))
    assert _callback(client, config, status=6, url=third, users=[me, colleague]).json() == {
        "error": 0
    }
    edited = researcher.post(_editor(world, run_id), json={}).json()["version_artifact_id"]
    approve = f"{API}/studies/{world.study_id()}/research/runs/{run_id}/report/approve"
    for author in (researcher, lead):
        refused = author.post(approve, json={"artifact_id": edited})
        assert refused.status_code == 403 and refused.json()["code"] == "separation_of_duties"
    approved = outsider.post(approve, json={"artifact_id": edited})
    assert approved.status_code == 200, approved.text


def test_a_co_author_who_approves_an_edited_version_is_recorded_as_self_approving(
    app: FastAPI,
    researcher: TestClient,
    lead: TestClient,
    client: TestClient,
    world: Any,
    saves: Saves,
) -> None:
    """Self-approval on (the default): the save's second author may approve, and the ledger
    says it is their own work, though the artifact records the first author as producer."""
    run_id, _ = _reported_run(app, researcher, world)
    config = researcher.post(_editor(world, run_id), json={}).json()["config"]
    lead.post(_editor(world, run_id), json={})
    users = [world.users["researcher"], world.users["lead"]]
    url = saves.put(_docx("50 %"))
    assert _callback(client, config, status=6, url=url, users=users).json() == {"error": 0}
    edited = researcher.post(_editor(world, run_id), json={}).json()["version_artifact_id"]
    approve = f"{API}/studies/{world.study_id()}/research/runs/{run_id}/report/approve"
    assert lead.post(approve, json={"artifact_id": edited}).status_code == 200
    said = lead.get(f"{_editor(world, run_id)}/versions/{edited}").json()["approval"]
    assert said["self_approved"] is True and said["by_you"] is True


def test_a_versions_download_reads_its_docx_alone_not_what_it_changed(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any, saves: Saves
) -> None:
    """The changes stored beside a version are the version route's to read: a download of
    an intact DOCX never fails because the changes beside it are damaged or gone."""
    run_id, _ = _reported_run(app, researcher, world)
    config = researcher.post(_editor(world, run_id), json={}).json()["config"]
    url = saves.put(_docx("50 %"))
    me = world.users["researcher"]
    assert _callback(client, config, status=6, url=url, users=[me]).json() == {"error": 0}
    edited = researcher.post(_editor(world, run_id), json={}).json()["version_artifact_id"]
    factory = create_session_factory(app.state.engine)
    with factory() as session:
        principal = AuthenticatedPrincipal(user_id=me, organization_id=world.organization_id)
        scope = ScopeResolver(session).study_context(principal, study_id=world.study_id())
        version = ReportEditorRepository(session, scope).version(edited)
        assert version is not None and version.changes_artifact_id is not None
        repo = research_artifacts(session, scope, app.state.artifact_store)
        key = repo.get(version.changes_artifact_id).storage_key
    versions = f"{_editor(world, run_id)}/versions/{edited}"
    for damage in ("tampered", "gone"):
        if damage == "tampered":
            app.state.artifact_store.put(key, b'{"tampered": true}')
        else:
            app.state.artifact_store.delete(key)
        got = researcher.get(f"{versions}/download")
        assert got.status_code == 200 and got.content == _docx("50 %"), damage
        assert researcher.get(versions).status_code == 409, damage  # said, never guessed


def test_the_first_version_serves_the_pdf_its_report_step_rendered(
    app: FastAPI, researcher: TestClient, world: Any
) -> None:
    """Koncept 1 of the run's own line is the report its step composed: its PDF is that
    report's, the step's rendering included, and never a second one asked for."""
    from aia_core.infrastructure.tables import WorkflowRunRow
    from sqlalchemy import func, select

    run_id, draft = _reported_run(app, researcher, world, with_pdf=True)
    researcher.post(_editor(world, run_id), json={})  # starts the line: koncept 1 is the draft
    versions = f"{_editor(world, run_id)}/versions/{draft}"
    said = researcher.get(versions).json()["pdf"]
    report = researcher.get(f"{API}/studies/{world.study_id()}/research/runs/{run_id}/report")
    assert said["state"] == "READY" and said == report.json()["pdf"]

    def jobs() -> int:
        with create_session_factory(app.state.engine)() as session:
            return int(session.scalar(select(func.count()).select_from(WorkflowRunRow)) or 0)

    before = jobs()
    asked = researcher.post(f"{versions}/pdf")
    assert asked.status_code == 200 and asked.json()["pdf"]["state"] == "READY"
    assert jobs() == before
    got = researcher.get(f"{versions}/pdf/download")
    assert got.status_code == 200 and got.content == b"%PDF-1.7 the step's rendering"
    assert got.headers["X-AIA-Artifact-Id"] == draft


def _stamp(at: Any) -> str:
    return at.isoformat(timespec="milliseconds").replace("+00:00", "Z")


def _ago(seconds: float) -> str:
    from datetime import UTC, datetime, timedelta

    return _stamp(datetime.now(UTC) - timedelta(seconds=seconds))


PAGE = "page-0123456789"


class _Line:
    """One run's report line open in the editor, as the RPT8f tests drive it."""

    def __init__(
        self, app: FastAPI, researcher: TestClient, client: TestClient, world: Any, saves: Saves
    ) -> None:
        self.run_id, self.draft = _reported_run(app, researcher, world)
        self.researcher, self.client, self.world, self.saves = researcher, client, world, saves
        self.config = researcher.post(_editor(world, self.run_id), json={}).json()["config"]
        self.key = self.config["document"]["key"]
        self.edits = f"{_editor(world, self.run_id)}/sessions/{self.key}/edits"
        self.approve_url = (
            f"{API}/studies/{world.study_id()}/research/runs/{self.run_id}/report/approve"
        )

    def report(self, seq: int, *, page: str = PAGE, by: TestClient | None = None) -> int:
        body = {"page": page, "seq": seq}
        return (by or self.researcher).post(self.edits, json=body).status_code

    def save(self, value: str, *, lastsave: str, by: str = "researcher") -> str:
        answer = _callback(
            self.client,
            self.config,
            status=6,
            url=self.saves.put(_docx(value)),
            users=[self.world.users[by]],
            lastsave=lastsave,
        )
        assert answer.json() == {"error": 0}, answer.text
        opened = self.researcher.post(_editor(self.world, self.run_id), json={}).json()
        return str(opened["version_artifact_id"])

    def approve(
        self, version: str, seq: int | None = None, *, newest: bool = True, page: str = PAGE
    ) -> Any:
        body: dict[str, Any] = {"artifact_id": version}
        if newest:
            body["newest_of_line"] = True
        if seq is not None:
            body["seen_edit"] = {"document_key": self.key, "page": page, "seq": seq}
        return self.researcher.post(self.approve_url, json=body)

    def verdict(self, version: str, seq: int | None = None) -> Any:
        params = {"seen_key": self.key, "seen_page": PAGE, "seen_seq": seq} if seq else {}
        url = f"{_editor(self.world, self.run_id)}/versions/{version}"
        return self.researcher.get(url, params=params).json()["edits"]


def _held(answer: Any, reason: str) -> None:
    assert answer.status_code == 409, answer.text
    assert answer.json()["code"] == "report_edits_unsaved"
    assert answer.json()["details"] == {"reason": reason}


def test_an_approval_waits_until_a_save_holds_every_edit_the_server_received(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any, saves: Saves
) -> None:
    """RPT8f: the server, not the page, decides whether the version holds the edits made."""
    line = _Line(app, researcher, client, world, saves)
    assert line.report(1) == 204
    older = line.save("50 %", lastsave=_ago(3600))  # its last change long before the edit's report
    _held(line.approve(older, 1), "edit_not_saved")
    assert line.verdict(older, 1) == {"saved": False, "reason": "edit_not_saved"}
    # The page made edit 2 and its report has not arrived: the approval leaving first waits.
    _held(line.approve(older, 2), "edit_not_received")
    assert line.verdict(older, 2) == {"saved": False, "reason": "edit_not_received"}
    url = f"{_editor(world, line.run_id)}/versions/{older}"
    assert researcher.get(url).json()["approval"] is None
    assert researcher.get(url, params={"seen_key": line.key}).status_code == 422  # all or none
    # Edit 2 reported, then a save whose last change follows it holds both.
    assert line.report(2) == 204
    newer = line.save("51 %", lastsave=_ago(0))
    assert line.report(1) == 204  # late or twice: the page's highest is kept
    assert line.verdict(newer, 2) == {"saved": True, "reason": None}
    assert line.approve(newer, 2).status_code == 200


def test_a_lastsave_ahead_of_aias_clock_counts_as_aias_receipt_and_never_holds_a_later_edit(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any, saves: Saves
) -> None:
    """A Document Server clock running ahead must not make a save hold an edit made after it:
    its lastsave is kept no later than AIA's receipt of the callback (fails closed)."""
    from datetime import UTC, datetime, timedelta

    line = _Line(app, researcher, client, world, saves)
    before = datetime.now(UTC)
    ahead = line.save("50 %", lastsave=_stamp(datetime.now(UTC) + timedelta(minutes=5)))
    said = researcher.get(f"{_editor(world, line.run_id)}/versions/{ahead}").json()
    assert before <= datetime.fromisoformat(said["content_at"]) <= datetime.now(UTC)
    assert line.report(1) == 204  # typed after that save
    _held(line.approve(ahead, 1), "edit_not_saved")


def test_an_edit_report_arriving_after_its_save_is_released_by_the_next_save_of_the_same_bytes(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any, saves: Saves
) -> None:
    """The wedge: the save holding an edit landed before the edit's report, so it does not count
    it. A later save of the same bytes stores no koncept but says a later lastsave, and releases
    it; one whose lastsave is still earlier (a Document Server clock behind) does not: fails
    closed."""
    line = _Line(app, researcher, client, world, saves)
    version = line.save("50 %", lastsave=_ago(1))
    assert line.report(1) == 204  # late
    _held(line.approve(version, 1), "edit_not_saved")
    assert line.save("50 %", lastsave=_ago(60)) == version  # the clock behind: still held
    _held(line.approve(version, 1), "edit_not_saved")
    assert line.save("50 %", lastsave=_ago(0)) == version  # no new koncept
    assert researcher.post(_editor(world, line.run_id), json={}).json()["koncept"] == 2
    assert line.verdict(version, 1) == {"saved": True, "reason": None}
    assert line.approve(version, 1).status_code == 200


def test_a_co_authors_edit_and_an_earlier_page_of_the_session_hold_the_approval_too(
    app: FastAPI,
    researcher: TestClient,
    lead: TestClient,
    client: TestClient,
    world: Any,
    saves: Saves,
) -> None:
    line = _Line(app, researcher, client, world, saves)
    version = line.save("50 %", lastsave=_ago(1))
    # Not signed into the session yet: the lead's report is no session of theirs.
    assert line.report(1, page="page-lead-0000", by=lead) == 404
    lead.post(_editor(world, line.run_id), json={})
    assert line.report(1, page="page-lead-0000", by=lead) == 204
    _held(line.approve(version), "edit_not_saved")
    assert line.verdict(version) == {"saved": False, "reason": "edit_not_saved"}
    # A reload is a new page: the edit the old page reported still waits for its save.
    assert line.report(1, page="page-old-00000") == 204
    _held(line.approve(version, page="page-new-00000"), "edit_not_saved")
    newest = line.save("52 %", lastsave=_ago(0), by="lead")
    assert line.approve(newest).status_code == 200


def test_every_edited_version_is_checked_whether_or_not_the_caller_asks_for_the_newest(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any, saves: Saves
) -> None:
    line = _Line(app, researcher, client, world, saves)
    version = line.save("50 %", lastsave=_ago(1))
    assert line.report(1) == 204
    _held(line.approve(version, newest=False), "edit_not_saved")


def test_an_edit_in_the_seen_session_still_holds_once_that_session_closed(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any, saves: Saves
) -> None:
    """The page's own session closed under it (everyone else left) before its edit was saved:
    the approval naming it is not let through by the close."""
    line = _Line(app, researcher, client, world, saves)
    version = line.save("50 %", lastsave=_ago(1))
    assert line.report(1) == 204
    _callback(client, line.config, status=4, users=[world.users["researcher"]])
    _held(line.approve(version, 1), "edit_not_saved")


def test_an_edit_is_reported_only_in_an_open_session_of_the_run_the_caller_was_signed_into(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any, saves: Saves
) -> None:
    line = _Line(app, researcher, client, world, saves)
    body = {"page": PAGE, "seq": 1}
    nowhere = f"{_editor(world, line.run_id)}/sessions/NO-SUCH/edits"
    assert researcher.post(nowhere, json=body).status_code == 404
    other = f"{API}/studies/{world.study_id()}/research/runs/RUN-0000000000000000/report/editor"
    assert researcher.post(f"{other}/sessions/{line.key}/edits", json=body).status_code in (
        404,
        422,
    )
    for bad in (
        {"page": "short", "seq": 1},
        {"page": PAGE, "seq": 0},
        {"page": PAGE, "seq": 1, "clean": True},
    ):
        assert researcher.post(line.edits, json=bad).status_code == 422
    # Everyone left: the session is closed and takes no edit.
    _callback(client, line.config, status=4, users=[world.users["researcher"]])
    assert researcher.post(line.edits, json=body).status_code == 404


def test_a_session_keeps_edit_reports_for_a_bounded_number_of_pages(
    app: FastAPI, researcher: TestClient, client: TestClient, world: Any, saves: Saves
) -> None:
    from aia_core.domain.report.editor import MAX_PAGES_PER_SESSION

    line = _Line(app, researcher, client, world, saves)
    pages = [f"page-{n:010d}" for n in range(MAX_PAGES_PER_SESSION + 1)]
    assert all(line.report(1, page=p) == 204 for p in pages[:-1])
    refused = researcher.post(line.edits, json={"page": pages[-1], "seq": 1})
    assert refused.status_code == 409 and refused.json()["code"] == "editor_pages_exhausted"
    assert line.report(2, page=pages[0]) == 204  # a page it keeps still reports
