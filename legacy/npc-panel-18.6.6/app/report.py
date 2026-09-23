"""
NPC PANEL — report.py
Export behu (run_dotaznik / run_pruzkum) do klientskeho XLSX.

    from report import export_xlsx
    export_xlsx(vysledky, "vystup.xlsx")

Listy:
  Prehled     — zadani, metodika, naklady, disclaimer o syntetickem panelu
  <ID otazky> — celkove rozlozeni + crosstaby podle segmentu
  Verbatimy   — otevrene odpovedi (jen kdyz jsou)
  Data        — klientsky bezpecny vyrez (bez latentnich D/M/F atributu)
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from runtime_config import RELEASE, PANEL_VERSION as DATA_PANEL_VERSION

FONT = "Arial"
NADPIS = Font(name=FONT, size=14, bold=True)
H2 = Font(name=FONT, size=11, bold=True, color="FFFFFF")
BEZNY = Font(name=FONT, size=10)
TUCNY = Font(name=FONT, size=10, bold=True)
MALY = Font(name=FONT, size=8, italic=True, color="808080")
FILL_H = PatternFill("solid", fgColor="1F3864")
FILL_ALT = PatternFill("solid", fgColor="EDF2F9")
FILL_VAR = PatternFill("solid", fgColor="FFF2CC")
RAMECEK = Border(*[Side(style="thin", color="BFBFBF")] * 4)

DISCLAIMER = (
    "Data pocházejí ze syntetického respondentního panelu (silicon sampling), "
    "ne z terénního sběru. Panel stojí na reálné sociodemografii ze 6 výzkumů ČR "
    "(CVVM / CHPS / Trendy Česka, archiv CSDA) a je kalibrován na populační "
    "marginály ČSÚ. Psychologické a dispoziční vrstvy jsou simulované, ne měřené — "
    "u každé dimenze panel eviduje zdroj (SPECIALIST = kalibrováno na konkrétní "
    "výzkum, OWN_ESTIMATE = odhad autorů bez měřeného marginálu). "
    "Výsledky jsou simulované odhady pro agregátní výzkumné použití a jejich "
    "prediktivní validita musí být ověřována na odděleném lidském holdoutu. "
    "Politické, zdravotní a další citlivé atributy jsou "
    "syntetické tendence pro výzkumnou simulaci, ne fakta o reálné osobě, a nesmí "
    "sloužit k rozhodování o jednotlivci. "
    "Před rozhodnutím s vysokou sázkou ověřte klíčová zjištění terénním sběrem "
    "(holdout validace proti reálným datům zatím neproběhla — viz validace.py)."
)

PANEL_VERSION = f"{DATA_PANEL_VERSION} / {RELEASE}"


def _sirky(ws, sirky: list[int]) -> None:
    for i, w in enumerate(sirky, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _hlavicka(ws, radek: int, hodnoty: list[str]) -> None:
    for j, h in enumerate(hodnoty, 1):
        c = ws.cell(radek, j, h)
        c.font, c.fill, c.border = H2, FILL_H, RAMECEK
        c.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")


def _list_prehled(wb: Workbook, v: dict) -> None:
    ws = wb.create_sheet("Přehled")
    _sirky(ws, [26, 90])
    ws["A1"], ws["A1"].font = v.get("nazev", v.get("otazka", "Průzkum")), NADPIS
    radky = [
        ("Velikost vzorku", f"{v.get('n_dotazano', '?')} respondentů"),
        ("Neplatných odpovědí", str(v.get("n_chyb", 0))),
        ("Filtry vzorku", str(v.get("filtry") or "žádné — reprezentativní ČR 18+")),
        ("Run ID", str(v.get("run_id", "n/a"))),
        ("Model", f"{v.get('model')} ({v.get('mode')})"),
        ("Response mode", str(v.get("response_mode", "choice"))),
        ("Execution", str(v.get("execution", "n/a"))),
        ("Doba běhu", f"{v.get('trvani_s')} s"),
        ("Náklady", f"${v.get('naklady_usd')}"),
        ("Tokeny in/out", f"{v.get('tokeny', {}).get('in')} / {v.get('tokeny', {}).get('out')}"),
        ("Seed", str(v.get("seed"))),
        ("Verze panelu/buildu", PANEL_VERSION),
    ]
    seg = v.get("segment") or {}
    if seg.get("mode") and seg.get("mode") != "none":
        radky += [
            ("Segment", str(seg.get("name", "learned segment"))),
            ("Segment confidence", str(seg.get("confidence_class", "n/a"))),
            ("Segment prevalence", f"{100*float(seg.get('prevalence_calibrated',0)):.2f} %"),
            ("Segment ESS", str(seg.get("ess_population", "n/a"))),
        ]
    r = 3
    for k, val in radky:
        ws.cell(r, 1, k).font = TUCNY
        ws.cell(r, 2, val).font = BEZNY
        r += 1
    r += 1
    ws.cell(r, 1, "Metodická poznámka").font = TUCNY
    c = ws.cell(r, 2, DISCLAIMER)
    c.font, c.fill = BEZNY, FILL_VAR
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 78

    r += 2
    ws.cell(r, 1, "Vážení").font = TUCNY
    if (v.get("segment") or {}).get("mode") not in (None, "none"):
        wt = "Vzorek je PPS-targetovaný podle kalibrační váhy × segment propensity. Segment je syntetický/odvozený podle uvedené confidence class; ESS je uvedeno výše."
    else:
        wt = "Vzorek je losován s pravděpodobností úměrnou kalibrační váze (self-weighting)."
    ws.cell(r, 2, wt).font = BEZNY


def _list_otazka(wb: Workbook, oid: str, otazka: dict, a: dict,
                 dmap: dict[str, str] | None = None, nrow: int = 0) -> None:
    nazev = oid[:31]
    ws = wb.create_sheet(nazev)
    _sirky(ws, [34] + [12] * 12)
    ws["A1"], ws["A1"].font = otazka["text"], NADPIS
    extra = ""
    if a.get("mean_response_entropy") is not None:
        extra += (f"   mean entropy={a['mean_response_entropy']}"
                  f"   mean maxP={a.get('mean_max_probability')}")
    ws["A2"] = (f"typ: {a['typ']}   eligible: {a.get('n_eligible', a['n_platnych'])}   "
                f"raw n={a['n_platnych']}   core donors={a.get('n_unique_core_donors','—')}   "
                f"layer={a.get('donor_layer','core')} / donors={a.get('n_unique_layer_donors','—')}   "
                f"donor-aware n={a.get('effective_n','—')}   support={a.get('support_status','—')}   chybí={a['n_chybi']}" +
                (f"   MC ±{a['mc95_pp']} p.b." if a.get("mc95_pp") else "") + extra)
    ws["A2"].font = MALY

    r = 4
    if a["typ"] == "skala":
        _hlavicka(ws, r, ["Ukazatel", "Hodnota"])
        r += 1
        for k, lbl in [("prumer", "Průměr vylosovaných odpovědí"), ("median", "Medián"), ("sd", "Sm. odchylka"),
                       ("top2box_pct", "Top-2-box (%)")]:
            ws.cell(r, 1, lbl).font = BEZNY
            ws.cell(r, 2, a.get(k)).font = BEZNY
            r += 1
        if a.get("expected_mean") is not None:
            ws.cell(r, 1, "Průměr z LLM pravděpodobností").font = TUCNY
            ws.cell(r, 2, a.get("expected_mean")).font = TUCNY
            r += 1
        r += 1
        _hlavicka(ws, r, ["Hodnota škály", "Počet"])
        r += 1
        for k, x in a.get("rozlozeni", {}).items():
            ws.cell(r, 1, k).font = BEZNY
            ws.cell(r, 2, x).font = BEZNY
            r += 1
        volby: list[str] = []
    elif a["typ"] == "otevrena":
        ws.cell(r, 1, f"{a['n_platnych']} verbatimů — list „Verbatimy“").font = BEZNY
        return
    else:
        volby = list(a["celkem_pct"].keys())
        vzorec = bool(dmap) and oid in dmap and a["typ"] == "vyber"
        rng = (f"Data!${dmap[oid]}$2:${dmap[oid]}${nrow}" if vzorec else "")
        _hlavicka(ws, r, ["Odpověď", "Podíl (%)", "n"])
        r += 1
        for k, p in sorted(a["celkem_pct"].items(), key=lambda x: -x[1]):
            ws.cell(r, 1, k).font = BEZNY
            if vzorec:
                c = ws.cell(r, 2, f'=IFERROR(COUNTIF({rng},$A{r})/COUNTIF({rng},"?*"),0)')
                ws.cell(r, 3, f'=COUNTIF({rng},$A{r})').font = BEZNY
            else:
                c = ws.cell(r, 2, p / 100)
            c.font, c.number_format = BEZNY, "0.0%"
            r += 1
        if a.get("expected_pct"):
            r += 1
            ws.cell(r, 1, "Průměr LLM pravděpodobností (bez externího losování)").font = TUCNY
            r += 1
            _hlavicka(ws, r, ["Odpověď", "Expected (%)"]); r += 1
            for k, pexp in sorted(a["expected_pct"].items(), key=lambda x: -x[1]):
                ws.cell(r, 1, k).font = BEZNY
                c = ws.cell(r, 2, pexp / 100); c.font = BEZNY; c.number_format = "0.0%"
                r += 1
        if a.get("pozn"):
            ws.cell(r, 1, a["pozn"]).font = MALY
            r += 1

    for klic, tab in [(k, v) for k, v in a.items() if k.startswith("podle_")]:
        seg = klic.replace("podle_", "")
        r += 2
        ws.cell(r, 1, f"Podle: {seg}").font = TUCNY
        r += 1
        sloupce = (["prumer"] if a["typ"] == "skala" else volby)
        _hlavicka(ws, r, [seg, "n"] + [str(s) for s in sloupce])
        hdr_r = r
        r += 1
        vz = bool(dmap) and oid in dmap and seg in dmap and a["typ"] == "vyber"
        if vz:
            qr = f"Data!${dmap[oid]}$2:${dmap[oid]}${nrow}"
            sr = f"Data!${dmap[seg]}$2:${dmap[seg]}${nrow}"
        for i, (hod, radek) in enumerate(sorted(tab.items(), key=lambda x: -x[1].get("n", 0))):
            ws.cell(r, 1, hod).font = BEZNY
            ws.cell(r, 2, f'=COUNTIFS({sr},$A{r},{qr},"?*")' if (vz and not radek.get("suppressed"))
                    else radek.get("n", "")).font = BEZNY
            if radek.get("suppressed"):
                ws.cell(r, 3, "málo pozorování").font = MALY
                r += 1
                continue
            for j, s in enumerate(sloupce, 3):
                val = radek.get(s)
                if vz:
                    col = get_column_letter(j)
                    c = ws.cell(r, j, f'=IFERROR(COUNTIFS({sr},$A{r},{qr},{col}${hdr_r})'
                                      f'/COUNTIFS({sr},$A{r},{qr},"?*"),0)')
                else:
                    c = ws.cell(r, j, (val / 100)
                                if (val is not None and a["typ"] != "skala") else val)
                c.font = BEZNY
                c.number_format = "0.0%" if a["typ"] != "skala" else "0.00"
            if i % 2:
                for j in range(1, len(sloupce) + 3):
                    ws.cell(r, j).fill = FILL_ALT
            r += 1
    for w in a.get("varovani", []):
        r += 1
        c = ws.cell(r, 1, "Pozor: " + w)
        c.font, c.fill = MALY, FILL_VAR
    ws.freeze_panes = "A5"


def _list_verbatimy(wb: Workbook, v: dict, detail: pd.DataFrame) -> None:
    otevrene = [o for o in v.get("otazky", []) if o["typ"] == "otevrena"]
    if not otevrene:
        return
    ws = wb.create_sheet("Verbatimy")
    _sirky(ws, [10, 14, 14, 24, 22, 70])
    _hlavicka(ws, 1, ["ID otázky", "synthetic_case_id", "věková skupina", "kraj", "vzdělání", "syntetická ilustrační odpověď"])
    r = 2
    for o in otevrene:
        if o["id"] not in detail.columns:
            continue
        sub = detail[detail[o["id"]].notna()]
        for _, row in sub.iterrows():
            for j, val in enumerate([o["id"], row.get("synthetic_case_id"), row.get("vek_skupina"),
                                     row.get("kraj"), row.get("vzdelani"),
                                     row.get(o["id"])], 1):
                c = ws.cell(r, j, val)
                c.font = BEZNY
                if j == 6:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
            r += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{max(r - 1, 2)}"


def _list_qc(wb: Workbook, q: dict) -> None:
    ws = wb.create_sheet("Kontrola kvality")
    _sirky(ws, [14, 14, 70, 12])
    ws["A1"], ws["A1"].font = f"Kontrola kvality: {q['uroven']}", NADPIS
    c = ws["A2"] = q["verdikt"]
    ws["A2"].font = BEZNY
    if q["uroven"] != "OK":
        ws["A2"].fill = FILL_VAR
    r = 4
    _hlavicka(ws, r, ["Úroveň", "Kde", "Nález", "Hodnota"])
    r += 1
    for nl in sorted(q["nalezy"], key=lambda x: x["uroven"] != "KRITICKE"):
        for j, val in enumerate([nl["uroven"], nl["kde"], nl["nalez"], nl["hodnota"]], 1):
            cc = ws.cell(r, j, val)
            cc.font = BEZNY
            if nl["uroven"] == "KRITICKE":
                cc.fill = FILL_VAR
        r += 1
    if not q["nalezy"]:
        ws.cell(r, 1, "Bez nálezů.").font = BEZNY
        r += 1
    r += 2
    ws.cell(r, 1, "Souhrnné metriky").font = TUCNY
    r += 1
    m = q["metriky"]
    for k, lbl in [("chybovost", "Chybovost volání"),
                   ("prumerne_max_cramer_v", "Průměrné max Cramérovo V (demografická diferenciace)"),
                   ("podil_duplicitnich_vektoru", "Podíl duplicitních vektorů odpovědí"),
                   ("straightlining", "Straightlining na škálách")]:
        if k in m:
            ws.cell(r, 1, lbl).font = BEZNY
            ws.cell(r, 3, m[k]).font = BEZNY
            r += 1
    r += 1
    ws.cell(r, 1, "Cramérovo V měří jen demografickou diferenciaci. Není důkazem, že "
                  "LLM použilo dispoziční personu; to vyžaduje ablation benchmark.").font = MALY


def _list_segment(wb: Workbook, seg: dict) -> None:
    if not seg or seg.get("mode") in (None, "none"):
        return
    ws = wb.create_sheet("Segment")
    _sirky(ws, [30, 28, 18, 18, 18])
    ws["A1"], ws["A1"].font = f"Segment: {seg.get('name','')}", NADPIS
    rows = [
        ("Confidence class", seg.get("confidence_class")),
        ("Membership source", seg.get("membership_source")),
        ("Raw prevalence", seg.get("prevalence_raw")),
        ("Calibrated prevalence", seg.get("prevalence_calibrated")),
        ("External prevalence target", seg.get("prevalence_target")),
        ("Prevalence source", seg.get("prevalence_source")),
        ("Population ESS", seg.get("ess_population")),
        ("CV AUC", seg.get("cv_auc")),
        ("Train n / positive", f"{seg.get('n_train')} / {seg.get('n_positive')}"),
    ]
    r=3
    for k,val in rows:
        ws.cell(r,1,k).font=TUCNY; ws.cell(r,2,val).font=BEZNY; r+=1
    r+=1; _hlavicka(ws,r,["Driver","Koeficient","Směr"]); r+=1
    for x in seg.get("top_drivers",[])[:25]:
        ws.cell(r,1,x.get("feature")).font=BEZNY; ws.cell(r,2,x.get("coef")).font=BEZNY; ws.cell(r,3,x.get("direction")).font=BEZNY; r+=1
    r+=1; _hlavicka(ws,r,["Profil feature","Hodnota","Populace","Segment","Lift/effect"]); r+=1
    for x in seg.get("profile",[])[:30]:
        ws.cell(r,1,x.get("feature")).font=BEZNY; ws.cell(r,2,x.get("value","")).font=BEZNY
        ws.cell(r,3,x.get("population")).font=BEZNY; ws.cell(r,4,x.get("segment")).font=BEZNY
        ws.cell(r,5,x.get("lift",x.get("effect"))).font=BEZNY; r+=1
    r+=1
    ws.cell(r,1,"Interpretace").font=TUCNY
    cls=seg.get("confidence_class")
    note={"A":"Membership/labels jsou opřené o human/microdata evidence.",
          "B":"Tvar membership je synteticky naučený, prevalence je ukotvena externím číslem.",
          "C":"Explorační syntetický segment bez externí prevalence; nepresentovat jako změřenou populaci."}.get(cls,"Odvozený segment.")
    ws.cell(r,2,note).font=MALY


def _list_discovery(wb: Workbook, disc: dict) -> None:
    if not disc:
        return
    ws=wb.create_sheet("Audience discovery"); _sirky(ws,[34,22,20,20])
    ws["A1"],ws["A1"].font=f"Reverse discovery: {disc.get('target_question','')}",NADPIS
    ws["A2"]="Popisuje asociace v syntetických odpovědích; není to kauzální ani human-measured segment."; ws["A2"].font=MALY
    r=4; _hlavicka(ws,r,["Driver","Koeficient","Směr"]); r+=1
    for x in disc.get("top_drivers",[])[:30]:
        ws.cell(r,1,x.get("feature")).font=BEZNY; ws.cell(r,2,x.get("coef")).font=BEZNY; ws.cell(r,3,x.get("direction")).font=BEZNY; r+=1


def export_xlsx(v: dict, cesta: str | Path = "report.xlsx", vc_dat: bool = True,
                internal_data: bool = False) -> Path:
    """Klientsky export je defaultne allowlist-only.

    internal_data=True je explicitni debug/export pro metodiky; nikdy ne default.
    """
    detail: pd.DataFrame = v.get("detail", pd.DataFrame())
    wb = Workbook()
    wb.remove(wb.active)

    dmap: dict[str, str] = {}
    nrow = 0
    if vc_dat and len(detail):
        ws = wb.create_sheet("Data")
        if internal_data:
            sloupce = [c for c in detail.columns if not c.startswith("_")
                       and c != "persona_text"]
        else:
            qids = [o["id"] for o in v.get("otazky", []) if o.get("id") in detail.columns]
            if not qids and "odpoved" in detail.columns:
                qids = ["odpoved"]
            safe = ["synthetic_case_id", "pohlavi", "vek_skupina", "vzdelani",
                    "kraj", "trida_spolecenska"]
            sloupce = [c for c in safe if c in detail.columns] + qids
        dmap = {c: get_column_letter(i) for i, c in enumerate(sloupce, 1)}
        nrow = len(detail) + 1
        _hlavicka(ws, 1, sloupce)
        for i, (_, row) in enumerate(detail[sloupce].iterrows(), 2):
            for j, c in enumerate(sloupce, 1):
                ws.cell(i, j, str(row[c]) if isinstance(row[c], list) else row[c]).font = BEZNY
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(sloupce))}{nrow}"
        _sirky(ws, [16] * len(sloupce))

    _list_prehled(wb, v)
    _list_segment(wb, v.get("segment") or {})
    _list_discovery(wb, v.get("audience_discovery") or {})
    if v.get("qc"):
        _list_qc(wb, v["qc"])

    if "otazky" in v:                       # vystup run_dotaznik
        for o in v["otazky"]:
            _list_otazka(wb, o["id"], o, v["vysledky"][o["id"]], dmap, nrow)
        _list_verbatimy(wb, v, detail)
    else:                                   # vystup run_pruzkum (1 otazka)
        a = {k: val for k, val in v.items() if k.startswith("podle_")}
        a.update({"typ": "vyber", "n_platnych": v["n_platnych"],
                  "n_chybi": v["n_chyb"], "celkem_pct": v["celkem_pct"],
                  "mc95_pp": v.get("mc95_pp"), "varovani": v.get("varovani", [])})
        _list_otazka(wb, "odpoved", {"text": v["otazka"]}, a, dmap, nrow)

    if "Data" in wb.sheetnames:
        wb.move_sheet("Data", offset=len(wb.sheetnames))

    if str(v.get("mode","")).lower()=="dry":
        from openpyxl.styles import PatternFill, Font, Alignment
        warn=wb.create_sheet("INVALID_DRY_RUN",0)
        warn["A1"]="INVALID – DRY RUN"; warn["A2"]="TECHNICKÝ TEST. Tento export není výzkumný výsledek a nesmí být interpretován jako populace."
        warn["A1"].font=Font(bold=True,size=24,color="FFFFFF"); warn["A1"].fill=PatternFill("solid",fgColor="8B0000")
        warn["A2"].font=Font(bold=True,color="8B0000"); warn.column_dimensions["A"].width=100
        for ws in wb.worksheets:
            ws.oddHeader.center.text="INVALID – DRY RUN · TECHNICKÝ TEST"
            ws.evenHeader.center.text="INVALID – DRY RUN · TECHNICKÝ TEST"

    p = Path(cesta)
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(p)
    return p



def export_dataset_csv_complete(v: dict, cesta: str | Path = "dataset_complete.csv") -> Path:
    """Owner/audit respondent export for 10.12.

    Unlike the client-safe CSV, this intentionally preserves the complete respondent
    record used by the engine: stable panel id, sampling fields, D_/M_ dimensions,
    P_ deep-persona fields including media minutes, persona_text, responses, reasons
    and language QC. It is the reproducibility dataset for the research owner.
    """
    detail: pd.DataFrame = v.get("detail", pd.DataFrame())
    if detail is None or not len(detail):
        raise ValueError("Běh neobsahuje respondent-level detail; kompletní dataset nelze exportovat.")
    # Keep every actual detail column. This is a deliberate owner/audit export, not
    # the client-safe report worksheet. Object dtype and nullable columns are left
    # unchanged to avoid losing provenance or exact prompt conditioning.
    p = Path(cesta); p.parent.mkdir(parents=True, exist_ok=True)
    detail.to_csv(p, index=False, encoding="utf-8-sig")
    return p

def export_dataset_csv(v: dict, cesta: str | Path = "dataset.csv", *, internal_data: bool = False) -> Path:
    """Export the respondent-level dataset as a separate CSV.

    The default mirrors the client-safe `Data` worksheet from export_xlsx so the UI
    can offer an obvious dataset download instead of hiding it inside the workbook.
    """
    detail: pd.DataFrame = v.get("detail", pd.DataFrame())
    if detail is None or not len(detail):
        raise ValueError("Běh neobsahuje respondent-level detail; dataset nelze exportovat.")
    if internal_data:
        cols = [c for c in detail.columns if not c.startswith("_") and c != "persona_text"]
    else:
        qids = [o["id"] for o in v.get("otazky", []) if o.get("id") in detail.columns]
        if not qids and "odpoved" in detail.columns:
            qids = ["odpoved"]
        safe = ["synthetic_case_id", "pohlavi", "vek_skupina", "vzdelani", "kraj", "trida_spolecenska"]
        cols = [c for c in safe if c in detail.columns] + qids
    if not cols:
        raise ValueError("Po bezpečnostním allowlistu nezbyly žádné sloupce datasetu.")
    p = Path(cesta); p.parent.mkdir(parents=True, exist_ok=True)
    detail[cols].to_csv(p, index=False, encoding="utf-8-sig")
    return p
