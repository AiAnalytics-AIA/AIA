from __future__ import annotations
from pathlib import Path
from typing import Any

def export_workflow_report(path:str|Path,*,project:dict[str,Any],result_summary:dict[str,Any],verification:dict[str,Any]|None,interpretation:dict[str,Any]|None=None)->Path:
 from docx import Document
 p=Path(path);p.parent.mkdir(parents=True,exist_ok=True);d=Document();d.add_heading(project.get('title') or 'NPC Panel — Final Client Report',0)
 d.add_paragraph('NPC Panel 17.2.1 UX Research OS. Syntetický/modelovaný výzkum; nejde o lidský terénní sběr.')
 d.add_heading('Shrnutí',1)
 if interpretation:
  for x in interpretation.get('highlights') or []:d.add_paragraph(str(x),style='List Bullet')
 else:d.add_paragraph('Automatická interpretační vrstva nebyla spuštěna; následují přímo agregované výsledky.')
 d.add_heading('Výsledky',1)
 for qid,r in list((result_summary.get('vysledky') or {}).items())[:60]:
  d.add_heading(str(qid),2)
  if isinstance(r,dict):
   if isinstance(r.get('celkem_pct'),dict):d.add_paragraph('; '.join(f'{k}: {float(v):.1f} %' for k,v in r['celkem_pct'].items()))
   elif r.get('prumer') is not None:d.add_paragraph(f"Průměr: {r.get('prumer')}")
 d.add_heading('Externí ověření / triangulace',1)
 if verification and (verification.get('by_target') or verification.get('findings')):
  d.add_paragraph('verification_status=COMPLETED. Externí evidence poskytuje kontext a nikdy nepřepisuje NPC výsledek.')
 else:d.add_paragraph('verification_status=NOT_RUN. Tento report neobsahuje externí triangulaci a nesmí ji předstírat.')
 ds=result_summary.get('donor_support') or {}
 if ds:
  d.add_heading('Donor-aware support',1);d.add_paragraph(f"raw N: {ds.get('n','—')} · effective N: {ds.get('effective_n_combined',ds.get('effective_n_kish','—'))} · support: {ds.get('support_status','—')}")
 d.add_heading('Metodický status',1);d.add_paragraph('Externí predictive certification zůstává pending. OOS a LLM ablation status se vydáním 17.2.1 nemění.')
 d.save(p);return p
