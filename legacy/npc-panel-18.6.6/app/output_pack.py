from __future__ import annotations
import csv,html,json,time,hashlib,zipfile
from pathlib import Path
from typing import Any
def _e(x):return html.escape(str(x or ''))
def _now():return time.strftime('%Y-%m-%dT%H:%M:%SZ',time.gmtime())
def _write(p,b):p=Path(p);p.parent.mkdir(parents=True,exist_ok=True);p.write_text(b,encoding='utf-8');return str(p)
def _page(t,b,k='NPC Panel · Research OS'):return f'<!doctype html><meta charset="utf-8"><style>body{{font:15px/1.6 Inter,Arial;background:#f3f5f8;margin:0;color:#172033}}main{{max-width:1080px;margin:auto;background:white;padding:48px 60px}}h1{{font-size:34px}}h2{{margin-top:32px;border-top:1px solid #e4e7eb;padding-top:20px}}.card{{border:1px solid #e1e5ea;border-radius:12px;padding:14px;margin:10px 0}}.mut{{color:#667085}}table{{width:100%;border-collapse:collapse}}td,th{{padding:9px;border-bottom:1px solid #e5e7eb;text-align:left;vertical-align:top}}</style><main><div class="mut">{_e(k)}</div><h1>{_e(t)}</h1>{b}</main>'
def export_research_brief(path,project,research):
 a=research.get('accepted') or [];q=research.get('quarantined') or [];body=f'<p><b>{len(a)}</b> accepted context items · <b>{len(q)}</b> quarantined outcome-like items.</p><h2>Kontext</h2>'+''.join(f'<div class="card"><b>{_e(x.get("claim"))}</b><br><span class="mut">{_e(x.get("source_title"))}</span></div>' for x in a)+'<h2>Anti-leakage karanténa</h2>'+''.join(f'<div class="card">{_e(x.get("claim") or x.get("source_title"))}</div>' for x in q[:30]);return _write(path,_page((project.get('title') or 'Výzkum')+' · Research Brief',body))
def export_analysis_readout(path,project,analysis):
 body=f'<p>{_e(analysis.get("executive_answer"))}</p><h2>Findings</h2>'+''.join(f'<div class="card"><h3>{_e(x.get("headline"))}</h3><p>{_e(x.get("finding"))}</p><p><b>Meaning:</b> {_e(x.get("meaning"))}</p></div>' for x in analysis.get('key_findings') or []);return _write(path,_page((project.get('title') or 'Výzkum')+' · Analysis',body))
def export_validation_readout(path,project,verification):
 f=verification.get('findings') or [];body='<p>Externí evidence je post-result triangulace a nepřepisuje NPC výsledek.</p><table><tr><th>Claim</th><th>Direction</th><th>Source</th></tr>'+''.join(f'<tr><td>{_e(x.get("claim_summary"))}</td><td>{_e(x.get("direction"))}</td><td>{_e(x.get("source_title"))}</td></tr>' for x in f)+'</table>';return _write(path,_page((project.get('title') or 'Výzkum')+' · Validation',body))
def export_executive_readout(path,project,report,maps=None):
 body=f'<p style="font-size:21px"><b>{_e(report.get("executive_summary"))}</b></p><h2>Decision answer</h2><p>{_e(report.get("decision_answer"))}</p><h2>Findings</h2>'+''.join(f'<div class="card"><h3>{_e(x.get("headline"))}</h3><p>{_e(x.get("meaning"))}</p></div>' for x in report.get('key_findings') or [])+'<h2>Doporučení</h2><ol>'+''.join(f'<li><b>{_e(x.get("action"))}</b> — {_e(x.get("why"))}</li>' for x in report.get('implications') or [])+'</ol>';return _write(path,_page((project.get('title') or 'Výzkum')+' · Executive Readout',body))
def export_questionnaire_html(path,project):
 qs=[]
 for sec in project.get('sections') or []:qs.extend(sec.get('questions') or [])
 if not qs:qs=project.get('questions') or []
 body='<table><tr><th>ID</th><th>Otázka</th><th>Typ</th></tr>'+''.join(f'<tr><td>{_e(q.get("id"))}</td><td>{_e(q.get("text") or q.get("otazka"))}</td><td>{_e(q.get("type") or q.get("typ"))}</td></tr>' for q in qs)+'</table>';return _write(path,_page((project.get('title') or 'Výzkum')+' · Dotazník',body))
def export_internal_report(path,project,report,analysis,research,verification,alignment,workflow_id=''):
 meta=report.get('_meta') or {};body=f'<p><b>Workflow:</b> {_e(workflow_id)}</p><p><b>Provider/model:</b> {_e(meta.get("provider"))} · {_e(meta.get("model"))}</p><h2>Report QA</h2><pre>{_e(json.dumps(meta.get("quality_gate") or {},ensure_ascii=False,indent=2))}</pre><h2>Evidence validation</h2><pre>{_e(json.dumps(meta.get("report_evidence_validation") or {},ensure_ascii=False,indent=2))}</pre><h2>Governance</h2><p>External predictive validation: NOT_VALIDATED / PENDING.</p>';return _write(path,_page((project.get('title') or 'Výzkum')+' · Internal report',body,k='NPC Panel · Internal'))
def export_handoff_documentation(path,project,report,files):
 body='<p>Klientský balík je oddělený od interní metodiky. Začněte Executive Readoutem, reportem nebo management deckem.</p><h2>Artefakty</h2><ul>'+''.join(f'<li><b>{_e(k)}</b>: {_e(v)}</li>' for k,v in files.items() if v)+'</ul>';return _write(path,_page((project.get('title') or 'Výzkum')+' · Předávací dokumentace',body))
def export_handoff_protocol(path,project,report,workflow_id=''):
 body=f'<p>Workflow {_e(workflow_id)}.</p><ul><li>Client report generated through evidence gate.</li><li>External predictive validation status remains pending.</li><li>Final human sign-off required before external send.</li></ul>';return _write(path,_page((project.get('title') or 'Výzkum')+' · Předávací protokol',body))
def export_management_deck_html(path,project,report,maps=None):
 slides=[('Executive answer',report.get('decision_answer') or report.get('executive_summary'))]+[(x.get('headline'),x.get('meaning')) for x in report.get('key_findings') or []]+[('Doporučení','; '.join(str(x.get('action')) for x in report.get('implications') or []))];body=''.join(f'<section style="min-height:440px;border-bottom:1px solid #ddd;padding:28px 0"><h2>{_e(a)}</h2><p style="font-size:22px">{_e(b)}</p></section>' for a,b in slides);return _write(path,_page((project.get('title') or 'Výzkum')+' · Management Deck',body))
def export_management_deck_pptx(path,project,report,maps=None):
 from pptx import Presentation
 from pptx.util import Inches,Pt
 from pptx.dml.color import RGBColor
 from pptx.enum.shapes import MSO_SHAPE
 p=Path(path);p.parent.mkdir(parents=True,exist_ok=True);prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5);INK=RGBColor(24,28,35);ACC=RGBColor(23,58,120);MUT=RGBColor(102,112,128)
 def slide(title,text,kicker='NPC PANEL'):
  sl=prs.slides.add_slide(prs.slide_layouts[6]);tb=sl.shapes.add_textbox(Inches(.8),Inches(.6),Inches(11.7),Inches(.35));tf=tb.text_frame;tf.text=kicker;r=tf.paragraphs[0].runs[0];r.font.size=Pt(10);r.font.bold=True;r.font.color.rgb=MUT;tb=sl.shapes.add_textbox(Inches(.8),Inches(1.15),Inches(11.6),Inches(1.3));tf=tb.text_frame;tf.word_wrap=True;tf.text=str(title or '');r=tf.paragraphs[0].runs[0];r.font.size=Pt(28);r.font.bold=True;r.font.color.rgb=INK;tb=sl.shapes.add_textbox(Inches(.85),Inches(2.8),Inches(11.3),Inches(2.7));tf=tb.text_frame;tf.word_wrap=True;tf.text=str(text or '');r=tf.paragraphs[0].runs[0];r.font.size=Pt(19);r.font.color.rgb=INK;return sl
 slide(project.get('title') or report.get('title'),report.get('decision_answer') or report.get('executive_summary'),'CLIENT RESEARCH READOUT')
 slide('Co potřebujete vědět',report.get('executive_summary'))
 for i,x in enumerate((report.get('key_findings') or [])[:7],1):slide(x.get('headline'),str(x.get('finding') or '')+'\n\nSO WHAT: '+str(x.get('meaning') or ''),f'FINDING {i}')
 slide('Co doporučujeme udělat teď','\n'.join(f'{i}. {x.get("action")} — {x.get("why")}' for i,x in enumerate((report.get('implications') or [])[:5],1)),'ACTION')
 slide('Jak silně závěrům věřit',str(report.get('confidence_summary') or '')+'\n\n'+'\n'.join('• '+str(x) for x in report.get('limitations') or []),'CONFIDENCE')
 prs.save(p);return str(p)
def export_evidence_pack(path,project,analysis,research,verification,report):
 obj={'kind':'npc_project_evidence_pack_v1','generated_at':_now(),'project_title':project.get('title'),'research_background':research.get('accepted') or [],'quarantined_outcome_benchmarks':research.get('quarantined') or [],'post_result_validation':verification.get('findings') or [],'analysis_evidence_validation':analysis.get('_meta',{}).get('evidence_validation') or {},'report_evidence_validation':report.get('_meta',{}).get('report_evidence_validation') or {},'report_partner_review':report.get('_meta',{}).get('partner_review') or {},'validation_status':'NOT_VALIDATED'};p=Path(path);p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(obj,ensure_ascii=False,indent=2,default=str),encoding='utf-8');return str(p)
def export_verbatims(csv_path,html_path,project,summary):
 rows=[]
 for qid,r in (summary.get('vysledky') or {}).items():
  for i,t in enumerate((r or {}).get('verbatimy') or [],1):rows.append({'question_id':qid,'verbatim_no':i,'verbatim':str(t)})
 pc=Path(csv_path);pc.parent.mkdir(parents=True,exist_ok=True)
 with pc.open('w',encoding='utf-8-sig',newline='') as f:w=csv.DictWriter(f,fieldnames=['question_id','verbatim_no','verbatim']);w.writeheader();w.writerows(rows)
 body=''.join(f'<div class="card"><b>{_e(x["question_id"])}</b><p>{_e(x["verbatim"])}</p></div>' for x in rows) or '<p>Bez otevřených odpovědí.</p>';_write(html_path,_page((project.get('title') or 'Výzkum')+' · Verbatimy',body));return {'csv':str(pc),'html':str(html_path),'count':len(rows)}
def build_output_manifest(path,files,*,project,workflow_id):
 obj={'workflow_id':workflow_id,'project_title':project.get('title'),'generated_at':_now(),'outputs':files,'method_status':'FINAL_ENGINEERING_RELEASE_EXTERNAL_PREDICTIVE_CERTIFICATION_PENDING'};p=Path(path);p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(obj,ensure_ascii=False,indent=2,default=str),encoding='utf-8');return str(p)


def export_client_delivery_zip(path, files, *, project=None, workflow_id=''):
 """Create one client-safe delivery ZIP with per-file SHA-256 checksums.

 Only existing files explicitly supplied by the final-report stage are included.
 Internal runtime databases, hidden persona columns and provider logs are never swept
 implicitly into the archive.
 """
 p=Path(path);p.parent.mkdir(parents=True,exist_ok=True)
 selected=[]
 for label,value in (files or {}).items():
  if not value: continue
  q=Path(str(value))
  if not q.is_file(): continue
  selected.append((str(label),q))
 checks=[]
 with zipfile.ZipFile(p,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=6) as z:
  for label,q in selected:
   arc=q.name
   z.write(q,arc)
   h=hashlib.sha256(q.read_bytes()).hexdigest();checks.append(f'{h}  {arc}')
  meta={'kind':'npc_client_delivery_v1','workflow_id':workflow_id,'project_title':(project or {}).get('title'),'generated_at':_now(),'file_count':len(selected)}
  z.writestr('DELIVERY_METADATA.json',json.dumps(meta,ensure_ascii=False,indent=2))
  z.writestr('DELIVERY_SHA256SUMS.txt','\n'.join(checks)+'\n')
 return str(p)
