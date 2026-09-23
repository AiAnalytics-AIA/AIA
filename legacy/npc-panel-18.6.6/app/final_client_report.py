"""Final client report composer — requires a completed external verification bundle."""
from __future__ import annotations
from pathlib import Path
from typing import Any
import json


def _table(doc, headers, rows):
    t=doc.add_table(rows=1,cols=len(headers)); t.style="Light Grid Accent 1"
    for i,h in enumerate(headers):
        t.rows[0].cells[i].text=str(h)
        for r in t.rows[0].cells[i].paragraphs[0].runs:r.bold=True
    for row in rows:
        c=t.add_row().cells
        for i,v in enumerate(row): c[i].text="" if v is None else str(v)


def export_final_client_report(path:str|Path, *, project:dict[str,Any], result_summary:dict[str,Any], verification:dict[str,Any], support:dict[str,Any]|None=None)->Path:
    if not verification or not (verification.get("by_target") or verification.get("findings")):
        raise ValueError("Final Client Report se vytváří až po externím ověření výsledků.")
    from docx import Document
    p=Path(path);p.parent.mkdir(parents=True,exist_ok=True)
    d=Document(); title=str(project.get("title") or result_summary.get("nazev") or "NPC Panel — Final Client Report")
    d.add_heading(title,0)
    d.add_paragraph("FINAL CLIENT REPORT · syntetický respondentní panel · externí triangulace zahrnuta")
    d.add_heading("Executive summary",1)
    d.add_paragraph(str((project.get("research_plan") or {}).get("problem_summary") or project.get("goal") or "Výsledky syntetické studie jsou interpretovány společně s dostupným externím kontextem."))
    d.add_paragraph("Externí shoda sama o sobě není důkaz prediktivní validity. NPC výsledek se externím srovnáním nepřepisuje.")
    d.add_heading("Hlavní výsledky",1)
    rows=[]
    for qid,r in (result_summary.get("vysledky") or {}).items():
        if not isinstance(r,dict):continue
        if r.get("celkem_pct"):
            top=sorted((r.get("celkem_pct") or {}).items(), key=lambda kv:-float(kv[1]))[:3]
            rows.append([qid,"; ".join(f"{k}: {float(v):.1f}%" for k,v in top)])
        elif r.get("prumer") is not None: rows.append([qid,f"průměr {float(r['prumer']):.2f}"])
    if rows:_table(d,["Otázka","Výsledek"],rows)
    else:d.add_paragraph("Výsledek nemá standardní číselné agregace vhodné pro tabulku.")
    d.add_heading("Externí kontext / triangulace",1)
    vr=[]
    for tid,x in (verification.get("by_target") or {}).items():
        t=x.get("target") or {}; best=x.get("best_direct") or {}
        vr.append([t.get("question_text") or tid,t.get("npc_value","—"),x.get("status","—"),best.get("benchmark_value","—"),best.get("source_title","—"),best.get("rationale","—")])
    if vr:_table(d,["Výsledek","NPC","Srovnatelnost","Benchmark","Zdroj","Interpretace"],vr[:30])
    else:d.add_paragraph("Nebyl nalezen přímo srovnatelný číselný benchmark; dostupná evidence je pouze kontextová.")
    if support:
        d.add_heading("Efektivní podpora vzorku",1)
        _table(d,["raw n","Kish n","core donoři","layer","layer donoři","effective n","support"],[[support.get("n"),support.get("effective_n_kish"),support.get("n_unique_core_donors"),support.get("donor_layer"),support.get("n_unique_layer_donors"),support.get("effective_n_combined"),support.get("support_status")]])
    d.add_heading("Metodické hranice",1)
    for x in [
        "Jde o syntetickou simulaci, nikoli lidský terénní sběr.",
        "Raw syntetické N se nesmí zaměňovat za počet nezávislých reálných donorů.",
        "Modelované behaviorální priory nejsou v LIVE promptu defaultně aktivní bez úspěšné A/B/C/D ablace.",
        "Externí triangulace výsledek nepřepisuje a není certifikací ekvivalence s lidským výzkumem.",
    ]:d.add_paragraph(x,style="List Bullet")
    d.add_heading("Audit",1)
    d.add_paragraph(f"Verification SHA-256: {verification.get('sha256','—')}")
    d.add_paragraph(f"Run ID: {result_summary.get('run_id','—')}")
    d.save(p);return p
