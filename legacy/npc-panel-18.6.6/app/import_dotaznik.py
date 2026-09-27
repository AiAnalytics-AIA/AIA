"""
NPC PANEL — import_dotaznik.py
Nahraje HOTOVY dotaznik (Word, Excel, CSV, PDF, txt) a prevede ho na brief.json,
ktery jde primo do run.py. Na rozdil od navrh.py (ktery otazky VYMYSLI ze
zadani) tenhle modul otazky POUZE PREPISE do strukturovaneho tvaru — nesmi
nic pridat ani zmenit smysl.

    python import_dotaznik.py muj_dotaznik.docx -o brief.json
    python run.py brief.json --dry

DVA REZIMY PARSOVANI
  1. STRUKTUROVANY (xlsx/csv s rozpoznatelnymi sloupci: otazka, typ, moznosti...)
     -> primy prevod bez LLM. Levne, presne, zadne riziko halucinace.
  2. VOLNY TEXT (docx, pdf, txt, nebo xlsx/csv bez jasne struktury)
     -> LLM extrahuje otazky z textu. System prompt EXPLICITNE zakazuje
     vymyslet cokoli navic — jen prepsat, co uz v dokumentu je.
"""

from __future__ import annotations
from provider_auth import get_ai_provider

from provider_auth import create_anthropic_client

import argparse
import json
import re
import sys
from pathlib import Path

from runtime_config import MODELS, resolve_model, RUN_DEFAULTS
from anthropic_compat import create_message

import pandas as pd

from navrh import POVOLENE, _validuj

# ---------------------------------------------------------------- cteni souboru

def _cti_docx(cesta: Path) -> str:
    import docx
    d = docx.Document(str(cesta))
    casti = [p.text for p in d.paragraphs if p.text.strip()]
    for tab in d.tables:
        for row in tab.rows:
            casti.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(casti)


def _cti_pdf(cesta: Path) -> str:
    import pypdf
    r = pypdf.PdfReader(str(cesta))
    return "\n".join(p.extract_text() or "" for p in r.pages)


def _cti_tabulku(cesta: Path) -> pd.DataFrame:
    if cesta.suffix.lower() == ".csv":
        return pd.read_csv(cesta)
    if cesta.suffix.lower() in (".xlsx", ".xls"):
        return pd.read_excel(cesta)
    raise ValueError(f"Neumim precist tabulku: {cesta.suffix}")


def nacti(cesta: str | Path) -> tuple[str, pd.DataFrame | None]:
    """Vrati (surovy_text, tabulka_nebo_None). Tabulka jen kdyz je xlsx/csv."""
    p = Path(cesta)
    suf = p.suffix.lower()
    if suf == ".docx":
        return _cti_docx(p), None
    if suf == ".pdf":
        return _cti_pdf(p), None
    if suf in (".txt", ".md"):
        return p.read_text(encoding="utf-8", errors="ignore"), None
    if suf in (".csv", ".xlsx", ".xls"):
        df = _cti_tabulku(p)
        return df.to_csv(index=False), df
    raise ValueError(f"Nepodporovana pripona: {suf}. "
                     "Podporovano: .docx .pdf .txt .md .csv .xlsx .xls")


# ---------------------------------------------------------------- strukturovany rezim

# alternativni nazvy sloupcu, ktere rozpoznam bez LLM
ALIASY = {
    "text": {"text", "otazka", "question", "znění", "zneni", "otázka"},
    "typ": {"typ", "type"},
    "kategorie": {"kategorie", "moznosti", "možnosti", "options", "odpovedi", "odpovědi"},
    "skala_min": {"skala_min", "min", "škála_min"},
    "skala_max": {"skala_max", "max", "škála_max"},
    "id": {"id", "kod", "kód"},
}


def _najdi_sloupec(df: pd.DataFrame, klic: str) -> str | None:
    aliasy = ALIASY[klic]
    for c in df.columns:
        if str(c).strip().lower() in aliasy:
            return c
    return None


def zkus_strukturovany_prevod(df: pd.DataFrame) -> list[dict] | None:
    """Pokud tabulka ma rozpoznatelny sloupec s textem otazky, prevede primo.

    Vraci None, kdyz struktura neni jasna — pak se pouzije LLM rezim.
    """
    col_text = _najdi_sloupec(df, "text")
    if col_text is None:
        return None
    col_typ = _najdi_sloupec(df, "typ")
    col_kat = _najdi_sloupec(df, "kategorie")
    col_id = _najdi_sloupec(df, "id")

    otazky = []
    for i, row in df.iterrows():
        text = str(row[col_text]).strip()
        if not text or text.lower() == "nan":
            continue
        typ_raw = str(row[col_typ]).strip().lower() if col_typ else ""
        typ = {"vyber": "vyber", "single": "vyber", "výběr": "vyber",
              "multi": "multi", "multiple": "multi", "více": "multi",
              "skala": "skala", "škála": "skala", "scale": "skala",
              "otevrena": "otevrena", "otevřená": "otevrena", "open": "otevrena",
              "text": "otevrena"}.get(typ_raw, "vyber")

        o: dict = {"id": str(row[col_id]).strip() if col_id else f"O{i+1}",
                  "text": text, "typ": typ}
        if typ in ("vyber", "multi") and col_kat:
            surove = str(row[col_kat])
            # oddelovac ; nebo | nebo ,
            for sep in (";", "|", ","):
                if sep in surove:
                    o["kategorie"] = [x.strip() for x in surove.split(sep) if x.strip()]
                    break
            else:
                o["kategorie"] = [surove.strip()] if surove.strip() else []
            if len(o.get("kategorie", [])) < 2:
                return None  # nejednoznacne, radsi LLM
        elif typ == "skala":
            mn = _najdi_sloupec(df, "skala_min")
            mx = _najdi_sloupec(df, "skala_max")
            o["skala"] = [int(row[mn]) if mn and pd.notna(row[mn]) else 1,
                         int(row[mx]) if mx and pd.notna(row[mx]) else 10]
        otazky.append(o)

    return otazky if otazky else None


# ---------------------------------------------------------------- LLM rezim (volny text)

SYSTEM_IMPORT = """Jsi asistent, který přepisuje existující dotazník do strukturovaného
formátu. NEJSI autor — nic nevymýšlíš, nic nepřidáváš, neopravuješ formulace ani
nezlepšuješ metodiku. Tvým jediným úkolem je věrně přepsat otázky, které jsou
v dokumentu už napsané.

Pravidla:
- Zachovej přesné znění otázek, jen odstraň číslování a formátovací balast.
- Rozpoznej typ otázky ze způsobu odpovídání v dokumentu:
  "vyber" — jedna z nabízených možností
  "multi" — lze zaškrtnout více možností
  "skala" — číselná škála (1-5, 1-10, Likert)
  "otevrena" — volná textová odpověď
- Pokud dokument u škály neuvádí čísla, ale slovní stupně (rozhodně souhlasím...
  rozhodně nesouhlasím), převeď na typ "skala" s odpovídajícím rozsahem a
  popisky_skaly = [nejnižší stupeň, nejvyšší stupeň].
- Instrukce, hlavičky, souhlas s GDPR, kontaktní údaje respondenta NEJSOU otázky
  pro panel — vynech je.
- Ke každé přepsané otázce přidej pole "topics": 1–4 tematické štítky pouze z tohoto seznamu:
  cena, potraviny, retail, nakup, online, doruceni, znacka, reklama, media, zdravi, sport, eko, obaly, energie, doprava, auto, bydleni, finance, politika, hodnoty, verejne, regulace, sluzby, predplatne, premium, moda, prace, volny_cas, B2B, loajalita, technologie, dovednosti, socialni_site, vztahy.
  Topics jsou metadata; nesmějí změnit znění ani smysl otázky.
- Pokud je dokument nejednoznačný nebo poškozený, přepiš jen to, čemu rozumíš
  s jistotou, a do "poznamka_metodika" napiš, co jsi musel vynechat nebo kde
  sis nebyl jistý.

Odpovíš výhradně jedním JSON objektem, bez markdownu a bez komentáře:
{"nazev": "...", "otazky": [ {...}, ... ], "poznamka_metodika": "co bylo nejasné nebo vynechané"}"""

UZIV_IMPORT = """DOKUMENT K PŘEPISU:
{text}

Odpověz JSON se strukturovanými otázkami z tohoto dokumentu."""


def parsuj_llm(text: str, model: str = MODELS["sonnet"].model_id, provider: str | None = None) -> dict:
    model = resolve_model(model)
    # dlouhe dokumenty oriznout, aby se vesly do kontextu s rezervou na odpoved
    text_orez = text[:40000]
    tool = {
        "name": "submit_imported_questionnaire",
        "description": "Submit the questionnaire reconstructed from the supplied document.",
        "input_schema": {
            "type": "object",
            "properties": {
                "nazev": {"type": "string"},
                "poznamka_metodika": {"type": "string"},
                "otazky": {
                    "type": "array", "minItems": 1,
                    "items": {
                        "type": "object",
                        "properties": {
                            "id": {"type": "string"}, "text": {"type": "string"},
                            "typ": {"type": "string", "enum": ["vyber", "multi", "skala", "otevrena"]},
                            "kategorie": {"type": "array", "items": {"type": "string"}},
                            "skala": {"type": "array", "minItems": 2, "maxItems": 2, "items": {"type": "integer"}},
                            "popisky_skaly": {"type": "array", "minItems": 2, "maxItems": 2, "items": {"type": "string"}},
                            "max_slov": {"type": "integer", "minimum": 5, "maximum": 100},
                            "povolit_nevim": {"type": "boolean"},
                            "filtr": {"type": "string"},
                            "topics": {"type": "array", "minItems": 1, "maxItems": 4,
                                       "items": {"type": "string", "enum": ["cena","potraviny","retail","nakup","online","doruceni","znacka","reklama","media","zdravi","sport","eko","obaly","energie","doprava","auto","bydleni","finance","politika","hodnoty","verejne","regulace","sluzby","predplatne","premium","moda","prace","volny_cas","B2B","loajalita","technologie","dovednosti","socialni_site","vztahy"]}}
                        },
                        "required": ["id", "text", "typ", "topics"],
                        "additionalProperties": False
                    }
                }
            },
            "required": ["nazev", "otazky"],
            "additionalProperties": False
        }
    }
    from ai_router import call_structured
    rr=call_structured(system=SYSTEM_IMPORT,messages=[{"role":"user","content":UZIV_IMPORT.format(text=text_orez)}],
                       schema=tool["input_schema"],schema_name=tool["name"],anthropic_model=model,max_tokens=4000,prefer=(provider or get_ai_provider()),allow_fallback=False)
    out=dict(rr["data"]); out["_ai"]={"provider":rr.get("provider"),"model":rr.get("model"),"fallback_used":rr.get("fallback_used",False)}
    return out


# ---------------------------------------------------------------- hlavni funkce

def importuj_dotaznik(
    cesta: str | Path,
    *,
    nazev: str | None = None,
    n: int = 500,
    mode: str = RUN_DEFAULTS["mode"],
    model: str = RUN_DEFAULTS["model"],
    vystup_xlsx: str | None = None,
    provider: str | None = None,
) -> dict:
    """Nacte soubor a vrati brief pripraveny pro run.py."""
    text, df = nacti(cesta)
    cesta = Path(cesta)

    otazky = None
    zpusob = "LLM (volny text)"
    if df is not None:
        otazky = zkus_strukturovany_prevod(df)
        if otazky is not None:
            zpusob = "strukturovana tabulka (bez LLM)"
    if otazky is None:
        brief = parsuj_llm(text, model, provider=provider)
        otazky = brief.get("otazky", [])
        poznamka = brief.get("poznamka_metodika")
        nazev = nazev or brief.get("nazev")
    else:
        poznamka = None

    brief = {"otazky": otazky}
    brief = _validuj(brief)  # sdilena validace s navrh.py
    nazev = nazev or cesta.stem
    brief.update({
        "nazev": nazev, "n": n, "mode": mode, "seed": 42, "filtry": {},
        "model": resolve_model(model), "response_mode": RUN_DEFAULTS["response_mode"], "persona_mode": RUN_DEFAULTS["persona_mode"],
        "vystup": vystup_xlsx or f"vystupy/{re.sub(r'[^a-zA-Z0-9]+', '_', nazev)[:40]}.xlsx",
        "_import_zdroj": str(cesta), "_import_zpusob": zpusob,
    })
    if poznamka:
        brief["poznamka_metodika"] = poznamka
    return brief


def main() -> int:
    ap = argparse.ArgumentParser(description="Import existujiciho dotazniku -> brief.json")
    ap.add_argument("soubor", help=".docx / .pdf / .txt / .csv / .xlsx")
    ap.add_argument("-o", "--out", default="brief.json")
    ap.add_argument("-n", type=int, default=500)
    ap.add_argument("--model", default=MODELS["sonnet"].model_id)
    a = ap.parse_args()

    brief = importuj_dotaznik(a.soubor, n=a.n, model=a.model)
    Path(a.out).write_text(json.dumps(brief, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"[import] {a.soubor} -> {len(brief['otazky'])} otazek "
         f"({brief['_import_zpusob']})")
    for o in brief["otazky"]:
        print(f"  [{o['id']}] ({o['typ']}) {o['text'][:70]}")
    if brief.get("poznamka_metodika"):
        print(f"\n  POZOR: {brief['poznamka_metodika']}")
    print(f"\n[ulozeno] {a.out}")
    print(f"Zkontroluj otazky v {a.out}, pak: python run.py {a.out} --dry")
    return 0


if __name__ == "__main__":
    sys.exit(main())
