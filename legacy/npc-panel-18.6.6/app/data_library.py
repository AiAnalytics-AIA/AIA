from __future__ import annotations

import hashlib
import io
import json
import re
import sqlite3
import time
from pathlib import Path
from typing import Any

ROOT=Path(__file__).resolve().parent
DB_PATH=ROOT/'data'/'knowledge_library.sqlite'
FILES_DIR=ROOT/'data'/'knowledge_library_files'
OVERLAY_PATH=ROOT/'data'/'dimension_library_overlays.json'

SOURCE_TYPES={
    'research':'Výzkum',
    'scientific_article':'Vědecký článek',
    'thesis':'Diplomová / disertační práce',
    'report':'Odborný report',
    'dataset':'Dataset / tabulka',
    'web_research':'Deep Research zdroj',
    'workflow_evidence':'Evidence z dokončeného projektu',
    'other':'Jiný zdroj',
}

DIMENSION_HINTS={
    'media':['média','media','televiz','zpravodaj','internet','sociální sítě','social media'],
    'nakup':['nákup','nakup','shopping','purchase','spotřebitel','consumer'],
    'finance':['finance','příjem','prijem','dluh','hypoték','úvěr','uver','spořen','investic'],
    'cena':['cena','price','cenová citlivost','value for money'],
    'hodnoty':['hodnot','values','morál','moral','svoboda','solidarit'],
    'duvera':['důvěr','duver','trust','instituc'],
    'politika':['politik','volb','stran','ideolog'],
    'ekologie':['ekolog','environment','klima','recykl','udržitel'],
    'technologie':['technolog','digital','ai','umělá inteligence','umela inteligence'],
    'znacka':['značk','znack','brand'],
    'reklama':['reklam','advertis','kampan'],
    'zdravi':['zdrav','health','diabet','celiak','mental'],
    'prace':['práce','prace','employment','zaměstn','zamestn','profes'],
    'rodina':['rodin','děti','deti','household','domácnost','domacnost'],
    'vztahy':['vztah','partner','relationship'],
    'socialni_site':['tiktok','instagram','facebook','youtube','social media','sociální sí'],
    'online':['online','e-commerce','eshop','e-shop'],
}


def _now()->str:
    return time.strftime('%Y-%m-%dT%H:%M:%SZ',time.gmtime())


def _connect():
    DB_PATH.parent.mkdir(parents=True,exist_ok=True)
    cx=sqlite3.connect(DB_PATH)
    cx.row_factory=sqlite3.Row
    cx.executescript('''
    CREATE TABLE IF NOT EXISTS library_entries(
      entry_id TEXT PRIMARY KEY,
      created_at TEXT NOT NULL,
      source_type TEXT NOT NULL,
      title TEXT NOT NULL,
      year TEXT,
      author TEXT,
      source_url TEXT,
      filename TEXT,
      stored_path TEXT,
      sha256 TEXT,
      notes TEXT,
      text_excerpt TEXT,
      summary TEXT,
      topics_json TEXT NOT NULL DEFAULT '[]',
      dimensions_json TEXT NOT NULL DEFAULT '[]',
      status TEXT NOT NULL DEFAULT 'IMPORTED',
      analysis_status TEXT NOT NULL DEFAULT 'NOT_ANALYZED',
      provenance_json TEXT NOT NULL DEFAULT '{}'
    );
    CREATE TABLE IF NOT EXISTS dimension_proposals(
      proposal_id TEXT PRIMARY KEY,
      entry_id TEXT,
      created_at TEXT NOT NULL,
      dimension_id TEXT NOT NULL,
      dimension_label TEXT NOT NULL,
      action TEXT NOT NULL,
      rationale TEXT NOT NULL,
      evidence_summary TEXT,
      confidence REAL,
      status TEXT NOT NULL DEFAULT 'PROPOSED',
      applied_at TEXT,
      FOREIGN KEY(entry_id) REFERENCES library_entries(entry_id)
    );
    CREATE TABLE IF NOT EXISTS proposal_specs(
      proposal_id TEXT PRIMARY KEY,
      origin TEXT NOT NULL DEFAULT 'source',
      runtime_status TEXT NOT NULL DEFAULT 'KNOWLEDGE_ONLY',
      spec_json TEXT NOT NULL DEFAULT '{}',
      FOREIGN KEY(proposal_id) REFERENCES dimension_proposals(proposal_id)
    );
    CREATE TABLE IF NOT EXISTS learning_cycles(
      cycle_id TEXT PRIMARY KEY,
      workflow_id TEXT UNIQUE,
      created_at TEXT NOT NULL,
      project_title TEXT,
      status TEXT NOT NULL,
      summary TEXT,
      gaps_json TEXT NOT NULL DEFAULT '[]',
      research_topics_json TEXT NOT NULL DEFAULT '[]',
      proposal_ids_json TEXT NOT NULL DEFAULT '[]',
      evidence_count INTEGER NOT NULL DEFAULT 0
    );
    CREATE TABLE IF NOT EXISTS calibration_history(
      event_id TEXT PRIMARY KEY,
      proposal_id TEXT,
      dimension_id TEXT,
      created_at TEXT NOT NULL,
      before_json TEXT NOT NULL DEFAULT '{}',
      after_json TEXT NOT NULL DEFAULT '{}',
      FOREIGN KEY(proposal_id) REFERENCES dimension_proposals(proposal_id)
    );
    CREATE TABLE IF NOT EXISTS library_project_links(
      project_id TEXT NOT NULL,
      source_ref TEXT NOT NULL,
      source_kind TEXT NOT NULL,
      usage_role TEXT NOT NULL DEFAULT 'context',
      created_at TEXT NOT NULL,
      PRIMARY KEY(project_id,source_ref,source_kind)
    );
    ''')
    cx.commit(); return cx


def _id(prefix:str,payload:str)->str:
    return prefix+'-'+hashlib.sha256((payload+str(time.time_ns())).encode()).hexdigest()[:14]


def _safe_filename(name:str)->str:
    name=Path(str(name or 'source')).name
    return re.sub(r'[^A-Za-z0-9._ -]+','_',name)[:160] or 'source'


def _extract_text(raw:bytes,filename:str)->str:
    suffix=Path(filename).suffix.lower()
    try:
        if suffix in {'.txt','.md','.csv','.json','.tsv'}:
            return raw.decode('utf-8',errors='replace')[:250_000]
        if suffix=='.pdf':
            from pypdf import PdfReader
            reader=PdfReader(io.BytesIO(raw))
            parts=[]
            for page in reader.pages[:80]:
                try: parts.append(page.extract_text() or '')
                except Exception: pass
            return '\n'.join(parts)[:250_000]
        if suffix=='.docx':
            from docx import Document
            doc=Document(io.BytesIO(raw))
            return '\n'.join(p.text for p in doc.paragraphs)[:250_000]
        if suffix in {'.xlsx','.xlsm'}:
            from openpyxl import load_workbook
            wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
            out=[]
            for ws in wb.worksheets[:12]:
                out.append(f'### {ws.title}')
                for row in ws.iter_rows(min_row=1,max_row=min(ws.max_row,300),values_only=True):
                    out.append(' | '.join('' if v is None else str(v) for v in row[:30]))
            return '\n'.join(out)[:250_000]
    except Exception as exc:
        return f'[TEXT_EXTRACTION_FAILED: {exc}]'
    return ''


def _hint_dimensions(text:str)->list[str]:
    t=str(text or '').lower()
    scored=[]
    for dim,terms in DIMENSION_HINTS.items():
        score=sum(t.count(x.lower()) for x in terms)
        if score: scored.append((score,dim))
    return [d for _,d in sorted(scored,reverse=True)[:8]]


def _runtime_status_for_action(action:str)->str:
    action=str(action or '')
    if action=='calibration_update': return 'CALIBRATION_CANDIDATE'
    if action in {'refine_dimension','relationship_update'}: return 'KNOWLEDGE_CONTEXT'
    if action=='new_dimension': return 'NEEDS_QUANT_ANCHOR'
    return 'KNOWLEDGE_ONLY'


def _save_proposal_spec(proposal_id:str, *, origin:str='source', action:str='', spec:dict[str,Any]|None=None)->None:
    cx=_connect()
    try:
        cx.execute('INSERT OR REPLACE INTO proposal_specs(proposal_id,origin,runtime_status,spec_json) VALUES(?,?,?,?)',(
            proposal_id,str(origin or 'source')[:80],_runtime_status_for_action(action),json.dumps(spec or {},ensure_ascii=False,default=str)))
        cx.commit()
    finally: cx.close()


def _proposal_spec(proposal_id:str)->dict[str,Any]:
    cx=_connect()
    try:r=cx.execute('SELECT * FROM proposal_specs WHERE proposal_id=?',(proposal_id,)).fetchone()
    finally:cx.close()
    if not r:return {'origin':'legacy','runtime_status':'KNOWLEDGE_ONLY','spec':{}}
    d=dict(r)
    try:d['spec']=json.loads(d.pop('spec_json') or '{}')
    except Exception:d['spec']={}
    return d


def knowledge_context_for_project(project:dict[str,Any], *, limit:int=12)->dict[str,Any]:
    """Return approved external knowledge relevant to this project.

    Population-level evidence is context, never an invented individual attribute.
    A new dimension therefore stays knowledge-only until it has a measured/calibrated
    respondent-level implementation.
    """
    dims=active_dimensions(); approved=set(str(x).lower() for x in ((project.get('persona_dimensions') or {}).get('approved') or []))
    text=' '.join(str(x or '') for x in [project.get('title'),project.get('goal'),project.get('decision_use'),json.dumps(project.get('briefing') or {},ensure_ascii=False),json.dumps(project.get('research_plan') or {},ensure_ascii=False)]).lower()
    scored=[]
    for did,d in dims.items():
        label=str(d.get('label') or did); score=5 if did.lower() in approved else 0
        for tok in re.findall(r'[a-zá-ž0-9_]{4,}',(did+' '+label+' '+str(d.get('rationale') or '')).lower()):
            if tok in text: score+=1
        if score or len(dims)<=limit: scored.append((score,did,d))
    selected=[]
    for _,did,d in sorted(scored,key=lambda x:(-x[0],x[1]))[:max(1,min(30,int(limit)))]:
        selected.append({'dimension_id':did,'label':d.get('label') or did,'action':d.get('last_action'),'rationale':d.get('rationale'),'evidence_summary':d.get('evidence_summary'),'confidence':d.get('confidence'),'runtime_status':d.get('runtime_status','KNOWLEDGE_ONLY'),'source_entry_id':d.get('source_entry_id')})
    try:
        from library_system_catalog import relevant_sources
        system_sources=relevant_sources(text,limit=limit)
    except Exception:
        system_sources=[]
    try:
        from population_context import project_population_context
        population=project_population_context(project)
    except Exception:
        population={}
    return {'dimensions':selected,'count':len(selected),'system_sources':system_sources,'system_source_count':len(system_sources),
            'population':population,'library_summary':summary_basic()}


def summary_basic()->dict[str,Any]:
    cx=_connect()
    try:
        total=cx.execute('SELECT COUNT(*) FROM library_entries').fetchone()[0]
        analyzed=cx.execute("SELECT COUNT(*) FROM library_entries WHERE analysis_status!='NOT_ANALYZED'").fetchone()[0]
        approved=cx.execute("SELECT COUNT(*) FROM dimension_proposals WHERE status='APPROVED'").fetchone()[0]
        cycles=cx.execute('SELECT COUNT(*) FROM learning_cycles').fetchone()[0]
    finally:cx.close()
    return {'sources':total,'analyzed':analyzed,'approved_changes':approved,'learning_cycles':cycles}


def public_entry(entry:dict[str,Any])->dict[str,Any]:
    d=dict(entry or {})
    d.pop('stored_path',None);d.pop('text_excerpt',None)
    return d


def add_source(*,raw:bytes,filename:str,source_type:str='other',title:str='',year:str='',author:str='',source_url:str='',notes:str='')->dict[str,Any]:
    if not raw: raise ValueError('Soubor je prázdný.')
    if len(raw)>50*1024*1024: raise ValueError('Jeden zdroj může mít maximálně 50 MB. U větších datasetů použijte respondentní ingest / vlastní audience.')
    if source_type not in SOURCE_TYPES: source_type='other'
    FILES_DIR.mkdir(parents=True,exist_ok=True)
    sha=hashlib.sha256(raw).hexdigest(); safe=_safe_filename(filename)
    # Content-addressed deduplication: the same file uploaded twice is one Library source.
    cx=_connect()
    try:
        old=cx.execute('SELECT entry_id FROM library_entries WHERE sha256=? ORDER BY created_at LIMIT 1',(sha,)).fetchone()
    finally: cx.close()
    if old:
        d=public_entry(get_entry(str(old['entry_id']))); d['deduplicated']=True; return d
    entry_id=_id('SRC',sha)
    path=FILES_DIR/f'{entry_id}_{safe}'; path.write_bytes(raw)
    text=_extract_text(raw,safe)
    hinted=_hint_dimensions(text+' '+title+' '+notes)
    cx=_connect()
    try:
        cx.execute('''INSERT INTO library_entries(entry_id,created_at,source_type,title,year,author,source_url,filename,stored_path,sha256,notes,text_excerpt,topics_json,dimensions_json,status,analysis_status,provenance_json)
        VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',(
            entry_id,_now(),source_type,(title or Path(safe).stem)[:300],str(year or '')[:20],str(author or '')[:300],str(source_url or '')[:1000],safe,str(path),sha,str(notes or '')[:4000],text[:12000],json.dumps([],ensure_ascii=False),json.dumps(hinted,ensure_ascii=False),'IMPORTED','NOT_ANALYZED',json.dumps({'origin':'user_upload','bytes':len(raw)},ensure_ascii=False)
        ))
        cx.commit()
    finally: cx.close()
    return public_entry(get_entry(entry_id))


def add_web_evidence(item:dict[str,Any],*,topic:str)->dict[str,Any]:
    payload=json.dumps(item,ensure_ascii=False,sort_keys=True)
    sha=hashlib.sha256(payload.encode()).hexdigest(); entry_id='WEB-'+sha[:14]
    cx=_connect()
    try:
        existing=cx.execute('SELECT entry_id FROM library_entries WHERE entry_id=?',(entry_id,)).fetchone()
        if not existing:
            claim=str(item.get('claim') or '')
            dims=_hint_dimensions(topic+' '+claim+' '+str(item.get('topics') or ''))
            cx.execute('''INSERT INTO library_entries(entry_id,created_at,source_type,title,year,author,source_url,filename,stored_path,sha256,notes,text_excerpt,summary,topics_json,dimensions_json,status,analysis_status,provenance_json)
            VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',(
                entry_id,_now(),'web_research',str(item.get('source_title') or 'Deep Research source')[:300],str(item.get('source_date') or '')[:20],'',str(item.get('source_url') or '')[:1000],'','',sha,'',claim[:12000],claim[:4000],json.dumps(item.get('topics') or [topic],ensure_ascii=False),json.dumps(dims,ensure_ascii=False),'IMPORTED','RESEARCH_IMPORTED',json.dumps({'origin':'deep_research','topic':topic,'source_quality':item.get('source_quality')},ensure_ascii=False)
            ));cx.commit()
    finally: cx.close()
    return public_entry(get_entry(entry_id))


def get_entry(entry_id:str)->dict[str,Any]:
    cx=_connect()
    try:r=cx.execute('SELECT * FROM library_entries WHERE entry_id=?',(entry_id,)).fetchone()
    finally:cx.close()
    if not r: raise KeyError(entry_id)
    d=dict(r)
    for k in ('topics_json','dimensions_json','provenance_json'):
        try:d[k[:-5] if k.endswith('_json') else k]=json.loads(d.pop(k) or '[]')
        except Exception:d[k[:-5]]= [] if k!='provenance_json' else {}
    return d


def list_entries(limit:int=200)->list[dict[str,Any]]:
    cx=_connect()
    try:rows=cx.execute('SELECT * FROM library_entries ORDER BY created_at DESC LIMIT ?',(max(1,min(1000,int(limit))),)).fetchall()
    finally:cx.close()
    out=[]
    for r in rows:
        d=dict(r)
        for k,new,default in [('topics_json','topics',[]),('dimensions_json','dimensions',[]),('provenance_json','provenance',{})]:
            try:d[new]=json.loads(d.pop(k) or ('{}' if isinstance(default,dict) else '[]'))
            except Exception:d[new]=default
        d.pop('stored_path',None);d.pop('text_excerpt',None)
        out.append(d)
    return out


def _proposal_from_ai(entry:dict[str,Any],*,model:str='sonnet',provider:str='claude_code_subscription')->dict[str,Any]:
    from ai_router import call_structured
    schema={'type':'object','properties':{
      'summary':{'type':'string'},'topics':{'type':'array','items':{'type':'string'},'maxItems':12},
      'dimension_impacts':{'type':'array','maxItems':12,'items':{'type':'object','properties':{
         'dimension_id':{'type':'string'},'dimension_label':{'type':'string'},
         'action':{'type':'string','enum':['new_dimension','refine_dimension','calibration_update','relationship_update','no_change']},
         'rationale':{'type':'string'},'evidence_summary':{'type':'string'},'confidence':{'type':'number','minimum':0,'maximum':1},
         'population_scope':{'type':'string'},'calibration_note':{'type':'string'},'relationship_note':{'type':'string'},
         'measurement_kind':{'type':'string','enum':['binary','scale_1_10','knowledge_only']},
         'target_prevalence':{'type':['number','null']},'target_mean':{'type':['number','null']},
         'quantitative_anchor_note':{'type':'string'},
         'predictors':{'type':'array','maxItems':6,'items':{'type':'object','properties':{
             'column':{'type':'string'},'direction':{'type':'string','enum':['positive','negative']},'strength':{'type':'number','minimum':0.05,'maximum':3}
         },'required':['column','direction','strength'],'additionalProperties':False}}
      },'required':['dimension_id','dimension_label','action','rationale','evidence_summary','confidence','measurement_kind','target_prevalence','target_mean','quantitative_anchor_note','predictors'],'additionalProperties':False}}
    },'required':['summary','topics','dimension_impacts'],'additionalProperties':False}
    excerpt=str(entry.get('text_excerpt') or '')[:18000]
    from audience_dimensions import CATEGORY_COLUMNS,DERIVED_COLUMNS
    predictor_ids=list(dict.fromkeys([c for cols in CATEGORY_COLUMNS.values() for c in cols]+list(DERIVED_COLUMNS)))
    payload={'source_type':entry.get('source_type'),'title':entry.get('title'),'year':entry.get('year'),'notes':entry.get('notes'),'text_excerpt':excerpt,'known_dimensions':sorted(DIMENSION_HINTS),'available_predictors':predictor_ids}
    from provider_auth import normalize_ai_provider
    from runtime_config import resolve_provider_model
    provider=normalize_ai_provider(provider); resolved_model=resolve_provider_model(provider,model)
    rr=call_structured(system='Jsi kurátor Data Library pro model české společnosti. Ze zdroje vytěž pouze poznatky, které mohou zpřesnit nebo rozšířit dimenze, vztahy mezi dimenzemi nebo kalibraci populace. Nic automaticky nepřepisuj. Každá změna je pouze návrh k lidskému schválení. Pro NOVOU respondentní dimenzi smíš vyplnit target_prevalence/target_mean jen když je kvantitativní populační kotva explicitně opřená o zdroj; jinak vrať null a knowledge_only. Predictory vybírej jen z available_predictors. Korelaci nevydávej za kauzalitu. Dimension_id piš krátce snake_case.',messages=[{'role':'user','content':json.dumps(payload,ensure_ascii=False)}],schema=schema,schema_name='npc_library_source_analysis',anthropic_model=resolved_model,openai_model=resolved_model,max_tokens=3500,prefer=provider,allow_fallback=False)
    return rr.get('data') or {}


def analyze_entry(entry_id:str,*,model:str='sonnet',provider:str='claude_code_subscription')->dict[str,Any]:
    entry=get_entry(entry_id)
    try:
        data=_proposal_from_ai(entry,model=model,provider=provider)
    except TypeError as exc:
        # Compatibility for local/test adapters created before the provider argument existed.
        # Production adapters are provider-aware; only retry when this exact keyword is unsupported.
        if 'provider' not in str(exc):
            raise
        data=_proposal_from_ai(entry,model=model)
    cx=_connect();created=[];pending_specs=[]
    try:
        cx.execute('UPDATE library_entries SET summary=?,topics_json=?,dimensions_json=?,analysis_status=? WHERE entry_id=?',(
            str(data.get('summary') or '')[:12000],json.dumps(data.get('topics') or [],ensure_ascii=False),json.dumps([str(x.get('dimension_id') or '') for x in (data.get('dimension_impacts') or []) if x.get('dimension_id')],ensure_ascii=False),'AI_ANALYZED',entry_id))
        for x in data.get('dimension_impacts') or []:
            if x.get('action')=='no_change': continue
            pid=_id('PROP',entry_id+str(x.get('dimension_id'))+str(x.get('action')))
            row=(pid,entry_id,_now(),str(x.get('dimension_id') or '')[:120],str(x.get('dimension_label') or x.get('dimension_id') or '')[:240],str(x.get('action') or '')[:60],str(x.get('rationale') or '')[:6000],str(x.get('evidence_summary') or '')[:6000],float(x.get('confidence') or 0),'PROPOSED',None)
            cx.execute('INSERT INTO dimension_proposals VALUES(?,?,?,?,?,?,?,?,?,?,?)',row);created.append(pid)
            pending_specs.append((pid,str(x.get('action') or ''),{'population_scope':x.get('population_scope'),'calibration_note':x.get('calibration_note'),'relationship_note':x.get('relationship_note'),'kind':x.get('measurement_kind'),'target_prevalence':x.get('target_prevalence'),'target_mean':x.get('target_mean'),'quantitative_anchor_note':x.get('quantitative_anchor_note'),'predictors':x.get('predictors') or []}))
        cx.commit()
    finally:cx.close()
    for pid,action,spec in pending_specs:_save_proposal_spec(pid,origin='source',action=action,spec=spec)
    return {'entry':public_entry(get_entry(entry_id)),'proposal_ids':created,'proposals':list_proposals(status='PROPOSED')}


def list_proposals(*,status:str|None=None,limit:int=300)->list[dict[str,Any]]:
    cx=_connect()
    try:
        if status:rows=cx.execute('SELECT * FROM dimension_proposals WHERE status=? ORDER BY created_at DESC LIMIT ?',(status,limit)).fetchall()
        else:rows=cx.execute('SELECT * FROM dimension_proposals ORDER BY created_at DESC LIMIT ?',(limit,)).fetchall()
    finally:cx.close()
    out=[]
    for r in rows:
        d=dict(r);d.update(_proposal_spec(d['proposal_id']));out.append(d)
    return out


def _load_overlays()->dict[str,Any]:
    if not OVERLAY_PATH.is_file():return {'version':'1','updated_at':'','dimensions':{}}
    try:return json.loads(OVERLAY_PATH.read_text(encoding='utf-8'))
    except Exception:return {'version':'1','updated_at':'','dimensions':{}}


def active_dimensions()->dict[str,dict[str,Any]]:
    return dict((_load_overlays().get('dimensions') or {}))


def decide_proposal(proposal_id:str,decision:str)->dict[str,Any]:
    decision=str(decision or '').upper()
    if decision not in {'APPROVE','REJECT'}:raise ValueError('decision musí být APPROVE nebo REJECT')
    cx=_connect()
    try:r=cx.execute('SELECT * FROM dimension_proposals WHERE proposal_id=?',(proposal_id,)).fetchone()
    finally:cx.close()
    if not r:raise KeyError(proposal_id)
    row=dict(r);status='APPROVED' if decision=='APPROVE' else 'REJECTED'
    overlays=_load_overlays()
    if status=='APPROVED':
        dims=overlays.setdefault('dimensions',{})
        did=row['dimension_id']
        cur=dict(dims.get(did) or {})
        before=dict(cur); ps=_proposal_spec(proposal_id)
        cur.update({'dimension_id':did,'label':row['dimension_label'],'last_action':row['action'],'rationale':row['rationale'],'evidence_summary':row['evidence_summary'],'confidence':row['confidence'],'proposal_id':proposal_id,'source_entry_id':row['entry_id'],'approved_at':_now(),'runtime_status':ps.get('runtime_status','KNOWLEDGE_ONLY'),'update_spec':ps.get('spec') or {},'origin':ps.get('origin','source')})
        dims[did]=cur;overlays['updated_at']=_now();OVERLAY_PATH.parent.mkdir(parents=True,exist_ok=True);OVERLAY_PATH.write_text(json.dumps(overlays,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
        hx=_connect()
        try:
            hx.execute('INSERT INTO calibration_history(event_id,proposal_id,dimension_id,created_at,before_json,after_json) VALUES(?,?,?,?,?,?)',(_id('CAL',proposal_id),proposal_id,did,_now(),json.dumps(before,ensure_ascii=False),json.dumps(cur,ensure_ascii=False)));hx.commit()
        finally:hx.close()
    cx=_connect()
    try:cx.execute('UPDATE dimension_proposals SET status=?,applied_at=? WHERE proposal_id=?',(status,_now() if status=='APPROVED' else None,proposal_id));cx.commit()
    finally:cx.close()
    return {'proposal':{**row,'status':status},'active_dimensions':active_dimensions()}


def deep_research(topic:str,*,model:str='sonnet',provider:str='claude_code_subscription',max_sources:int=8,progress=None)->dict[str,Any]:
    topic=str(topic or '').strip()
    if len(topic)<4:raise ValueError('Napište oblast, kterou má Deep Research doplnit.')
    from research_context import ResearchConfig,run_dual_research
    cfg=ResearchConfig(enabled=True,topic=topic,anthropic_model=model,max_sources_per_agent=max(3,min(12,int(max_sources))),max_context_blocks=10)
    bundle=run_dual_research(cfg,[],provider_override=provider,progress=progress)
    entries=[]
    for item in bundle.accepted or []:entries.append(add_web_evidence(item,topic=topic))
    # One compact AI synthesis over imported claims to create actionable proposals.
    synthetic={'entry_id':'RESEARCH-'+hashlib.sha256((topic+bundle.sha256).encode()).hexdigest()[:12],'source_type':'web_research','title':'Deep Research: '+topic,'year':'','notes':'','text_excerpt':'\n\n'.join(str(x.get('claim') or '')+'\n'+str(x.get('why_relevant') or '') for x in bundle.accepted[:20])}
    data=_proposal_from_ai(synthetic,model=model,provider=provider)
    cx=_connect();created=[];pending_specs=[]
    try:
        for x in data.get('dimension_impacts') or []:
            if x.get('action')=='no_change':continue
            pid=_id('PROP',synthetic['entry_id']+str(x.get('dimension_id')))
            cx.execute('INSERT INTO dimension_proposals VALUES(?,?,?,?,?,?,?,?,?,?,?)',(pid,None,_now(),str(x.get('dimension_id') or '')[:120],str(x.get('dimension_label') or x.get('dimension_id') or '')[:240],str(x.get('action') or '')[:60],str(x.get('rationale') or '')[:6000],str(x.get('evidence_summary') or '')[:6000],float(x.get('confidence') or 0),'PROPOSED',None));created.append(pid)
            pending_specs.append((pid,str(x.get('action') or ''),{'population_scope':x.get('population_scope'),'calibration_note':x.get('calibration_note'),'relationship_note':x.get('relationship_note'),'kind':x.get('measurement_kind'),'target_prevalence':x.get('target_prevalence'),'target_mean':x.get('target_mean'),'quantitative_anchor_note':x.get('quantitative_anchor_note'),'predictors':x.get('predictors') or []}))
        cx.commit()
    finally:cx.close()
    for pid,action,spec in pending_specs:_save_proposal_spec(pid,origin='deep_research',action=action,spec=spec)
    return {'topic':topic,'quality_status':bundle.quality_status,'sources_added':len(entries),'accepted_evidence':len(bundle.accepted),'quarantined_evidence':len(bundle.quarantined),'proposal_ids':created,'summary':str(data.get('summary') or ''),'entries':entries[:20],'proposals':list_proposals(status='PROPOSED')}


def list_learning_cycles(limit:int=100)->list[dict[str,Any]]:
    cx=_connect()
    try:rows=cx.execute('SELECT * FROM learning_cycles ORDER BY created_at DESC LIMIT ?',(max(1,min(500,int(limit))),)).fetchall()
    finally:cx.close()
    out=[]
    for r in rows:
        d=dict(r)
        for k,new in [('gaps_json','gaps'),('research_topics_json','research_topics'),('proposal_ids_json','proposal_ids')]:
            try:d[new]=json.loads(d.pop(k) or '[]')
            except Exception:d[new]=[]
        out.append(d)
    return out


def learn_from_completed_workflow(*,workflow_id:str,project:dict[str,Any],research:dict[str,Any]|None=None,verification:dict[str,Any]|None=None,model:str='sonnet',provider:str='claude_code_subscription')->dict[str,Any]:
    """Close the learning loop without teaching on synthetic answers.

    Synthetic survey results are deliberately excluded from calibration truth. The
    cycle may use the project's questions to identify gaps, while dimension-change
    proposals require external/library evidence from research or verification.
    """
    workflow_id=str(workflow_id or '').strip()
    if not workflow_id:raise ValueError('workflow_id je povinné')
    cx=_connect()
    try:existing=cx.execute('SELECT cycle_id FROM learning_cycles WHERE workflow_id=?',(workflow_id,)).fetchone()
    finally:cx.close()
    if existing:
        return next(x for x in list_learning_cycles(500) if x['cycle_id']==existing['cycle_id'])
    external=[]
    for x in (research or {}).get('accepted') or []:
        external.append({'claim':x.get('claim'),'why_relevant':x.get('why_relevant'),'source_title':x.get('source_title'),'source_url':x.get('source_url'),'source_date':x.get('source_date'),'source_quality':x.get('source_quality')})
    for x in (verification or {}).get('findings') or []:
        external.append({'claim':x.get('finding') or x.get('claim'),'why_relevant':x.get('comparison') or x.get('why_relevant'),'source_title':x.get('source_title'),'source_url':x.get('source_url'),'source_date':x.get('source_date'),'source_quality':x.get('source_quality')})
    active=active_dimensions()
    schema={'type':'object','properties':{
      'summary':{'type':'string'},
      'gaps':{'type':'array','maxItems':8,'items':{'type':'object','properties':{'topic':{'type':'string'},'why':{'type':'string'},'priority':{'type':'string','enum':['high','medium','low']}},'required':['topic','why','priority'],'additionalProperties':False}},
      'research_topics':{'type':'array','maxItems':6,'items':{'type':'string'}},
      'dimension_impacts':{'type':'array','maxItems':8,'items':{'type':'object','properties':{'dimension_id':{'type':'string'},'dimension_label':{'type':'string'},'action':{'type':'string','enum':['new_dimension','refine_dimension','calibration_update','relationship_update','no_change']},'rationale':{'type':'string'},'evidence_summary':{'type':'string'},'confidence':{'type':'number','minimum':0,'maximum':1}},'required':['dimension_id','dimension_label','action','rationale','evidence_summary','confidence'],'additionalProperties':False}}
    },'required':['summary','gaps','research_topics','dimension_impacts'],'additionalProperties':False}
    from ai_router import call_structured
    payload={'project':{'title':project.get('title'),'goal':project.get('goal'),'decision_use':project.get('decision_use'),'persona_dimensions':(project.get('persona_dimensions') or {}).get('approved') or [],'research_plan':project.get('research_plan'),'sections':project.get('sections')},'external_evidence':external[:40],'active_library_dimensions':active}
    from provider_auth import normalize_ai_provider
    from runtime_config import resolve_provider_model
    provider=normalize_ai_provider(provider); resolved_model=resolve_provider_model(provider,model)
    rr=call_structured(system=('Jsi kurátor kontinuálního učení NPC Panelu. Dokončený syntetický survey NESMÍŠ používat jako validační pravdu ani podle něj kalibrovat populaci. Z projektu smíš zjistit pouze, jaká témata/dimenze byly důležité. Návrh změny dimenze smí vzniknout jen pokud jej podporuje external_evidence. Když externí evidence chybí, dimension_impacts musí být prázdné a navrhni pouze research gaps. Nic automaticky neaplikuj.'),messages=[{'role':'user','content':json.dumps(payload,ensure_ascii=False,default=str)}],schema=schema,schema_name='npc_learning_cycle',anthropic_model=resolved_model,openai_model=resolved_model,max_tokens=3200,prefer=provider,allow_fallback=False)
    data=rr.get('data') or {};created=[]
    if external:
        cx=_connect();pending_specs=[]
        try:
            for x in data.get('dimension_impacts') or []:
                if x.get('action')=='no_change':continue
                pid=_id('LEARN',workflow_id+str(x.get('dimension_id'))+str(x.get('action')))
                cx.execute('INSERT INTO dimension_proposals VALUES(?,?,?,?,?,?,?,?,?,?,?)',(pid,None,_now(),str(x.get('dimension_id') or '')[:120],str(x.get('dimension_label') or x.get('dimension_id') or '')[:240],str(x.get('action') or '')[:60],str(x.get('rationale') or '')[:6000],str(x.get('evidence_summary') or '')[:6000],float(x.get('confidence') or 0),'PROPOSED',None));created.append(pid)
                pending_specs.append((pid,str(x.get('action') or ''),{'workflow_id':workflow_id,'evidence_count':len(external)}))
            cx.commit()
        finally:cx.close()
        for pid,action,spec in pending_specs:_save_proposal_spec(pid,origin='workflow_external_evidence',action=action,spec=spec)
    cid=_id('CYCLE',workflow_id);status='PROPOSALS_READY' if created else ('GAPS_IDENTIFIED' if (data.get('gaps') or data.get('research_topics')) else 'NO_CHANGE')
    cx=_connect()
    try:
        cx.execute('INSERT INTO learning_cycles(cycle_id,workflow_id,created_at,project_title,status,summary,gaps_json,research_topics_json,proposal_ids_json,evidence_count) VALUES(?,?,?,?,?,?,?,?,?,?)',(cid,workflow_id,_now(),str(project.get('title') or '')[:300],status,str(data.get('summary') or '')[:8000],json.dumps(data.get('gaps') or [],ensure_ascii=False),json.dumps(data.get('research_topics') or [],ensure_ascii=False),json.dumps(created,ensure_ascii=False),len(external)));cx.commit()
    finally:cx.close()
    return next(x for x in list_learning_cycles(500) if x['cycle_id']==cid)


def summary()->dict[str,Any]:
    cx=_connect()
    try:
        total=cx.execute('SELECT COUNT(*) FROM library_entries').fetchone()[0]
        analyzed=cx.execute("SELECT COUNT(*) FROM library_entries WHERE analysis_status!='NOT_ANALYZED'").fetchone()[0]
        proposed=cx.execute("SELECT COUNT(*) FROM dimension_proposals WHERE status='PROPOSED'").fetchone()[0]
        approved=cx.execute("SELECT COUNT(*) FROM dimension_proposals WHERE status='APPROVED'").fetchone()[0]
        project_links=cx.execute('SELECT COUNT(*) FROM library_project_links').fetchone()[0]
    finally:cx.close()
    cycles=list_learning_cycles(100)
    try:
        from library_system_catalog import catalog as system_catalog
        system_summary=system_catalog().get('summary') or {}
    except Exception: system_summary={}
    try:
        from population_context import summary as population_summary
        pop=population_summary()
    except Exception: pop={}
    try:
        from results_registry import summary as results_summary
        results=results_summary(sync=False)
    except Exception: results={}
    return {'sources':total,'user_sources':total,'system_sources':int(system_summary.get('source_groups') or 0),
            'total_visible_sources':total+int(system_summary.get('source_groups') or 0),
            'analyzed':analyzed,'proposals':proposed,'approved_changes':approved,'active_dimensions':active_dimensions(),
            'learning_cycles':len(cycles),'recent_learning_cycles':cycles[:12],'system_catalog':system_summary,
            'population':pop,'results_registry':results,'project_links':project_links,
            'contract':{'library_import_mutates_live':False,'synthetic_results_calibrate_live':False,'human_approval_required':True}}

# ---------------------------------------------------------------------------
# 17.9.3 evidence-backed custom dimensions
# ---------------------------------------------------------------------------
def create_dimension_request(*,label:str,dimension_id:str='',rationale:str='',source_strategy:str='document_or_research',spec:dict[str,Any]|None=None,entry_id:str|None=None,origin:str='user')->dict[str,Any]:
    label=str(label or '').strip()
    if len(label)<2:raise ValueError('Napište název nové dimenze.')
    did=re.sub(r'[^a-z0-9_]+','_',str(dimension_id or label).lower()).strip('_')[:100]
    if not did:raise ValueError('Nelze vytvořit dimension_id.')
    pid=_id('PROP',did+label+origin)
    rationale=str(rationale or f'Uživatel požaduje novou dimenzi: {label}')[:6000]
    cx=_connect()
    try:
        cx.execute('INSERT INTO dimension_proposals VALUES(?,?,?,?,?,?,?,?,?,?,?)',(pid,entry_id,_now(),did,label,'new_dimension',rationale,'',0.0,'PROPOSED',None));cx.commit()
    finally:cx.close()
    merged={'source_strategy':source_strategy,**(spec or {})}
    _save_proposal_spec(pid,origin=origin,action='new_dimension',spec=merged)
    return {'proposal_id':pid,'dimension_id':did,'label':label,'status':'PROPOSED','runtime_status':'NEEDS_QUANT_ANCHOR','spec':merged}


def update_dimension_spec(proposal_id:str,spec:dict[str,Any])->dict[str,Any]:
    cur=_proposal_spec(proposal_id);merged=dict(cur.get('spec') or {});merged.update(dict(spec or {}))
    _save_proposal_spec(proposal_id,origin=cur.get('origin','user'),action='new_dimension',spec=merged)
    return {'proposal_id':proposal_id,'spec':merged,'materializable':dimension_materialization_readiness(proposal_id)}


def dimension_materialization_readiness(proposal_id:str)->dict[str,Any]:
    cx=_connect()
    try:r=cx.execute('SELECT * FROM dimension_proposals WHERE proposal_id=?',(proposal_id,)).fetchone()
    finally:cx.close()
    if not r:return {'ready':False,'reasons':['proposal nenalezen']}
    row=dict(r);spec=_proposal_spec(proposal_id).get('spec') or {};reasons=[]
    if row.get('status')!='APPROVED':reasons.append('návrh ještě není schválen')
    kind=str(spec.get('kind') or spec.get('measurement_kind') or 'knowledge_only')
    if kind in {'binary','boolean','prevalence'}:
        try:float(spec.get('target_prevalence'))
        except Exception:reasons.append('chybí evidence-based target_prevalence')
    elif kind in {'scale_1_10','1_10','continuous'}:
        try:float(spec.get('target_mean'))
        except Exception:reasons.append('chybí evidence-based target_mean')
    else:reasons.append('measurement kind je knowledge_only')
    if not (spec.get('predictors') or []):reasons.append('chybí alespoň jeden existující predictor')
    return {'ready':not reasons,'reasons':reasons,'kind':kind,'spec':spec,'status':row.get('status'),'dimension_id':row.get('dimension_id')}


def materialize_dimension(proposal_id:str)->dict[str,Any]:
    ready=dimension_materialization_readiness(proposal_id)
    if not ready.get('ready'):raise ValueError('Dimenzi nelze dosimulovat: '+'; '.join(ready.get('reasons') or []))
    cx=_connect()
    try:r=cx.execute('SELECT * FROM dimension_proposals WHERE proposal_id=?',(proposal_id,)).fetchone()
    finally:cx.close()
    row=dict(r);spec=ready['spec']
    from pipeline import PANEL_PATH
    import pandas as pd
    from persona_depth import attach_overlay
    from audience_dimensions import attach_derived,materialize_evidence_dimension
    base=pd.read_csv(PANEL_PATH,low_memory=False,dtype={'occupation_isco08':'string'})
    base=attach_derived(attach_overlay(base))
    mat=materialize_evidence_dimension(base,dimension_id=row['dimension_id'],spec=spec)
    overlays=_load_overlays();dims=overlays.setdefault('dimensions',{});cur=dict(dims.get(row['dimension_id']) or {})
    cur.update({'dimension_id':row['dimension_id'],'label':row['dimension_label'],'last_action':'new_dimension','rationale':row['rationale'],'evidence_summary':row.get('evidence_summary') or spec.get('quantitative_anchor_note') or '',
                'confidence':row.get('confidence'),'proposal_id':proposal_id,'source_entry_id':row.get('entry_id'),'approved_at':cur.get('approved_at') or _now(),'runtime_status':'SIMULATED_FROM_EVIDENCE','origin':_proposal_spec(proposal_id).get('origin','source'),
                'overlay_path':mat['path'],'overlay_sha256':mat['sha256'],'materialized_at':_now(),'target_anchor':mat['target_anchor'],'predictors':mat['predictors'],'kind':mat['kind']})
    dims[row['dimension_id']]=cur;overlays['updated_at']=_now();OVERLAY_PATH.parent.mkdir(parents=True,exist_ok=True);OVERLAY_PATH.write_text(json.dumps(overlays,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    hx=_connect()
    try:
        hx.execute('INSERT INTO calibration_history(event_id,proposal_id,dimension_id,created_at,before_json,after_json) VALUES(?,?,?,?,?,?)',(_id('CAL',proposal_id+'materialize'),proposal_id,row['dimension_id'],_now(),json.dumps({},ensure_ascii=False),json.dumps(cur,ensure_ascii=False)))
        hx.execute('UPDATE proposal_specs SET runtime_status=? WHERE proposal_id=?',('SIMULATED_FROM_EVIDENCE',proposal_id));hx.commit()
    finally:hx.close()
    population_event=None
    try:
        from population_context import propose_calibration,decide_calibration,register_applied_overlay
        population_event=propose_calibration(source_ref=str(row.get('entry_id') or proposal_id),proposal_id=proposal_id,dimension_id=row['dimension_id'],
                                             change_type='evidence_overlay',spec={'overlay_path':mat['path'],'overlay_sha256':mat['sha256'],'target_anchor':mat['target_anchor'],'kind':mat['kind']},
                                             note='Explicitní materializace schválené Data Library dimenze do LIVE runtime overlay vrstvy.')
        decide_calibration(population_event['event_id'],'APPROVE')
        population_event=register_applied_overlay(event_id=population_event['event_id'])
    except Exception:
        population_event=None
    return {'proposal_id':proposal_id,'dimension_id':row['dimension_id'],'status':'MATERIALIZED','runtime_status':'SIMULATED_FROM_EVIDENCE',**mat,'active_dimension':cur,'population_event':population_event}



# ---------------------------------------------------------------------------
# 18.5 project/library linkage
# ---------------------------------------------------------------------------
def link_source_to_project(project_id:str,source_ref:str,*,source_kind:str='library_entry',usage_role:str='context')->dict[str,Any]:
    project_id=str(project_id or '').strip(); source_ref=str(source_ref or '').strip()
    if not project_id or not source_ref: raise ValueError('Chybí project_id nebo source_ref.')
    cx=_connect()
    try:
        cx.execute('INSERT OR REPLACE INTO library_project_links(project_id,source_ref,source_kind,usage_role,created_at) VALUES(?,?,?,?,?)',
                   (project_id,source_ref,str(source_kind or 'library_entry')[:60],str(usage_role or 'context')[:60],_now()));cx.commit()
    finally: cx.close()
    return {'project_id':project_id,'source_ref':source_ref,'source_kind':source_kind,'usage_role':usage_role}

def project_sources(project_id:str)->list[dict[str,Any]]:
    cx=_connect()
    try: rows=cx.execute('SELECT * FROM library_project_links WHERE project_id=? ORDER BY created_at DESC',(str(project_id),)).fetchall()
    finally: cx.close()
    return [dict(x) for x in rows]
