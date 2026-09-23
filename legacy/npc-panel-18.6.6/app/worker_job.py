from __future__ import annotations

import argparse
import json
import os
import sys
import time
import traceback
import threading
import re
import datetime as dt
from dataclasses import asdict
from pathlib import Path
from typing import Any

from job_store import JobStore
from cost_controller import CostController, WorkflowBudgetExceeded

ROOT = Path(__file__).resolve().parent
ART = ROOT / 'prototype_outputs' / 'research_os_artifacts'
ART.mkdir(exist_ok=True)

# Conservative stage estimates are used only when a paid API is explicitly selected.
# Claude Code subscription remains the primary zero-incremental-API-cost route.
_STAGE_API_ESTIMATES = {
    'background_research': {'anthropic': 0.75},
    'interpret_results': {'anthropic': 0.30},
    'analysis_module': {'anthropic': 0.08},
    'analysis_assemble': {'anthropic': 0.0},
    'external_verification': {'anthropic': 0.75},
    'reality_alignment': {'anthropic': 0.35},
    'final_report': {'anthropic': 0.35},
}


def _write(wid: str, key: str, obj: Any) -> str:
    p = ART / wid / f'{key}.json'
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(obj, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
    return str(p)


def _deps(store: JobStore, job: dict[str, Any]) -> list[dict[str, Any]]:
    with store.cx() as c:
        rows = c.execute(
            'SELECT j2.* FROM job_dependencies d JOIN jobs j2 ON j2.job_id=d.depends_on_job_id WHERE d.job_id=?',
            (job['job_id'],),
        ).fetchall()
    return [store._job(r) for r in rows]


def _dep_output(store: JobStore, job: dict[str, Any], node_key: str) -> dict[str, Any]:
    for d in _deps(store, job):
        if d['node_key'] == node_key:
            return d.get('output') or {}
    return {}


def _all_prior_outputs(store: JobStore, wid: str) -> dict[str, dict[str, Any]]:
    return {j['node_key']: j.get('output') or {} for j in store.get_workflow(wid)['jobs']}


def _project(job: dict[str, Any]) -> dict[str, Any]:
    return dict(job['input'].get('project') or {})


def _approved_provider(inp: dict[str, Any]) -> str | None:
    for d in reversed(inp.get('approval_decisions') or []):
        if d.get('option') in {'use_anthropic_api','continue_with_claude_api'}:
            return 'anthropic'
    return None


def _provider_for(project: dict[str, Any], inp: dict[str, Any]) -> str:
    from provider_auth import normalize_ai_provider
    phase=inp.get('phase_policy') or {}
    return normalize_ai_provider(_approved_provider(inp) or phase.get('provider') or (project.get('run_policy') or {}).get('provider') or 'claude_code_subscription')

def _model_for(project: dict[str, Any], inp: dict[str, Any], default='sonnet') -> str:
    return str((inp.get('phase_policy') or {}).get('model') or project.get('model') or default)


def _is_claude_subscription_failure(exc: Exception) -> bool:
    txt = str(exc)
    return (
        'SUBSCRIPTION_PROVIDER_' in txt
        or 'ClaudeCodeLimit' in type(exc).__name__
        or 'ClaudeCodeUnavailable' in type(exc).__name__
        or ('claude code' in txt.lower() and any(x in txt.lower() for x in ('limit', 'usage', 'auth', 'unavailable', 'busy')))
    )


def _is_subscription_quota_failure(exc: Exception) -> bool:
    """True only for a real Claude subscription usage/session cap."""
    txt=str(exc or ''); low=txt.lower()
    # Do not use the generic API 429 classifier here: Anthropic API capacity is
    # WAITING_CAPACITY, not a Claude Code subscription reset.
    return any(x in low for x in ('session limit','usage limit','usage cap','out_of_credits','rate_limit_event','subscription_provider_limit','claudecodelimit'))

def _is_api_capacity_failure(exc: Exception) -> bool:
    status=getattr(exc,'status_code',None) or getattr(exc,'status',None); low=str(exc or '').lower()
    return status in {429,503} or any(x in low for x in ('[quota] anthropic','rate_limit','rate limit','overloaded','503','service unavailable','api_capacity'))

def _quota_reset_at(exc: Exception) -> tuple[str,str|None]:
    """Return local scheduler timestamp + raw provider reset epoch when available."""
    txt=str(exc or '')
    m=re.search(r'resetsAt[\"\'\s:=]+(\d{10,13})',txt,re.I)
    raw=m.group(1) if m else None
    if raw:
        ts=int(raw); ts=ts/1000 if ts>20_000_000_000 else ts
        when=dt.datetime.fromtimestamp(ts)+dt.timedelta(seconds=75)
    else:
        # No exact reset timestamp: retry one hour later rather than burning attempts.
        when=dt.datetime.now()+dt.timedelta(hours=1)
    return when.strftime('%Y-%m-%dT%H:%M:%S'),raw

def _park_for_subscription_quota(store: JobStore, wid: str, job_id: str, exc: Exception, *, run_dir: Path|None=None) -> int:
    """Park a durable workflow on Claude Pro quota and auto-resume after reset.

    Quota is availability, not a failed attempt. Completed respondent journal/checkpoint
    data stay untouched. Attempts are reset for the next availability window.
    """
    reset_at,raw_reset=_quota_reset_at(exc)
    msg=f'Claude Pro session limit dosažen · automaticky pokračuji po {reset_at.replace("T"," ")}'
    err={'error':'WAITING_FOR_CLAUDE_QUOTA','provider':'claude_code_subscription','reason':str(exc)[:1800],
         'auto_resume_at':reset_at,'provider_resetsAt':raw_reset,'checkpoint':str(run_dir) if run_dir else None}
    store.transition(job_id,'WAITING_CREDITS',error=err,message=msg,force=True)
    # WAITING_CREDITS must not consume the retry budget or keep a stale lease.
    with store.cx() as c:
        c.execute("UPDATE jobs SET attempt=0,lease_owner=NULL,lease_until=NULL,cancel_requested=0,updated_at=? WHERE job_id=?",(dt.datetime.now().strftime('%Y-%m-%dT%H:%M:%S'),job_id))
    store.event(job_id,'PROGRESS',msg,{'phase':msg,'provider_stage':'waiting_subscription_reset','auto_resume_at':reset_at,
                'resume_available':bool(run_dir and (run_dir/'checkpoint.pkl').is_file()),'checkpoint':str(run_dir) if run_dir else None},'WARN')
    # Dedupe one-time resume schedules for this workflow.
    try:
        from scheduler import Scheduler
        existing=False
        with store.cx() as c:
            rows=c.execute("SELECT schedule_id,workflow_template_json,next_run_at FROM schedules WHERE enabled=1 AND next_run_at IS NOT NULL").fetchall()
        for r in rows:
            try:
                tpl=json.loads(r['workflow_template_json'] or '{}')
                if str(tpl.get('resume_workflow_id') or '')==wid:
                    existing=True; break
            except Exception:
                pass
        if not existing:
            wf=store.get_workflow(wid) or {}
            Scheduler(store).create({'project_id':wf.get('project_id') or 'quota-resume',
                'project_revision':wf.get('project_revision'),'use_latest_revision':False,'schedule_type':'one_time',
                'schedule_expr':reset_at,'next_run_at':reset_at,'timezone':'Europe/Prague','misfire_policy':'run_once',
                'workflow_template':{'resume_workflow_id':wid,'reason':'claude_subscription_quota','checkpoint':str(run_dir) if run_dir else None}})
        store.event(job_id,'AUTO_RESUME_SCHEDULED','Claude subscription reset auto-resume scheduled',{'resume_at':reset_at,'workflow_id':wid})
    except Exception as sched_exc:
        # Parking is still safe if scheduler metadata fails; user can resume manually.
        store.event(job_id,'WARN','Quota auto-resume scheduling failed; checkpoint remains resumable',{'error':str(sched_exc)[:700],'resume_at':reset_at},'WARN')
    store.refresh_workflow(wid)
    return 0

def _respondent_retryable_failure(exc: Exception) -> bool:
    """Transient fieldwork failures safe to replay from the durable respondent journal."""
    txt=str(exc or ''); low=txt.lower()
    if any(x in low for x in ('job_cancelled','authentication','not authenticated','login required','usage limit','credit')):
        return False
    if any(x in txt for x in ('RESPONDENT_BATCH_', 'RESPONDENT_QUESTION_WATCHDOG_TIMEOUT', 'SUBSCRIPTION_PROVIDER_BUSY', 'SUBSCRIPTION_PROVIDER_START_TIMEOUT', 'SUBSCRIPTION_PROVIDER_TIMEOUT')):
        return True
    return any(x in low for x in ('timed out','timeout','temporarily unavailable','connection reset','connection aborted','provider process exited','backend connection'))


def _stage_estimates(kind: str) -> dict[str, float]:
    return dict(_STAGE_API_ESTIMATES.get(kind) or {'anthropic': 0.50})


def _reserve_stage_if_paid(store: JobStore, wid: str, job_id: str, kind: str, provider: str) -> tuple[CostController, str | None, float]:
    cc = CostController(store, wid)
    if provider != 'anthropic':
        return cc, None, 0.0
    amount = float(_stage_estimates(kind).get('anthropic',0.50))
    # Project-level paid API budget is checked before the external request.
    job=store.get_job(job_id) or {}; pid=job.get('project_id'); rev=job.get('project_revision')
    if pid:
        try:
            from project_store import ProjectStore
            from provider_runtime import api_budget_check
            ps=ProjectStore(ROOT/'data'/'project_store.sqlite')
            try:
                pr=ps.get(pid,rev) or ps.get(pid)
                check=api_budget_check(spent_usd=float((pr or {}).get('api_spent_usd') or 0),estimate_usd=amount,max_api_cost_usd=float((pr or {}).get('max_api_cost_usd') or 10.0))
                if not check['allowed']:
                    store.create_approval(job_id,'Claude API krok by překročil maximální API rozpočet projektu.',[{'option':'increase_budget'},{'option':'continue_later'},{'option':'cancel'}],{'reason':'PROJECT_API_BUDGET_LIMIT','provider':'anthropic','stage':kind,**check})
                    ps.set_stage(pid,int(rev or pr.get('revision') or 1),job.get('stage_id') or 'ANALYSIS','WAITING_USER',waiting_reason='PROJECT_API_BUDGET_LIMIT',current_job_id=job_id)
                    # The project and the durable job must agree. Returning from the
                    # worker with a RUNNING job would leave a zombie lease and make
                    # the explicit budget decision impossible to resume reliably.
                    store.transition(job_id,'WAITING_USER',error={'error':'PROJECT_API_BUDGET_LIMIT','provider':'anthropic','stage':kind,**check},message='Claude API budget requires explicit user decision',force=True)
                    store.refresh_workflow(wid)
                    return cc,'__WAITING__',amount
            finally: ps.close()
        except Exception:
            # Workflow budget remains a second independent safety gate.
            pass
    try:
        return cc, cc.reserve(job_id, amount), amount
    except WorkflowBudgetExceeded as exc:
        store.create_approval(job_id,'Workflow budget nestačí pro tento placený Claude API krok.',[{'option':'increase_budget'},{'option':'continue_later'},{'option':'cancel'}],{'reason':str(exc),'new_cost_estimate_usd':amount,'provider':'anthropic','stage':kind})
        store.transition(job_id,'WAITING_CREDITS',error={'reason':str(exc),'new_cost_estimate_usd':amount,'provider':'anthropic','stage':kind},message='waiting for budget/credits decision',force=True)
        return cc,'__WAITING__',amount

def _settle_stage(cc: CostController, reservation: str | None, estimated: float) -> None:
    if reservation and reservation != '__WAITING__':
        # Until provider token accounting is normalized across APIs, settle the conservative
        # user-approved estimate. Subscription calls have no API reservation at all.
        cc.settle(reservation, estimated)


def _cancel_stage(cc: CostController, reservation: str | None) -> None:
    if reservation and reservation != '__WAITING__':
        cc.cancel(reservation)


def _compiled_questions(project: dict[str, Any]) -> list[dict[str, Any]]:
    from research_project import compile_project
    return list((compile_project(project).get('brief') or {}).get('otazky') or [])


def _research_topic(project: dict[str, Any]) -> str:
    parts = [project.get('title'), project.get('goal'), project.get('decision_use')]
    rp = project.get('research_plan') or {}
    parts += list(rp.get('research_questions') or [])[:4]
    parts += list(rp.get('objectives') or [])[:4]
    return ' | '.join(str(x).strip() for x in parts if str(x or '').strip())[:4000]


def _collect_maps(prior: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    result = (prior.get('run') or {}).get('result') or {}
    maps = []
    for b in result.get('batteries') or []:
        if not isinstance(b, dict) or b.get('error'):
            continue
        files = b.get('files') or {}
        png = files.get('map_png')
        html = files.get('map_html')
        matrix = files.get('relation_matrix')
        if not (png or html or matrix):
            continue
        maps.append({
            'id': b.get('id'),
            'title': b.get('title') or b.get('id') or 'Relační mapa',
            'png': png,
            'html': html,
            'relation_matrix': matrix,
            'nearest_pairs': b.get('nearest_pairs') or [],
            'map_metrics': b.get('map_metrics') or {},
        })
    return maps



def _fallback_report_html(path: Path, report: dict[str,Any], project: dict[str,Any]) -> str:
    import html as _html
    p=Path(path);p.parent.mkdir(parents=True,exist_ok=True)
    e=lambda x:_html.escape(str(x or ''))
    findings=''.join(f"<section><h3>{e(x.get('headline'))}</h3><p>{e(x.get('finding'))}</p><p><b>Co to znamená:</b> {e(x.get('meaning'))}</p></section>" for x in (report.get('key_findings') or []))
    limits=''.join(f'<li>{e(x)}</li>' for x in (report.get('limitations') or []))
    txt=("<!doctype html><meta charset='utf-8'><style>body{font:16px/1.6 Arial;margin:0;background:#f4f6f8}main{max-width:1000px;margin:auto;background:white;padding:48px}h1{font-size:34px}.warn{padding:12px;background:#fff4d6;border:1px solid #e7c46a}</style><main>"
         +"<div class='warn'>Výstup byl dokončen v nouzovém exportním režimu. Výsledková data a evidence zůstávají zachované.</div>"
         +f"<h1>{e(report.get('title') or project.get('title') or 'Výsledky')}</h1><h2>Executive summary</h2><p>{e(report.get('executive_summary'))}</p><h2>Odpověď pro rozhodnutí</h2><p>{e(report.get('decision_answer'))}</p><h2>Co jsme zjistili</h2>{findings}<h2>Limity</h2><ul>{limits}</ul></main>")
    p.write_text(txt,encoding='utf-8');return str(p)


def _fallback_report_docx(path: Path, report: dict[str,Any], project: dict[str,Any]) -> str:
    from docx import Document
    p=Path(path);p.parent.mkdir(parents=True,exist_ok=True);d=Document()
    d.add_heading(str(report.get('title') or project.get('title') or 'Výsledky'),0)
    d.add_paragraph('Nouzový exportní režim: výsledková data a evidence zůstávají zachované; stylistické/sekundární exporty mohou vyžadovat kontrolu.')
    d.add_heading('Executive summary',1);d.add_paragraph(str(report.get('executive_summary') or ''))
    d.add_heading('Odpověď pro rozhodnutí',1);d.add_paragraph(str(report.get('decision_answer') or ''))
    d.add_heading('Co jsme zjistili',1)
    for x in report.get('key_findings') or []:
        d.add_heading(str(x.get('headline') or 'Zjištění'),2);d.add_paragraph(str(x.get('finding') or ''));d.add_paragraph('Co to znamená: '+str(x.get('meaning') or ''))
    d.add_heading('Limity',1)
    for x in report.get('limitations') or []: d.add_paragraph(str(x),style='List Bullet')
    d.save(p);return str(p)


def _fallback_manifest(path: Path, files: dict[str,Any], project: dict[str,Any], workflow_id: str, warnings: list[dict[str,Any]]) -> str:
    p=Path(path);p.parent.mkdir(parents=True,exist_ok=True)
    obj={'workflow_id':workflow_id,'project_title':project.get('title'),'generated_at':time.strftime('%Y-%m-%dT%H:%M:%SZ',time.gmtime()),'status':'RESULTS_READY_DEGRADED_EXPORT' if warnings else 'RESULTS_READY','outputs':files,'artifact_warnings':warnings}
    p.write_text(json.dumps(obj,ensure_ascii=False,indent=2,default=str),encoding='utf-8');return str(p)

def execute(job_id: str, worker_id: str = 'worker-job') -> int:
    s = JobStore()
    job = s.get_job(job_id)
    if not job:
        return 2

    stop = threading.Event()

    def beat():
        s.heartbeat(job_id, job.get('lease_owner') or worker_id)
        while not stop.wait(8):
            s.heartbeat(job_id, job.get('lease_owner') or worker_id)

    th = threading.Thread(target=beat, daemon=True)
    th.start()

    def progress(payload):
        s.event(job_id, 'PROGRESS', str(payload.get('phase') or 'progress'), payload)
        s.heartbeat(job_id, job.get('lease_owner') or worker_id)

    def cancel():
        return s.cancel_requested(job_id)

    from ai_execution_context import bind_runtime, reset_runtime
    _ai_runtime_tokens = bind_runtime(progress, cancel)

    try:
        if s.cancel_requested(job_id):
            s.transition(job_id,'CANCELLED',error={'error':'JOB_CANCELLED'},message='cancelled before task dispatch')
            return 0
        kind = job['kind']
        inp = job['input']
        wid = job['workflow_id']
        project = _project(job)
        out: dict[str, Any] = {}
        project_id=str(job.get('project_id') or inp.get('project_id') or '').strip()
        project_revision=int(job.get('project_revision') or inp.get('project_revision') or 0)
        stage_id=str(job.get('stage_id') or inp.get('stage_id') or '').strip()
        artifact_target=str(job.get('artifact_target') or inp.get('artifact_target') or kind).strip()
        input_fp=str(job.get('input_fingerprint') or inp.get('input_fingerprint') or '')
        _ps=None; _astore=None
        if project_id and project_revision and stage_id:
            from project_store import ProjectStore
            from artifact_store import ArtifactStore
            _ps=ProjectStore(ROOT/'data'/'project_store.sqlite'); _astore=ArtifactStore(_ps,ROOT/'data'/'project_artifacts')
            reused=_ps.reusable_artifact(project_id,project_revision,stage_id,artifact_target,input_fp) if input_fp else None
            # A new immutable revision may intentionally preserve a completed
            # upstream stage by referencing artifacts from its parent revision.
            # Reuse that verified artifact rather than recomputing it merely
            # because the registry row belongs to the older revision.
            if not reused and input_fp:
                st=_ps.stage(project_id,project_revision,stage_id) or {}
                if st.get('status') in {'DONE','DONE_WITH_WARNINGS'}:
                    for aid in reversed(st.get('artifact_ids') or []):
                        cand=_ps.get_artifact(aid)
                        if cand and cand.get('artifact_type')==artifact_target and str(cand.get('input_fingerprint') or '')==input_fp and _astore.verify(aid):
                            reused=cand; break
            if not reused and input_fp:
                reused=_ps.reusable_artifact_any_revision(project_id,stage_id,artifact_target,input_fp)
            if reused:
                try: out=json.loads(Path(reused['path']).read_text(encoding='utf-8'))
                except Exception: out={'artifact_id':reused['artifact_id'],'path':reused['path'],'reused':True}
                out['_project_artifact_reused']=True;out['_artifact_id']=reused['artifact_id']
                _ps.set_stage(project_id,project_revision,stage_id,'DONE',input_fingerprint=input_fp,provider=job.get('provider_policy'),model=job.get('model'),current_job_id=job_id,artifacts=[reused['artifact_id']])
                s.event(job_id,'ARTIFACT_REUSED','Durable project artifact reused; expensive stage not recomputed.',{'artifact_id':reused['artifact_id'],'stage_id':stage_id,'input_fingerprint':input_fp})
                s.transition(job_id,'COMPLETED',output=out);s.queue_ready_jobs(wid);_ps.close();return 0
            _ps.set_stage(project_id,project_revision,stage_id,'RUNNING',input_fingerprint=input_fp,provider=job.get('provider_policy'),model=job.get('model'),current_job_id=job_id)

        if kind == 'project_compile':
            from research_project import compile_project
            c = compile_project(project)
            out = {
                'compiled_question_count': len((c.get('brief') or {}).get('otazky') or []),
                'project': c.get('project'),
            }

        elif kind == 'background_research':
            mode = str(inp.get('mode') or 'dry')
            rp=project.get('run_policy') or {}
            existing=project.get('pre_research') or {}
            has_existing=bool((existing.get('accepted') or []) or (existing.get('agents') or []))
            final_research=bool(rp.get('final_research', True))
            update_research=bool(rp.get('update_research_before_run', False))
            if mode == 'dry':
                out = {
                    'research_status': 'NOT_RUN_DRY',
                    'quality_status': 'NO_RESEARCH',
                    'bundle': {'enabled': False, 'accepted': [], 'quarantined': [], 'agents': []},
                    'provider': None,
                    'api_cost_usd': 0.0,
                }
            elif has_existing and not update_research:
                out={'research_status':'REUSED_PROJECT_RESEARCH','quality_status':existing.get('quality_status') or 'REUSED',
                     'bundle':existing,'accepted_count':len(existing.get('accepted') or []),'quarantined_count':len(existing.get('quarantined') or []),
                     'provider':'reused','api_cost_usd':0.0,'research_context_reused':True}
            elif not final_research:
                out={'research_status':'SKIPPED_BY_USER','quality_status':'NO_RESEARCH','bundle':{'enabled':False,'accepted':[],'quarantined':[],'agents':[]},'provider':None,'api_cost_usd':0.0}
            else:
                provider = _provider_for(project, inp)
                cc, reservation, est = _reserve_stage_if_paid(s, wid, job_id, kind, provider)
                if reservation == '__WAITING__':
                    return 0
                try:
                    progress({'phase': 'Research Agent · hledám kontext a zdroje', 'provider': provider, 'agent': 'Research Agent'})
                    from research_context import ResearchConfig, run_dual_research
                    cfg = ResearchConfig(
                        enabled=True,
                        topic=_research_topic(project),
                        max_sources_per_agent=10,
                        strict_consensus=True,
                        allow_single_agent_primary=True,
                        allow_degraded_single_agent=True,
                        max_context_blocks=18,
                    )
                    def _research_progress(text):
                        progress({'phase':str(text),'provider':provider,'agent':'Research Agent','research_stage':'background_research'})
                    bundle = run_dual_research(cfg, _compiled_questions(project), provider_override=provider, progress=_research_progress)
                    obj = asdict(bundle)
                    artifact = _write(wid, 'background_research', obj)
                    _settle_stage(cc, reservation, est)
                    out = {
                        'research_status': 'COMPLETED',
                        'quality_status': bundle.quality_status,
                        'bundle': obj,
                        'artifact': artifact,
                        'accepted_count': len(bundle.accepted),
                        'quarantined_count': len(bundle.quarantined),
                        'provider': provider,
                        'api_cost_usd': 0.0 if provider == 'claude_code_subscription' else est,
                    }
                except Exception as exc:
                    _cancel_stage(cc, reservation)
                    if provider == 'claude_code_subscription' and _is_claude_subscription_failure(exc):
                        raise RuntimeError(f'CLAUDE_CODE_STAGE_FAILED[{kind}]: {exc}') from exc
                        return 0
                    raise

        elif kind == 'project_stage_snapshot':
            sid=stage_id or str(inp.get('stage_id') or 'RESEARCH_DESIGN')
            selected={}
            if sid=='RESEARCH_DESIGN': selected={'research_plan':project.get('research_plan'),'goal':project.get('goal'),'decision_use':project.get('decision_use'),'study_type':project.get('study_type')}
            elif sid=='QUESTIONNAIRE': selected={'sections':project.get('sections'),'instrument_library':project.get('instrument_library'),'questionnaire_version':project.get('schema_version')}
            elif sid=='AUDIENCE': selected={'audience':project.get('audience')}
            elif sid=='DIMENSIONS': selected={'persona_mode':project.get('persona_mode'),'persona_dimensions':project.get('persona_dimensions')}
            elif sid=='SAMPLE_PLAN': selected={'n':project.get('n'),'panel_mode':project.get('panel_mode'),'discovery':project.get('discovery')}
            elif sid=='SCENARIO_CONTRACT': selected={'scenario_contract':project.get('scenario_contract'),'scenario':project.get('scenario'),'simulation':project.get('simulation')}
            elif sid=='BASELINE': selected={'baseline':project.get('baseline'),'goal':project.get('goal'),'decision_use':project.get('decision_use')}
            elif sid=='VARIANTS': selected={'variants':project.get('variants'),'n':project.get('n')}
            else: selected={'project':project}
            out={'stage':sid,'approved_snapshot':selected,'input_fingerprint':input_fp,'created_at':time.strftime('%Y-%m-%dT%H:%M:%S'),'local_deterministic':True}

        elif kind == 'project_preflight':
            from ui_server import project_preflight
            pf = project_preflight(project)
            out = pf
            blockers = [x for x in pf.get('issues', []) if x.get('level') == 'BLOCKER']
            if blockers:
                msg='Projekt vyžaduje úpravu před spuštěním: '+ '; '.join(str(x.get('user_message') or x.get('message') or x.get('code')) for x in blockers[:4])
                s.event(job_id,'PREFLIGHT_BLOCKED',msg,{'blockers':blockers[:20],'count':len(blockers)},'WARN')
                s.transition(job_id,'WAITING_USER',output={'preflight':pf,'blockers':blockers},error={'error':'PROJECT_PREFLIGHT_BLOCKED','blockers':blockers},message=msg)
                return 0

        elif kind == 'respondent_run':
            mode = str(inp.get('mode') or 'dry')
            confirm = bool(inp.get('confirm_live'))
            p = json.loads(json.dumps(project))
            provider = _provider_for(p, inp)
            p.setdefault('run_policy', {})['provider'] = provider
            p['run_policy']['allow_provider_fallback'] = False
            # Claude Code subscription is a valid production transport. The historical
            # parity lock remains only for switching to a cheaper respondent MODEL;
            # until parity passes, the normal quality model is retained.
            if provider == 'claude_code_subscription':
                p['run_policy']['run_purpose'] = 'production' if mode != 'dry' else 'testing'
                try:
                    from provider_parity import load_status
                    if not load_status().get('economy_respondent_default_allowed'):
                        p.setdefault('runtime_notes', {})['cheap_respondent_model'] = 'LOCKED_PARITY_NOT_RUN_USING_REFERENCE_QUALITY'
                except Exception:
                    pass

            reserve = None
            cc = CostController(s, wid)
            if mode != 'dry' and provider == 'anthropic':
                from research_project import compile_project
                from cost_estimator import estimate_range
                brief = compile_project(p)['brief']
                est_obj = estimate_range(brief)
                amount = float(est_obj.get('usd_high') or est_obj.get('usd_typical') or 0)
                try:
                    reserve = cc.reserve(job_id, amount)
                except WorkflowBudgetExceeded as exc:
                    s.create_approval(
                        job_id,
                        'Workflow budget nestačí pro placený respondentní běh.',
                        [{'option': 'increase_budget'}, {'option': 'reduce_n'}, {'option': 'continue_later'}, {'option': 'cancel'}],
                        {'reason': str(exc), 'new_cost_estimate_usd': amount},
                    )
                    return 0

            run_dir = ROOT / 'runs' / f'research_os_{wid}'
            resume = run_dir if (run_dir / 'manifest.json').is_file() and (run_dir / 'checkpoint.pkl').is_file() else None
            research_obj = (_dep_output(s, job, 'research').get('bundle') or None)
            from project_engine import run_project
            try:
                progress({'phase': 'Synthetic fieldwork · respondenti', 'provider': provider, 'agent': 'Respondent Engine'})
                r = run_project(
                    p,
                    mode=mode,
                    confirm_live=confirm,
                    run_dir=run_dir,
                    resume_dir=resume,
                    progress_callback=progress,
                    cancel_check=cancel,
                    workflow_id=wid,
                    job_id=job_id,
                    precomputed_research=research_obj,
                )
            except Exception as exc:
                if reserve:
                    cc.cancel(reserve)
                # Subscription respondent fieldwork is idempotent at respondent level:
                # every completed case is journaled before parsing and the run has an
                # initial checkpoint even on question 1. Transient transport/watchdog
                # failures therefore retry automatically from the exact durable point.
                if provider == 'claude_code_subscription' and _is_subscription_quota_failure(exc):
                    # The fieldwork branch handles quota locally in order to retain
                    # respondent journals. Keep the project source-of-truth in sync
                    # before parking the durable job.
                    try:
                        if _ps and _astore and project_id and project_revision and stage_id:
                            from project_artifact_sync import sync_fieldwork
                            sync_fieldwork(_astore,project_id=project_id,revision=project_revision,input_fingerprint=input_fp,provider=provider,model=str(job.get('model') or ''),job_id=job_id,workflow_id=wid,run_dir=run_dir)
                            reset_at,_raw=_quota_reset_at(exc)
                            _ps.set_stage(project_id,project_revision,stage_id,'WAITING_CREDITS',input_fingerprint=input_fp,current_job_id=job_id,last_checkpoint=str(run_dir),waiting_reason=str(exc)[:1000],quota_reset_at=reset_at)
                    except Exception:
                        pass
                    return _park_for_subscription_quota(s,wid,job_id,exc,run_dir=run_dir)
                if provider == 'claude_code_subscription' and _respondent_retryable_failure(exc) and int(job.get('attempt') or 0) < int(job.get('max_attempts') or 3):
                    msg=f'Respondent Engine · dočasná chyba, pokračuji z checkpointu ({int(job.get("attempt") or 0)}/{int(job.get("max_attempts") or 3)})'
                    s.event(job_id,'RESPONDENT_AUTO_RETRY',msg,{'error':str(exc)[:1200],'run_dir':str(run_dir),'resume_available':(run_dir/'checkpoint.pkl').is_file()},'WARN')
                    s.transition(job_id,'RETRYING',error={'error':str(exc),'recoverable':True,'run_dir':str(run_dir)},message=msg)
                    s.queue_ready_jobs(wid)
                    return 0
                if provider == 'claude_code_subscription' and _is_claude_subscription_failure(exc):
                    raise RuntimeError(f'CLAUDE_CODE_STAGE_FAILED[{kind}]: {exc}') from exc
                raise
            actual = float(((r.get('main') or {}).get('summary') or {}).get('naklady_usd') or 0)
            if reserve:
                cc.settle(reserve, actual)
            out = {
                'result': r,
                'artifact': _write(wid, 'respondent_result', r),
                'resume_used': bool(resume),
                'provider': provider,
                'actual_cost_usd': actual,
                'research_context_reused': bool(research_obj),
            }

        elif kind == 'aggregate_results':
            r = _dep_output(s, job, 'run').get('result') or {}
            summary = ((r.get('main') or {}).get('summary') or {})
            out = {
                'summary': summary,
                'batteries': r.get('batteries') or [],
                'artifact': _write(wid, 'aggregate', {'summary': summary, 'batteries': r.get('batteries') or []}),
            }

        elif kind == 'donor_qc':
            agg = _dep_output(s, job, 'aggregate')
            summary = agg.get('summary') or {}
            ds = summary.get('donor_support') or {}
            out = {
                'donor_support': ds,
                'support_status': ds.get('support_status') or 'NOT_AVAILABLE',
                'passed': str(ds.get('support_status') or '').upper() not in {'SUPPRESSED', 'FAIL'},
            }

        elif kind == 'analysis_module':
            prior = _all_prior_outputs(s, wid)
            summary = (prior.get('aggregate') or {}).get('summary') or {}
            research = (prior.get('research') or {}).get('bundle') or {}
            donor = (prior.get('donor_qc') or {}).get('donor_support') or {}
            module = str(inp.get('analysis_module') or '').strip()
            mode = str(inp.get('mode') or 'dry')
            provider = _provider_for(project, inp)
            if mode == 'dry':
                module_result={'module':module,'content':{},'provider':None,'model':None,'tok_in':0,'tok_out':0,'created_at':dt.datetime.now().isoformat()}
                out={'analysis_module_status':'NOT_RUN_DRY','analysis_module':module_result,'provider':'','api_cost_usd':0.0,'input_tokens':0,'output_tokens':0}
            else:
                cc, reservation, est = _reserve_stage_if_paid(s, wid, job_id, kind, provider)
                if reservation == '__WAITING__': return 0
                try:
                    progress({'phase':f'Analysis Agent · modul {module}', 'provider':provider, 'agent':'Analysis Agent', 'analysis_module':module})
                    from analysis_agent import analyze_module
                    module_result=analyze_module(module,project,summary,research_bundle=research,donor_support=donor,battery_results=(prior.get('aggregate') or {}).get('batteries') or [],provider_override=provider,model_override=_model_for(project,inp))
                    _settle_stage(cc,reservation,est)
                    out={'analysis_module_status':'COMPLETED','analysis_module':module_result,'provider':module_result.get('provider') or provider,'model':module_result.get('model') or _model_for(project,inp),'api_cost_usd':0.0 if provider=='claude_code_subscription' else est,'input_tokens':int(module_result.get('tok_in') or 0),'output_tokens':int(module_result.get('tok_out') or 0)}
                except Exception:
                    _cancel_stage(cc,reservation); raise

        elif kind == 'analysis_assemble':
            prior=_all_prior_outputs(s,wid)
            mode=str(inp.get('mode') or 'dry')
            summary=(prior.get('aggregate') or {}).get('summary') or {}
            if mode=='dry':
                analysis={'executive_answer':'Dry/demo workflow nevolá Analysis Agenta.','research_question_answers':[],'key_findings':[],'implications':[],'limitations':[],'_meta':{'provider':None,'analysis_mode':'DRY_NOT_RUN'}}
                out={'analysis_status':'NOT_RUN_DRY','analysis':analysis,'api_cost_usd':0.0,'provider':''}
            else:
                modules={}
                for node_key,node_out in prior.items():
                    if not str(node_key).startswith('analysis_'): continue
                    m=(node_out or {}).get('analysis_module') or {}
                    if isinstance(m,dict) and m.get('module'): modules[str(m['module'])]=m
                from analysis_agent import MODULE_ORDER, assemble_modules
                missing=[m for m in MODULE_ORDER if m not in modules]
                if missing: raise RuntimeError('ANALYSIS_MODULES_MISSING: '+','.join(missing))
                analysis=assemble_modules(modules,project,summary,battery_results=(prior.get('aggregate') or {}).get('batteries') or [])
                out={'analysis_status':'COMPLETED','analysis':analysis,'artifact':_write(wid,'analysis',analysis),'provider':(analysis.get('_meta') or {}).get('provider') or 'mixed','model':(analysis.get('_meta') or {}).get('model') or 'mixed','api_cost_usd':0.0,'module_count':len(modules)}

        elif kind == 'interpret_results':
            prior = _all_prior_outputs(s, wid)
            summary = (prior.get('aggregate') or {}).get('summary') or {}
            research = (prior.get('research') or {}).get('bundle') or {}
            donor = (prior.get('donor_qc') or {}).get('donor_support') or {}
            mode = str(inp.get('mode') or 'dry')
            if mode == 'dry':
                # Demo/dry artifacts remain precomputed/local; no fake AI claim.
                out = {
                    'analysis_status': 'NOT_RUN_DRY',
                    'analysis': {'executive_answer': 'Dry/demo workflow nevolá Analysis Agenta.', 'research_question_answers': [], 'key_findings': [], 'implications': [], '_meta': {'provider': None}},
                    'api_cost_usd': 0.0,
                }
            else:
                provider = _provider_for(project, inp)
                cc, reservation, est = _reserve_stage_if_paid(s, wid, job_id, kind, provider)
                if reservation == '__WAITING__':
                    return 0
                try:
                    progress({'phase': 'Analysis Agent · odpovídám na výzkumné otázky', 'provider': provider, 'agent': 'Analysis Agent'})
                    from analysis_agent import analyze_results
                    analysis = analyze_results(project, summary, research_bundle=research, donor_support=donor, battery_results=(prior.get('aggregate') or {}).get('batteries') or [], provider_override=provider, model_override=_model_for(project, inp), checkpoint_dir=ROOT/'prototype_outputs'/'workflows'/wid/'analysis_checkpoints', cost_mode=inp.get('cost_mode'))
                    artifact = _write(wid, 'analysis', analysis)
                    _settle_stage(cc, reservation, est)
                    out = {
                        'analysis_status': 'COMPLETED',
                        'analysis': analysis,
                        'artifact': artifact,
                        'provider': provider,
                        'api_cost_usd': 0.0 if provider == 'claude_code_subscription' else est,
                    }
                except Exception as exc:
                    _cancel_stage(cc, reservation)
                    if provider == 'claude_code_subscription' and _is_claude_subscription_failure(exc):
                        raise RuntimeError(f'CLAUDE_CODE_STAGE_FAILED[{kind}]: {exc}') from exc
                        return 0
                    raise

        elif kind == 'external_verification':
            if not inp.get('verification_requested'):
                out = {'verification_status': 'NOT_RUN', 'by_target': {}, 'findings': [], 'api_cost_usd': 0.0}
            else:
                prior = _all_prior_outputs(s, wid)
                summary = (prior.get('aggregate') or {}).get('summary') or {}
                provider = _provider_for(project, inp)
                cc, reservation, est = _reserve_stage_if_paid(s, wid, job_id, kind, provider)
                if reservation == '__WAITING__':
                    return 0
                try:
                    progress({'phase': 'Validation Agent · porovnávám s externí realitou', 'provider': provider, 'agent': 'Validation Agent'})
                    from result_context import verify_results
                    verification = verify_results(project, summary, include_mapped=False, max_sources_per_agent=8, provider_override=provider)
                    artifact = _write(wid, 'external_verification', verification)
                    _settle_stage(cc, reservation, est)
                    out = {
                        'verification_status': verification.get('quality_status') or 'COMPLETED',
                        'verification': verification,
                        'artifact': artifact,
                        'by_target': verification.get('by_target') or {},
                        'findings': verification.get('findings') or [],
                        'provider': provider,
                        'api_cost_usd': 0.0 if provider == 'claude_code_subscription' else est,
                    }
                except ValueError as exc:
                    # Some studies have no comparable numeric target. This is a valid, explicit state.
                    if 'žádný číselný výsledek' in str(exc).lower() or 'zadny ciselny' in str(exc).lower():
                        _cancel_stage(cc, reservation)
                        out = {'verification_status': 'NO_COMPARABLE_TARGET', 'by_target': {}, 'findings': [], 'note': str(exc), 'api_cost_usd': 0.0}
                    else:
                        _cancel_stage(cc, reservation)
                        raise
                except Exception as exc:
                    _cancel_stage(cc, reservation)
                    if provider == 'claude_code_subscription' and _is_claude_subscription_failure(exc):
                        raise RuntimeError(f'CLAUDE_CODE_STAGE_FAILED[{kind}]: {exc}') from exc
                        return 0
                    raise

        elif kind == 'reality_alignment':
            prior = _all_prior_outputs(s, wid)
            summary = (prior.get('aggregate') or {}).get('summary') or {}
            analysis = (prior.get('interpret') or {}).get('analysis') or {}
            research = (prior.get('research') or {}).get('bundle') or {}
            verify_out = prior.get('verify') or {}
            verification = verify_out.get('verification') or verify_out
            mode = str(inp.get('mode') or 'dry')
            if mode == 'dry':
                out = {'alignment_status':'NOT_RUN_DRY','api_cost_usd':0.0}
            else:
                provider = _provider_for(project, inp)
                cc, reservation, est = _reserve_stage_if_paid(s, wid, job_id, kind, provider)
                if reservation == '__WAITING__':
                    return 0
                try:
                    progress({'phase':'Reality Alignment Agent · porovnávám research, výsledek a realitu','provider':provider,'agent':'Reality Alignment Agent'})
                    from reality_alignment import compare, build_calibration_profile, export_profile, export_alignment_html
                    alignment = compare(project, summary, analysis, research, verification, provider_override=provider)
                    profile = build_calibration_profile(alignment, project=project, workflow_id=wid)
                    progress({'phase':'Reality Alignment Agent · AI porovnání hotové, exportuji kalibrační profil','provider':provider,'agent':'Reality Alignment Agent','alignment_stage':'content_ready'})
                    out_dir = ROOT / 'prototype_outputs' / 'final_reports' / wid
                    out_dir.mkdir(parents=True, exist_ok=True)
                    jp, cp = export_profile(profile, out_dir / f'{wid}_AI_PANEL_CALIBRATION.json', out_dir / f'{wid}_AI_PANEL_CALIBRATION.csv')
                    ah = export_alignment_html(out_dir / f'{wid}_REALITY_ALIGNMENT.html', project, alignment, profile)
                    artifact = _write(wid, 'reality_alignment', alignment)
                    _settle_stage(cc, reservation, est)
                    out = {'alignment_status':'COMPLETED','alignment':alignment,'artifact':artifact,'alignment_html':ah,
                           'calibration_profile':profile,'calibration_profile_json':jp,'calibration_profile_csv':cp,
                           'provider':provider,'api_cost_usd':0.0 if provider=='claude_code_subscription' else est}
                except Exception as exc:
                    _cancel_stage(cc, reservation)
                    if provider == 'claude_code_subscription' and _is_claude_subscription_failure(exc):
                        raise RuntimeError(f'CLAUDE_CODE_STAGE_FAILED[{kind}]: {exc}') from exc
                        return 0
                    raise

        elif kind == 'final_report':
            prior = _all_prior_outputs(s, wid)
            missing=[k for k in ('research','interpret','verify','alignment') if not (prior.get(k) or {})]
            if missing: raise RuntimeError('REPORT_WORKFLOW_PRECONDITION_MISSING: '+','.join(missing))
            agg = prior.get('aggregate') or {}
            summary = agg.get('summary') or {}
            analysis = (prior.get('interpret') or {}).get('analysis') or {}
            if not analysis: raise RuntimeError('REPORT_WORKFLOW_PRECONDITION_MISSING: interpret.analysis')
            research = (prior.get('research') or {}).get('bundle') or {}
            verify_out = prior.get('verify') or {}
            verification = verify_out.get('verification') or verify_out
            donor = (prior.get('donor_qc') or {}).get('donor_support') or {}
            alignment_out = prior.get('alignment') or {}
            alignment = alignment_out.get('alignment') or {}
            run_result = (prior.get('run') or {}).get('result') or {}
            run_main = run_result.get('main') or {}
            def _local_client_file(v):
                if not v: return None
                sv=str(v)
                if sv.startswith('/files/'):
                    q=ROOT/'prototype_outputs'/sv[len('/files/'):]
                    return str(q) if q.is_file() else None
                q=Path(sv)
                return str(q) if q.is_file() else None
            respondent_dataset_csv=_local_client_file(run_main.get('dataset_csv'))
            respondent_results_xlsx=_local_client_file(run_main.get('xlsx'))
            respondent_run_html=_local_client_file(run_main.get('html'))
            maps = _collect_maps(prior)
            mode = str(inp.get('mode') or 'dry')

            if mode == 'dry':
                out = {
                    'report_status': 'NOT_RUN_DRY',
                    'note': 'Dry/demo používá předpočítané showcase artefakty a nevydává AI report jako LIVE.',
                    'maps': maps,
                }
            else:
                provider = _provider_for(project, inp)
                cc, reservation, est = _reserve_stage_if_paid(s, wid, job_id, kind, provider)
                if reservation == '__WAITING__':
                    return 0
                try:
                    progress({'phase': 'Client Report Agent · skládám finální report', 'provider': provider, 'agent': 'Client Report Agent'})
                    from client_report_v2 import compose_report, export_docx, export_html
                    report = compose_report(
                        project,
                        summary,
                        analysis,
                        research=research,
                        verification=verification,
                        donor_support=donor,
                        reality_alignment=alignment,
                        battery_results=agg.get('batteries') or [], provider_override=provider, model_override=_model_for(project, inp, 'sonnet'), checkpoint_dir=ROOT/'prototype_outputs'/'workflows'/wid/'report_checkpoints',
                    )
                    out_dir = ROOT / 'prototype_outputs' / 'final_reports' / wid
                    out_dir.mkdir(parents=True,exist_ok=True)
                    docx = out_dir / f'{wid}_FINAL_CLIENT_REPORT.docx'
                    html = out_dir / f'{wid}_FINAL_CLIENT_REPORT.html'
                    js = out_dir / f'{wid}_FINAL_CLIENT_REPORT.json'
                    artifact_warnings=[]
                    # Persist the machine-readable report FIRST. From this point on,
                    # secondary rendering failures may degrade presentation but may
                    # not erase valid analytical results.
                    js.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
                    progress({'phase':'Výsledky · analytické jádro uloženo, generuji klientské artefakty','agent':'Report Export','report_stage':'results_core_saved'})
                    try:
                        export_docx(docx, report, project=project, maps=maps)
                    except Exception as exc:
                        artifact_warnings.append({'artifact':'client_report_docx','error':str(exc)[:800]})
                        _fallback_report_docx(docx,report,project)
                    progress({'phase':'Report export · DOCX hotový, vytvářím HTML','agent':'Report Export','report_stage':'docx_done'})
                    try:
                        export_html(html, report, project=project, maps=maps)
                    except Exception as exc:
                        artifact_warnings.append({'artifact':'client_report_html','error':str(exc)[:800]})
                        _fallback_report_html(html,report,project)
                    progress({'phase':'Report export · HTML hotový, připravuji další výstupy','agent':'Report Export','report_stage':'html_done'})
                    from output_pack import export_executive_readout, export_research_brief, export_analysis_readout, export_validation_readout, export_questionnaire_html, export_internal_report, export_handoff_protocol, export_management_deck_html, export_management_deck_pptx, export_evidence_pack, export_verbatims, build_output_manifest, export_client_delivery_zip
                    def _optional(label, fn, default=None):
                        try:return fn()
                        except Exception as exc:
                            artifact_warnings.append({'artifact':label,'error':str(exc)[:800]})
                            progress({'phase':f'Report export · {label} odložen, pokračuji','agent':'Report Export','report_stage':'optional_degraded','artifact':label,'level':'WARN'})
                            return default
                    executive = _optional('executive_readout',lambda:export_executive_readout(out_dir / f'{wid}_EXECUTIVE_READOUT.html', project, report, maps=maps))
                    if not executive: executive=_fallback_report_html(out_dir/f'{wid}_EXECUTIVE_READOUT.html',report,project)
                    research_brief = _optional('research_brief',lambda:export_research_brief(out_dir / f'{wid}_RESEARCH_BRIEF.html', project, research))
                    analysis_html = _optional('analysis_html',lambda:export_analysis_readout(out_dir / f'{wid}_ANALYSIS.html', project, analysis))
                    validation_html = _optional('validation_html',lambda:export_validation_readout(out_dir / f'{wid}_VALIDATION.html', project, verification))
                    questionnaire_html=_optional('questionnaire_html',lambda:export_questionnaire_html(out_dir/f'{wid}_QUESTIONNAIRE.html',project))
                    if not questionnaire_html:
                        qpath=out_dir/f'{wid}_QUESTIONNAIRE.html';qpath.write_text('<!doctype html><meta charset="utf-8"><h1>Dotazník</h1><pre>'+json.dumps(project.get('sections') or project.get('questions') or [],ensure_ascii=False,indent=2,default=str)+'</pre>',encoding='utf-8');questionnaire_html=str(qpath)
                    management_deck=_optional('management_deck_html',lambda:export_management_deck_html(out_dir/f'{wid}_MANAGEMENT_DECK.html',project,report,maps=maps))
                    management_deck_pptx=_optional('management_deck_pptx',lambda:export_management_deck_pptx(out_dir/f'{wid}_MANAGEMENT_DECK.pptx',project,report,maps=maps))
                    progress({'phase':'Report export · hlavní report hotový, balím evidence a data','agent':'Report Export','report_stage':'core_reports_done'})
                    evidence_pack=_optional('evidence_pack',lambda:export_evidence_pack(out_dir/f'{wid}_EVIDENCE_PACK.json',project,analysis,research,verification,report))
                    verbatims=_optional('verbatims',lambda:export_verbatims(out_dir/f'{wid}_VERBATIMS.csv',out_dir/f'{wid}_VERBATIMS.html',project,summary),default={'csv':None,'html':None,'count':0}) or {'csv':None,'html':None,'count':0}
                    internal_report=_optional('internal_report',lambda:export_internal_report(out_dir/f'{wid}_INTERNAL_REPORT.html',project,report,analysis,research,verification,alignment,workflow_id=wid))
                    handoff_protocol=_optional('handoff_protocol',lambda:export_handoff_protocol(out_dir/f'{wid}_HANDOFF_PROTOCOL.html',project,report,workflow_id=wid))
                    learning_cycle={}; learning_file=None
                    try:
                        progress({'phase':'Model Learning · hledám nové externě podložené mezery','provider':provider,'agent':'Learning Curator'})
                        from data_library import learn_from_completed_workflow
                        learning_cycle=learn_from_completed_workflow(workflow_id=wid,project=project,research=research,verification=verification,model=_model_for(project, inp, 'sonnet'))
                        learning_file=out_dir/f'{wid}_MODEL_LEARNING.json'
                        learning_file.write_text(json.dumps(learning_cycle,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
                    except Exception as learn_exc:
                        learning_cycle={'status':'DEFERRED','summary':'Učení modelu bylo odloženo; report a výsledky jsou dokončené.','technical_reason':str(learn_exc)[:800]}
                        artifact_warnings.append({'artifact':'model_learning','error':str(learn_exc)[:800]})
                    progress({'phase':'Report export · dokončuji manifest výstupů','agent':'Report Export','report_stage':'manifest'})
                    manifest_files={
                        'executive_readout': executive, 'research_brief': research_brief, 'analysis': analysis_html,
                        'validation': validation_html, 'reality_alignment': alignment_out.get('alignment_html'),
                        'ai_panel_calibration_json': alignment_out.get('calibration_profile_json'), 'ai_panel_calibration_csv': alignment_out.get('calibration_profile_csv'),
                        'client_report_docx': str(docx), 'client_report_html': str(html), 'client_report_json': str(js), 'management_deck_pptx':management_deck_pptx,'management_deck_html':management_deck,'questionnaire_html':questionnaire_html,'respondent_dataset_csv':respondent_dataset_csv,'respondent_results_xlsx':respondent_results_xlsx,'respondent_run_html':respondent_run_html,'evidence_pack_json':evidence_pack,'verbatims':verbatims,'internal_report':internal_report,'handoff_protocol':handoff_protocol,'model_learning_json':str(learning_file) if learning_file else None,'maps': maps,
                    }
                    try: manifest = build_output_manifest(out_dir / f'{wid}_OUTPUT_MANIFEST.json', manifest_files, project=project, workflow_id=wid)
                    except Exception as exc:
                        artifact_warnings.append({'artifact':'output_manifest','error':str(exc)[:800]})
                        manifest=_fallback_manifest(out_dir/f'{wid}_OUTPUT_MANIFEST.json',manifest_files,project,wid,artifact_warnings)
                    delivery_files={'client_report_docx':str(docx),'client_report_html':str(html),'client_report_json':str(js),'executive_readout':executive,
                        'management_deck_pptx':management_deck_pptx,'management_deck_html':management_deck,
                        'questionnaire_html':questionnaire_html,'respondent_dataset_csv':respondent_dataset_csv,
                        'respondent_results_xlsx':respondent_results_xlsx,'respondent_run_html':respondent_run_html,
                        'verbatims_csv':verbatims.get('csv'),'verbatims_html':verbatims.get('html'),'output_manifest':manifest}
                    try: client_delivery_zip=export_client_delivery_zip(out_dir/f'{wid}_CLIENT_DELIVERY.zip',delivery_files,project=project,workflow_id=wid)
                    except Exception as exc:
                        artifact_warnings.append({'artifact':'client_delivery_zip','error':str(exc)[:800]})
                        import zipfile,hashlib
                        zp=out_dir/f'{wid}_CLIENT_DELIVERY.zip'
                        with zipfile.ZipFile(zp,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=6) as z:
                            checks=[]
                            for label,value in delivery_files.items():
                                if not value:continue
                                q=Path(str(value))
                                if q.is_file():z.write(q,q.name);checks.append(f'{hashlib.sha256(q.read_bytes()).hexdigest()}  {q.name}')
                            z.writestr('DELIVERY_SHA256SUMS.txt','\n'.join(checks)+'\n')
                        client_delivery_zip=str(zp)
                    try:
                        _manifest_obj=json.loads(Path(manifest).read_text(encoding='utf-8'))
                        _manifest_obj.setdefault('outputs',{})['client_delivery_zip']=client_delivery_zip
                        _manifest_obj['artifact_warnings']=artifact_warnings
                        Path(manifest).write_text(json.dumps(_manifest_obj,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
                    except Exception: pass
                    progress({'phase':'Report export · klientský delivery ZIP hotový','agent':'Report Export','report_stage':'delivery_zip'})
                    # Client-ready means physical deliverables exist, not merely that
                    # the orchestration function returned without an exception.
                    _required_client_artifacts={
                        'client_report_docx':str(docx),'client_report_html':str(html),'client_report_json':str(js),
                        'executive_readout':executive,'questionnaire_html':questionnaire_html,
                        'respondent_dataset_csv':respondent_dataset_csv,'respondent_results_xlsx':respondent_results_xlsx,
                        'output_manifest':manifest,'client_delivery_zip':client_delivery_zip,
                    }
                    _optional_client_artifacts={'management_deck_pptx':management_deck_pptx,'management_deck_html':management_deck,
                                                'research_brief_html':research_brief,'evidence_pack_json':evidence_pack,
                                                'verbatims_csv':verbatims.get('csv'),'verbatims_html':verbatims.get('html')}
                    _missing_client=[]
                    for _name,_value in _required_client_artifacts.items():
                        try:
                            if not _value or not Path(str(_value)).is_file() or Path(str(_value)).stat().st_size<=0:
                                _missing_client.append(_name)
                        except Exception:
                            _missing_client.append(_name)
                    _optional_missing=[]
                    for _name,_value in _optional_client_artifacts.items():
                        try:
                            if not _value or not Path(str(_value)).is_file() or Path(str(_value)).stat().st_size<=0:_optional_missing.append(_name)
                        except Exception:_optional_missing.append(_name)
                    if _missing_client:
                        progress({'phase':'Report export · core artifact gate FAILED','agent':'Report Export','report_stage':'artifact_gate_failed','missing':_missing_client,'level':'ERROR'})
                        raise RuntimeError('CLIENT_DELIVERY_ARTIFACT_GATE_FAILED: '+','.join(_missing_client))
                    progress({'phase':'Report export · core artifact gate PASS','agent':'Report Export','report_stage':'artifact_gate_pass','artifact_count':len(_required_client_artifacts),'optional_missing':_optional_missing,'level':'WARN' if _optional_missing else 'INFO'})
                    _settle_stage(cc, reservation, est)
                    out = {
                        # NPC AI RUNTIME FIX: artefakty vzniknou vzdy; neprosla
                        # kvalitativni brana se propise do stavu, ne do ztraty behu.
                        'report_status': ('COMPLETED_DEGRADED_EXPORT' if artifact_warnings else ('COMPLETED_DEGRADED_FORMATTING' if (report.get('_meta') or {}).get('report_generation_mode')=='DETERMINISTIC_FROM_VALIDATED_ANALYSIS' else ('COMPLETED' if ((report.get('_meta') or {}).get('quality_gate') or {}).get('passed') else 'COMPLETED_REVIEW_REQUIRED'))),
                        'report': report,
                        'executive_readout': executive,
                        'research_brief_html': research_brief,
                        'analysis_html': analysis_html,
                        'validation_html': validation_html,
                        'reality_alignment': alignment, 'reality_alignment_html': alignment_out.get('alignment_html'),
                        'calibration_profile': alignment_out.get('calibration_profile'),
                        'calibration_profile_json': alignment_out.get('calibration_profile_json'), 'calibration_profile_csv': alignment_out.get('calibration_profile_csv'),
                        'output_manifest': manifest,
                        'report_docx': str(docx),
                        'report_html': str(html),
                        'report_json': str(js), 'management_deck_pptx':management_deck_pptx,'management_deck_html':management_deck,'questionnaire_html':questionnaire_html,'respondent_dataset_csv':respondent_dataset_csv,'respondent_results_xlsx':respondent_results_xlsx,'respondent_run_html':respondent_run_html,'client_delivery_zip':client_delivery_zip,'evidence_pack_json':evidence_pack,'verbatims_csv':verbatims.get('csv'),'verbatims_html':verbatims.get('html'),'internal_report_html':internal_report,'handoff_protocol_html':handoff_protocol,'learning_cycle':learning_cycle,'model_learning_json':str(learning_file) if learning_file else None,
                        'quality_gate': (report.get('_meta') or {}).get('quality_gate') or {},'artifact_warnings':artifact_warnings+([{'artifact':x,'error':'missing optional output'} for x in _optional_missing]),
                        'maps': maps,
                        'verification_status': verify_out.get('verification_status', 'NOT_RUN'),
                        'provider': provider,
                        'api_cost_usd': 0.0 if provider == 'claude_code_subscription' else est,
                    }
                except Exception as exc:
                    _cancel_stage(cc, reservation)
                    if provider == 'claude_code_subscription' and _is_claude_subscription_failure(exc):
                        raise RuntimeError(f'CLAUDE_CODE_STAGE_FAILED[{kind}]: {exc}') from exc
                        return 0
                    raise

        elif kind == 'project_delivery_finalize':
            prior=_all_prior_outputs(s,wid); rep=prior.get('report') or {}
            paths=[]
            def _walk(x):
                if isinstance(x,dict):
                    for v in x.values():_walk(v)
                elif isinstance(x,list):
                    for v in x:_walk(v)
                elif isinstance(x,str):
                    try:
                        pp=Path(x)
                        if pp.is_file(): paths.append(str(pp))
                    except Exception: pass
            _walk(rep)
            out={'delivery_status':'COMPLETED','report_status':rep.get('report_status'),'files':sorted(set(paths)),'file_count':len(set(paths)),'source_job':'report','results_guarantee':True}

        elif kind == 'condition_gate':
            from task_conditions import evaluate
            out = {'condition_met': evaluate(inp['condition'], inp.get('metrics') or {})}

        elif kind == 'legacy_task':
            payload = json.loads(json.dumps(inp.get('payload') or {}))
            legacy_kind = str(inp.get('legacy_kind') or 'run')
            override = _approved_provider(inp)
            if override:
                payload['provider'] = override
                if isinstance(payload.get('project'), dict):
                    payload['project'].setdefault('run_policy', {})['provider'] = override
                    payload['project']['run_policy']['allow_provider_fallback'] = False
                if isinstance(payload.get('spec'), dict):
                    payload['spec']['provider'] = override

            def legacy_phase(text):
                progress({'phase': str(text), 'legacy_kind': legacy_kind})

            from legacy_job_dispatch import execute_legacy
            try:
                obj = execute_legacy(legacy_kind, payload, job_id, legacy_phase)
            except Exception as exc:
                from ai_execution_context import is_cancel_exception
                if is_cancel_exception(exc):
                    raise
                provider = str(
                    payload.get('provider')
                    or (((payload.get('project') or {}).get('run_policy') or {}).get('provider'))
                    or ((payload.get('spec') or {}).get('provider'))
                    or ''
                )
                if provider == 'claude_code_subscription' and _is_claude_subscription_failure(exc):
                    estimates = {}
                    pr = payload.get('project') or {}
                    if pr:
                        try:
                            from research_project import compile_project
                            from cost_estimator import estimate_range
                            for prov in ('anthropic',):
                                pp = json.loads(json.dumps(pr))
                                pp.setdefault('run_policy', {})['provider'] = prov
                                pp['run_policy']['allow_provider_fallback'] = False
                                estimates[prov] = float(estimate_range(compile_project(pp)['brief']).get('usd_typical') or 0)
                        except Exception:
                            pass
                    if estimates:
                        raise RuntimeError(f'CLAUDE_CODE_STAGE_FAILED[{legacy_kind}]: {exc}') from exc
                    else:
                        raise RuntimeError(f'CLAUDE_CODE_STAGE_FAILED[{legacy_kind}]: {exc}') from exc
                    return 0
                raise
            out = {'result': obj, 'legacy_kind': legacy_kind, 'durable': True, 'provider': override or payload.get('provider')}

        else:
            raise ValueError(f'Unknown job kind: {kind}')

        if cancel():
            if _ps and project_id and project_revision and stage_id:
                _ps.set_stage(project_id,project_revision,stage_id,'READY',input_fingerprint=input_fp,waiting_reason='cancelled_by_user',clear_current_job=True)
            s.transition(job_id, 'CANCELLED', output=out, message='cancelled cooperatively')
        else:
            if _astore and _ps and project_id and project_revision and stage_id:
                provider=str(out.get('provider') or job.get('provider_policy') or '')
                model=str(out.get('model') or job.get('model') or '')
                # Project the proven engine-native checkpoints into granular durable
                # project artifacts before the stage is declared complete.
                from project_artifact_sync import sync_for_job
                sync_for_job(_astore,kind=kind,out=out,project_id=project_id,revision=project_revision,stage_id=stage_id,input_fingerprint=input_fp,provider=provider,model=model,job_id=job_id,workflow_id=wid)
                art=_astore.put_json(project_id=project_id,revision=project_revision,stage_type=stage_id,artifact_type=artifact_target,obj=out,filename=f'{artifact_target}.json',input_fingerprint=input_fp,provider=provider,model=model,metadata={'job_id':job_id,'workflow_id':wid,'kind':kind})
                out['_artifact_id']=art['artifact_id'];out['_artifact_sha256']=art['sha256']
                warn=bool(out.get('artifact_warnings')) or str(out.get('report_status') or '').startswith('COMPLETED_DEGRADED')
                # Eight analysis modules are separate durable jobs within one
                # logical ANALYSIS stage. A module completion is a checkpoint,
                # not completion of the whole stage; only deterministic assembly
                # advances the project to REPORT.
                if kind=='analysis_module':
                    mod=((out.get('analysis_module') or {}).get('module') or inp.get('analysis_module') or '')
                    _ps.set_stage(project_id,project_revision,stage_id,'RUNNING',input_fingerprint=input_fp,provider=provider,model=model,current_job_id=job_id,last_checkpoint=f'analysis_module:{mod}',artifacts=[art['artifact_id']])
                else:
                    _ps.set_stage(project_id,project_revision,stage_id,'DONE_WITH_WARNINGS' if warn else 'DONE',input_fingerprint=input_fp,provider=provider,model=model,current_job_id=job_id,artifacts=[art['artifact_id']])
                cost=float(out.get('api_cost_usd') or 0); tin=int(out.get('input_tokens') or 0); tout=int(out.get('output_tokens') or 0)
                if provider=='anthropic' and (cost or tin or tout):
                    _ps.record_provider_event(project_id,revision=project_revision,stage_type=stage_id,to_provider='anthropic',reason='stage_completed',input_tokens=tin,output_tokens=tout,estimated_cost_usd=cost,actual_cost_usd=cost,metadata={'job_id':job_id,'artifact_id':art['artifact_id'],'kind':kind,'analysis_module':inp.get('analysis_module')})
            s.transition(job_id, 'COMPLETED', output=out)
        s.queue_ready_jobs(wid)
        if _ps:_ps.close()
        return 0

    except Exception as exc:
        from ai_execution_context import is_cancel_exception
        if is_cancel_exception(exc):
            try:
                if locals().get('_ps') and project_id and project_revision and stage_id:
                    _ps.set_stage(project_id,project_revision,stage_id,'READY',input_fingerprint=input_fp,waiting_reason='cancelled_by_user',clear_current_job=True)
            except Exception:
                pass
            s.transition(job_id, 'CANCELLED', error={'error':'JOB_CANCELLED'}, message='cancelled cooperatively')
            if locals().get('_ps'):
                try:_ps.close()
                except Exception:pass
            return 0
        # 17.8.8: a Claude Pro session/usage limit is not a workflow failure.
        # Park every AI stage durably and continue automatically after provider reset.
        if _is_subscription_quota_failure(exc):
            _rd=ROOT/'runs'/f'research_os_{wid}' if kind=='respondent_run' else None
            try:
                if locals().get('_ps') and project_id and project_revision and stage_id:
                    if kind=='respondent_run' and locals().get('_astore') and _rd and _rd.is_dir():
                        try:
                            from project_artifact_sync import sync_fieldwork
                            sync_fieldwork(_astore,project_id=project_id,revision=project_revision,input_fingerprint=input_fp,provider='claude_code_subscription',model=str(job.get('model') or ''),job_id=job_id,workflow_id=wid,run_dir=_rd)
                        except Exception: pass
                    reset_at,_raw=_quota_reset_at(exc);_ps.set_stage(project_id,project_revision,stage_id,'WAITING_CREDITS',input_fingerprint=input_fp,current_job_id=job_id,last_checkpoint=str(_rd) if _rd else None,waiting_reason=str(exc)[:1000],quota_reset_at=reset_at)
            except Exception: pass
            return _park_for_subscription_quota(s,wid,job_id,exc,run_dir=_rd)
        # Claude API 429/503 is availability, not a failed project. ai_router has
        # already exhausted the explicit 5/15/45 second retry ladder.
        try:
            active_provider=_provider_for(project,inp)
        except Exception:
            active_provider=str(job.get('provider_policy') or '')
        if active_provider=='anthropic' and _is_api_capacity_failure(exc):
            try:
                if locals().get('_ps') and project_id and project_revision and stage_id:
                    _ps.set_stage(project_id,project_revision,stage_id,'WAITING_CAPACITY',input_fingerprint=input_fp,current_job_id=job_id,waiting_reason=str(exc)[:1000])
            except Exception: pass
            s.transition(job_id,'WAITING_CAPACITY',error={'error':'CLAUDE_API_CAPACITY','provider':'anthropic','reason':str(exc)[:1800]},message='Claude API je dočasně kapacitně nedostupné; hotové artefakty zůstávají uložené.',force=True)
            return 0
        try:
            if locals().get('_ps') and project_id and project_revision and stage_id:_ps.set_stage(project_id,project_revision,stage_id,'FAILED',input_fingerprint=input_fp,current_job_id=job_id,waiting_reason=str(exc)[:1000])
        except Exception: pass
        s.transition(job_id, 'FAILED', error={'error': str(exc), 'trace': traceback.format_exc(limit=12)})
        return 1
    finally:
        try:
            if locals().get('_ps'):_ps.close()
        except Exception: pass
        try: reset_runtime(_ai_runtime_tokens)
        except Exception: pass
        stop.set()
        th.join(timeout=1)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--job-id', required=True)
    ap.add_argument('--worker-id', default='worker-job')
    args = ap.parse_args()
    return execute(args.job_id, args.worker_id)


if __name__ == '__main__':
    raise SystemExit(main())
